#!/usr/bin/env python3
"""
build_arc_plc.py -- the ARC or PLC decision calculator (2026 and 2027) and its pages.

OFFLINE. Reads committed files only. MACHINE-OWNED OUTPUT: edit this script,
never the files below.

  data/arc-plc.json               what the calculator reads (components/arc-plc.js)
  arc-plc.html                    /arc-plc, the calculator, every key number baked
                                  into the HTML so crawlers see it without JS
  arc-plc/<state>.html            /arc-plc/<state>, FSA's official ARC-CO
                                  benchmark yields for every county in the state
  arc-plc/<state>/<county>.html   /arc-plc/<state>/<county>, one county: short
                                  answers, key numbers, FAQ; "Show the math" is
                                  built on open by components/arc-plc-math.js
  data/arc-plc/math/<ST>.json     what "Show the math" and the counter sheet read
  arc-plc/sheet.html              /arc-plc/sheet?c=<state>/<county>, the one
                                  printable counter sheet (noindex), built in the
                                  browser (components/arc-plc-sheet.js, qr.js)
  arc-plc/<state>/<county>-sheet.html  old sheet URLs: tiny pages that forward
  data/arc-plc/hub/<crop>.json    the crop pages' county lists, filled on open
  components/arc-plc-form.js      the calculator's inputs for the county pages
  embed/arc-plc.html              the embeddable calculator (noindex)
  sitemap-arc-plc.xml             every indexable URL above, lastmod per page
  llms.txt                        its two ARC/PLC lines are rewritten from the data

    python3 scripts/build_arc_plc.py              write everything
    python3 scripts/build_arc_plc.py --check      exit 1 if any output is stale
    python3 scripts/build_arc_plc.py --selftest   hand-checked cases, writes nothing

The build refuses to write when arc-plc/ would pass its size budget or any page
its cap (both read from scripts/check_site_budget.py): bulk goes in data files
the page loads on demand, not in 2,800 copies of a page.

=============================================================================
WHERE EVERY NUMBER COMES FROM (verified 2026-10-09)
=============================================================================
fsa.usda.gov, federalregister.gov, ecfr.gov and law.cornell.edu are not
reachable from the build sandbox. Program rules were checked through search
results quoting those pages, each against a second source. FSA's own data
files are fetched by .github/workflows/fetch-fsa-arcplc.yml into data/fsa/.

1. FSA OFFICIAL COUNTY DATA (data/fsa/arcco_<PY>_data*.xlsx)
   "ARC-CO trend adjusted yields, benchmark yields and guarantee revenues for
   program year PY", one row per county x sub-county x crop x practice
   ("ARC-CO Yield Designation": All, Irrigated or Nonirrigated). Columns: the
   five window years' trend-adjusted yields (county yield or 80% of T-yield),
   the PY benchmark yield (Olympic average), benchmark price, benchmark
   revenue, guarantee, maximum payment rate, and once the year is settled the
   actual county yield, national price, actual revenue and ARC-CO payment rate.
   The newest program year file is the official benchmark shown everywhere.
   Older files give each county's history (benchmark and payment rate by year).
   The benchmark price in the file must equal the one computed below for the
   same year, or the build refuses (two sources disagreeing in front of a reader).

   2027: FSA has not posted 2027 benchmark yields. Building them FSA's way needs
   the 2021-2025 county yields with 80% T-yield plugs, trend-adjusted to 2027
   with RMA's PY2027 trend factors, which are not published. So the 2027 view
   uses FSA's official 2026 benchmark yield, labeled as such ("the 2027
   benchmark shifts the window one year; FSA posts it later"). No NASS
   estimate is used anywhere: a county with no FSA row gets no page and the
   calculator asks for the benchmark from the county office.

2. PROGRAM RULES. Law: One Big Beautiful Bill Act, Pub. L. 119-21 (2025),
   Title I. FSA rule: Federal Register 2026-00313, Jan 12, 2026 [FR26].
   Definitions at 7 CFR 1412.3 [CFR].
   Statutory reference prices 2025-2030 [FR26][ERS]: corn $4.10 (was $3.70),
     soybeans $10.00 (was $8.40), wheat $6.35 (was $5.50); sorghum $4.40,
     barley $5.45, oats $2.65 [ERS][CWA]. DTN (2026-09-17) printed wheat $6.33;
     FR26, ERS and farmdoc say $6.35, which is used. From 2031: +0.5% a year,
     capped at 113% [FR26]. Iowa State CALT called the ERP cap 113%; the eCFR
     and FSA say 115% (113% is the 2031 escalator cap). 115% is used.
   ERP [CFR]: lesser of 115% of statutory, or greater of statutory and (from
     2025) 88% of the Olympic average national MYA price of the 5 most recent
     crop years. Program year PY uses PY-6..PY-2 (2026: 2020-2024; 2027:
     2021-2025): the lag that reproduces every published ERP in --selftest.
   PLC rate = ERP - max(MYA, loan rate), not below 0; paid on 85% of base x
     the farm's PLC yield [FSAF].
   Loan rates 2026 and later [LOAN]: corn $2.42, soybeans $6.82, wheat $3.72.
   ARC-CO [FR26][FD25]: guarantee 90% of benchmark revenue, payment capped at
     12% of benchmark revenue (2025-2031); benchmark price = Olympic average of
     the same 5 MYAs each raised to the program year's ERP (reproduces FSA's
     published 2023, 2024 and 2026 benchmark prices); actual revenue = county
     yield x max(MYA, loan rate); paid on 85% of base [FSAF]. A farm with
     irrigated and nonirrigated county figures is weighted by the farm's
     historical irrigated percentage, not by what it plants [FSA24][APPX].
   Payment limit [ALA][FR26][DTNL]: $155,000 per person or legal entity,
     indexed to CPI-U; USDA put 2025 at $160,000; the rule's own example shows
     $164,000 (factor 1.059). 2026 and 2027 not announced. Members of LLCs and
     S corporations who are actively engaged each get a limit under the 2025
     law [CCF][UMN]. Sequestration: 5.7% on recent payments [SEQ].
   AGI: $900,000 average AGI limit still applies to ARC and PLC; the 75%-of-
     gross-from-farming exception covers certain conservation and disaster
     programs, not ARC/PLC [CALT][CRS].
   10 base acres: farms with 10 or fewer base acres are not paid, except
     socially disadvantaged, limited resource, veteran and beginning farmers [ERS].
   New base acres (up to 30 million, 2019-2023 plantings, 3.69% cut on new
     acres only) take the farm's existing PLC yield for that crop, or the
     county average PLC yield if the farm has none [FD-BASE][FR26][DTN].
   2026 election: unanimous among producers on the farm; one-time for 2026;
     2027-2031 can be changed crop by crop each year [FR26][CFR71]. Multi-year
     contract 2026-2031 stays valid while farm records and producers do not
     change [FR26 via search].
   Rented ground [CFR54][TAMU]: cash lease, the tenant gets the payment unless
     agreed otherwise; share lease, landlord and tenant share it and both sign.
   SCO [RMA25][HOEV][RMA27]: under the 2025 law SCO can be bought whatever the
     ARC/PLC election, from policies with sales closing on or after July 1,
     2025; for 2026 the 90% band is sold as ECO, SCO goes to 90% in 2027;
     premium support 80%. RMA's 27-SCO summary lists removing the ARC
     restriction as a 2027 endorsement change, so the page says to confirm
     2026 with the agent.
   Enrollment [USDA0915 via DTN]: 2026 Sept 16 to Dec 11, 2026; 2027 Nov 2,
     2026 to Mar 15, 2027. No 2026 election: the 2025 election carries over
     and the farm gets no 2026 payment.

3. PRICES: data/nass/{corn,soy,wheat}-price-my.json, national MYA, final only
   (scope string must say so). 2025 corn ($4.16) and soybeans ($10.50) came
   out at the end of Sept 2026; corn 2025 is the low year and is dropped from
   the 2027 ERP.

Source URLs: SRC below, printed on the page. Stdlib + openpyxl.
"""
import argparse
import datetime as dt
import glob
import html
import io
import json
import math
import os
import re
import sys
import zipfile
from decimal import Decimal, ROUND_HALF_UP

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import arc_plc_crops as XC  # noqa: E402  every covered crop beyond corn, soybeans and wheat

SITE = "https://agsist.com"
YEARS = (2026, 2027)
OUT_JSON = "data/arc-plc.json"
OUT_STATE_JSON = "data/arc-plc"
OUT_MAIN = "arc-plc.html"
OUT_DIR = "arc-plc"
OUT_EMBED = "embed/arc-plc.html"
OUT_SITEMAP = "sitemap-arc-plc.xml"
FSA_DIR = "data/fsa"
GEO_SRC = "data/atlas/counties.geo.json"   # county shapes the Farmland Atlas already ships; cut per state for Use my location
FINAL_SCOPE = "final only"
STYLES_V = "23"
LOADER_V = "18"
ASOF_V = "1"
CALC_V = "8"   # 8: Show the math built on open, one shared counter sheet, slimmer pages
MIN_STATE_COUNTIES = 3
SCEN_YEARS = list(range(2015, 2026))   # years with both an FSA county yield and a price change; 2025 counts once its final MYA is out
LEAN_GAP = 3.00       # $/base acre: "Leans X" needs at least this expected gap, at least 1 SE, and MIN_WINS winning years
MIN_WINS = 3          # the leading program must have paid more in at least this many scenario years, for Pick or Leans
SNAPSHOT = "data/arc-plc-prices.json"  # futures snapshot the scenarios are centered on (--snapshot-prices)
SNAP_MAX_DAYS = 10    # a build refuses to write when the snapshot's futures close is older than this (weekly refresh + a long weekend)
SNAP_FRESH_DAYS = 4   # --snapshot-prices refuses data/prices.json quotes older than this (a holiday weekend)
BUILD_DATE = None     # date the signup copy is written for; main() sets it (today, UTC, or --today YYYY-MM-DD)
RMA_HIST_PAGE = "harvest-price-tracker.html"  # RMA projected and harvest prices 2011-2025, one copy on the site
MIN_SCEN = 8          # never recommend from fewer scenario years
MIN_GAP = 2.00        # $/base acre: "Pick" needs at least this expected gap, and more than t x SE
# two-sided 95% Student t by degrees of freedom (n - 1): Pick needs the gap above t x SE
T975 = {1: 12.706, 2: 4.303, 3: 3.182, 4: 2.776, 5: 2.571, 6: 2.447, 7: 2.365, 8: 2.306, 9: 2.262, 10: 2.228,
        11: 2.201, 12: 2.179, 13: 2.160, 14: 2.145, 15: 2.131, 16: 2.120, 17: 2.110, 18: 2.101, 19: 2.093, 20: 2.086,
        21: 2.080, 22: 2.074, 23: 2.069, 24: 2.064, 25: 2.060, 26: 2.056, 27: 2.052, 28: 2.048, 29: 2.045, 30: 2.042}
LEVELS = (("lower", 0.85), ("same", 1.00), ("higher", 1.15))   # the simple page's three PLC-yield buttons, x county average
MIN_STATE_N = 3       # counties a state-year needs to stand in for a county

# ---------------------------------------------------------------- sources
SRC = {
    "FSADATA": ("FSA ARC/PLC program data (official county benchmark files)",
                "https://www.fsa.usda.gov/resources/programs/arc-plc/program-data"),
    "FR26": ("Federal Register 2026-00313, Jan 12, 2026 (FSA rule for the 2025 law)",
             "https://www.federalregister.gov/documents/2026/01/12/2026-00313/changes-to-agriculture-risk-coverage-price-loss-coverage-and-dairy-margin-coverage-programs"),
    "CFR": ("7 CFR 1412.3, definitions (eCFR)",
            "https://www.ecfr.gov/current/title-7/subtitle-B/chapter-XIV/subchapter-B/part-1412/subpart-A"),
    "CFR54": ("7 CFR 1412.54, leases and division of payment (eCFR)",
              "https://www.ecfr.gov/current/title-7/subtitle-B/chapter-XIV/subchapter-B/part-1412/subpart-E/section-1412.54"),
    "ERS": ("USDA ERS, Title I crop commodity program provisions",
            "https://www.ers.usda.gov/topics/farm-economy/farm-commodity-policy/title-i-crop-commodity-program-provisions"),
    "FSAF": ("FSA ARC and PLC fact sheet, Sept 2025",
             "https://www.fsa.usda.gov/sites/default/files/2025-09/FSA_ARC%20&%20PLC_3pg_Fact%20Sheet-SEPT%202025_final.pdf"),
    "FSA24": ("FSA ARC and PLC fact sheet, Dec 2024 (irrigated weighting)",
              "https://www.fsa.usda.gov/sites/default/files/2024-12/fsa_arc_plc_factsheet_1223.pdf"),
    "LOAN25": ("FSA, 2025 national average loan rates",
               "https://www.fsa.usda.gov/sites/default/files/2025-04/2025-National-Average-Loan-Rates_0.pdf"),
    "LOAN": ("FSA, 2026 marketing assistance loan rates (Apr 8, 2026)",
             "https://www.fsa.usda.gov/news-events/news/04-08-2026/usda-announces-2026-marketing-assistance-loan-rates-wheat-feed-grains"),
    "APPX": ("FSA, ARC/PLC contract appendix (benchmark yield rules)",
             "https://www.fsa.usda.gov/sites/default/files/documents/appendix-to-arc_plc-contracts.pdf"),
    "FD24": ("farmdoc daily, Feb 2024 (2024 ERPs and benchmark prices)",
             "https://farmdocdaily.illinois.edu/2024/02/estimated-likelihoods-of-plc-and-arc-co-payments-for-2024.html"),
    "FD25": ("farmdoc daily, Jul 2025 (2025 law, commodity title)",
             "https://farmdocdaily.illinois.edu/2025/07/impacts-of-the-commodity-title-changes-under-the-one-big-beautiful-bill-act-obbba-for-midwestern-farms-in-2025.html"),
    "FDBASE": ("farmdoc daily, Jul 2025 (new base acre provisions)",
               "https://farmdocdaily.illinois.edu/2025/07/the-new-base-acre-provisions-in-the-2025-farm-bill.html"),
    "HOFF15": ("Hoffman, Etienne, Irwin, Colino and Toasa (2015), Forecast performance of WASDE price projections for U.S. corn",
               "https://www.researchgate.net/publication/282941394_Forecast_performance_of_WASDE_price_projections_for_US_corn"),
    "WASDE": ("USDA WASDE-676, Oct 9, 2026",
              "https://www.usda.gov/about-usda/general-information/staff-offices/office-chief-economist/commodity-markets/wasde-report"),
    "AFBF26": ("American Farm Bureau Market Intel (2026 ERPs and benchmark prices)",
               "https://www.fb.org/market-intel/risk-management-options-for-2026-corn-soybeans-and-wheat"),
    "ALA": ("Alabama Cooperative Extension, Title I program changes",
            "https://www.aces.edu/blog/topics/crop-production/title-i-commodity-program-changes/"),
    "DTNL": ("DTN, Sept 2, 2026: 2025 payment limit $160,000",
             "https://www.dtnpf.com/agriculture/web/ag/blogs/ag-policy-blog/blog-post/2026/09/02/usda-official-arc-plc-payments-will"),
    "CCF": ("Center for Commercial Agriculture: LLC and S corp payment limits",
            "https://ccf.us/a-big-change/"),
    "CALT": ("Iowa State CALT: new payment limitation and eligibility rules",
             "https://www.calt.iastate.edu/post/usda-issues-new-payment-limitation-and-eligibility-rules"),
    "SEQ": ("Farm CPA Report, sequestration and the 2025 limit",
            "https://www.farmcpareport.com/p/correction-the-2025-arcplc-limit"),
    "TAMU": ("Texas A&M AgriLife, lease payment structures",
             "https://agrilife.org/texasaglaw/2019/03/25/common-agricultural-lease-payment-structures/"),
    "RMA25": ("RMA bulletin MGR-25-006 (2025 law crop insurance changes)",
              "https://www.rma.usda.gov/policy-procedure/bulletins-memos/managers-bulletin/mgr-25-006-one-big-beautiful-bill-act-amendment"),
    "HOEV": ("Sen. Hoeven: RMA implements 2025 law (ECO 90% for 2026)",
             "https://www.hoeven.senate.gov/newsroom/press-releases/hoeven-rma-follows-through-on-one-big-beautiful-bill-implements-enhanced-crop-insurance"),
    "RMA27": ("RMA, 2027 SCO endorsement summary of changes",
              "https://www.rma.usda.gov/sites/default/files/2026-06/SCO%20Endorsement%2027-SCO.pdf"),
    "USDA0915": ("USDA, Sept 15, 2026: 2026 and 2027 ARC/PLC enrollment",
                 "https://www.fsa.usda.gov/news-events/news/09-15-2026/usda-announces-2026-2027-enrollment-key-price-revenue-safety-net"),
    "DTN": ("DTN, Sept 17, 2026: USDA opens ARC/PLC enrollment",
            "https://www.dtnpf.com/agriculture/web/ag/crops/article/2026/09/17/usda-opens-arcplc-enrollment-2026"),
    "CWA": ("Colorado Wheat, 2025 law provisions table",
            "https://coloradowheat.org/wp-content/uploads/FINAL-OBBBA-Agriculture-Provisions-06.09.2025-MA.pdf"),
}
FSA_PAGE = "https://www.fsa.usda.gov/resources/programs/arc-plc"
FSA_DATA = SRC["FSADATA"][1]
OFFICE = "https://offices.sc.egov.usda.gov/locator/app"

PARAMS = {
    "erp_pct": 0.88, "erp_cap_pct": 1.15,
    "arc_guarantee_pct": 0.90, "arc_max_pct": 0.12, "payment_acres_pct": 0.85,
    "payment_limit_base": 155000, "payment_limit_2025": 160000, "sequestration_recent_pct": 0.057,
}
YEAR_INFO = {
    2026: {"open": "2026-09-16", "close": "2026-12-11", "pay_after": "2027-10-01",
           "fut": {"corn": ("corn-dec26", "Dec 2026 corn futures"), "soybeans": ("beans-nov26", "Nov 2026 soybean futures"),
                   "wheat": ("wheat-dec26", "Dec 2026 Chicago wheat futures")},
           "my": {"corn": "Sept 2026 to Aug 2027", "soybeans": "Sept 2026 to Aug 2027", "wheat": "June 2026 to May 2027"}},
    2027: {"open": "2026-11-02", "close": "2027-03-15", "pay_after": "2028-10-01",
           "fut": {"corn": ("corn-dec27", "Dec 2027 corn futures"), "soybeans": ("beans-nov27", "Nov 2027 soybean futures"),
                   "wheat": ("wheat-jul27", "Jul 2027 Chicago wheat futures")},
           "my": {"corn": "Sept 2027 to Aug 2028", "soybeans": "Sept 2027 to Aug 2028", "wheat": "June 2027 to May 2028"}},
}
CROPS = {
    "corn": {"label": "Corn", "statutory": 4.10, "loan": 2.42, "mya": "corn-price-my", "fsa": "Corn",
             "grid": (3.20, 5.20, 0.20), "prior_statutory": 3.70},
    "soybeans": {"label": "Soybeans", "statutory": 10.00, "loan": 6.82, "mya": "soy-price-my", "fsa": "Soybeans",
                 "grid": (8.00, 13.00, 0.50), "prior_statutory": 8.40},
    "wheat": {"label": "Wheat", "statutory": 6.35, "loan": 3.72, "mya": "wheat-price-my", "fsa": "Wheat",
              "grid": (4.50, 7.50, 0.30), "prior_statutory": 5.50},
}
CROP_LC = {"corn": "corn", "soybeans": "soybeans", "wheat": "wheat"}
CROP_LC.update({k: v["lc"] for k, v in XC.XCROPS.items()})
ALL = list(CROPS) + list(XC.XCROPS)   # every crop with county data, corn, soybeans and wheat first
# 2018 law (crop years 2019-2024): statutory price, 85% escalator; loan rates 2019-2025 [ERS][LOAN25]
LAW2018 = {"corn": (3.70, 2.20), "soybeans": (8.40, 6.20), "wheat": (5.50, 3.38)}
BACKTEST_YEARS = list(range(2019, 2026))
OTHER_STATUTORY = [("Sorghum", 4.40, 3.95, 2.42), ("Barley", 5.45, 4.95, 2.75), ("Oats", 2.65, 2.40, 2.20)]
DESIG = {"all": "all", "irrigated": "irr", "nonirrigated": "non", "non-irrigated": "non"}
DLABEL = {"all": "All practices", "irr": "Irrigated", "non": "Non-irrigated"}
PRAC = {"all": "All", "irr": "Irr.", "non": "Non-irr."}   # the state tables' short practice column

STATE_NAMES = {
    "AL": "Alabama", "AK": "Alaska", "AZ": "Arizona", "AR": "Arkansas", "CA": "California",
    "CO": "Colorado", "CT": "Connecticut", "DE": "Delaware", "FL": "Florida", "GA": "Georgia",
    "HI": "Hawaii", "ID": "Idaho", "IL": "Illinois", "IN": "Indiana", "IA": "Iowa",
    "KS": "Kansas", "KY": "Kentucky", "LA": "Louisiana", "ME": "Maine", "MD": "Maryland",
    "MA": "Massachusetts", "MI": "Michigan", "MN": "Minnesota", "MS": "Mississippi",
    "MO": "Missouri", "MT": "Montana", "NE": "Nebraska", "NV": "Nevada", "NH": "New Hampshire",
    "NJ": "New Jersey", "NM": "New Mexico", "NY": "New York", "NC": "North Carolina",
    "ND": "North Dakota", "OH": "Ohio", "OK": "Oklahoma", "OR": "Oregon", "PA": "Pennsylvania",
    "RI": "Rhode Island", "SC": "South Carolina", "SD": "South Dakota", "TN": "Tennessee",
    "TX": "Texas", "UT": "Utah", "VT": "Vermont", "VA": "Virginia", "WA": "Washington",
    "WV": "West Virginia", "WI": "Wisconsin", "WY": "Wyoming",
}
FIPS_ST = {"01": "AL", "02": "AK", "04": "AZ", "05": "AR", "06": "CA", "08": "CO", "09": "CT", "10": "DE", "12": "FL",
           "13": "GA", "15": "HI", "16": "ID", "17": "IL", "18": "IN", "19": "IA", "20": "KS", "21": "KY", "22": "LA",
           "23": "ME", "24": "MD", "25": "MA", "26": "MI", "27": "MN", "28": "MS", "29": "MO", "30": "MT", "31": "NE",
           "32": "NV", "33": "NH", "34": "NJ", "35": "NM", "36": "NY", "37": "NC", "38": "ND", "39": "OH", "40": "OK",
           "41": "OR", "42": "PA", "44": "RI", "45": "SC", "46": "SD", "47": "TN", "48": "TX", "49": "UT", "50": "VT",
           "51": "VA", "53": "WA", "54": "WV", "55": "WI", "56": "WY"}
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


# ---------------------------------------------------------------- math
def rnd(v, k=2):
    """Half away from zero. round(v, 9) first: 4.10 x 1.15 is 4.7149999... in binary."""
    d = Decimal(repr(round(float(v), 9))).quantize(Decimal(1).scaleb(-k), rounding=ROUND_HALF_UP)
    return float(d)


def cents(x):
    """Dollars -> whole cents, half up, the way JS Math.round(x * 100) does for x >= 0."""
    return int(math.floor(x * 100 + 0.5))


def olympic(vals):
    if len(vals) != 5:
        raise ValueError("an Olympic average needs exactly 5 values")
    s = sorted(vals)
    return (s[1] + s[2] + s[3]) / 3.0, s[0], s[4]


def erp_calc(myas, statutory, pct=0.88, cap_pct=1.15):
    avg, lo, hi = olympic(myas)
    pv = rnd(avg * pct)
    cap = rnd(statutory * cap_pct)
    erp = min(cap, max(statutory, pv))
    binding = "cap" if erp == cap and pv > cap else ("formula" if pv > statutory else "statutory")
    return {"olympic_avg": rnd(avg, 4), "dropped_low": lo, "dropped_high": hi, "pct": pct,
            "pct_value": pv, "statutory": statutory, "cap_pct": cap_pct, "cap": cap, "erp": rnd(erp), "binding": binding}


def benchmark_price(myas, erp):
    used = [max(m, erp) for m in myas]
    avg, lo, hi = olympic(used)
    return rnd(avg), used, lo, hi


def plc_rate(erp, mya, loan):
    return max(0.0, erp - max(mya, loan))


def arc_rate(by, bp, y, mya, loan, g=0.90, cap=0.12):
    br = by * bp
    return min(max(0.0, g * br - y * max(mya, loan)), cap * br)


def per_base(rate_value, pa=0.85):
    return rate_value * pa


def ranges(erp, loan, plc_y, by, bp, y, g=0.90, cap=0.12, pa=0.85, scale=100):
    """Which program pays more at each tick of season-average price at county
    yield y. Runs [{w, lo, hi}] in ticks (1/scale dollars: cents for bushel
    crops, scale 10000 for FSA's pound crops), inclusive. The scan stops at
    max(ERP, benchmark price) x 1.5 so no input can make it run long.
    Mirrors AgArcPlc.ranges in components/arc-plc.js (test/arc-plc.test.mjs)."""
    lo_c = int(math.floor(loan * scale + 0.5))
    trig = (g * by * bp / y) if y > 0 else erp
    hi_c = min(int(math.ceil(max(erp, trig) * scale)) + 1, int(math.ceil(max(erp, bp) * (1.5 * scale))))
    runs = []
    for c in range(lo_c, max(hi_c, lo_c) + 1):
        p = c / float(scale)
        a = cents(per_base(arc_rate(by, bp, y, p, loan, g, cap), pa))
        b = cents(per_base(plc_rate(erp, p, loan) * plc_y, pa))
        w = "none" if a == 0 and b == 0 else ("same" if a == b else ("plc" if b > a else "arc"))
        if runs and runs[-1]["w"] == w:
            runs[-1]["hi"] = c
        else:
            runs.append({"w": w, "lo": c, "hi": c})
    return runs


def plc_below(runs):
    return runs[0]["hi"] if runs and runs[0]["w"] == "plc" else None


# ---------------------------------------------------------------- scenario engine
def tcrit(df):
    """Two-sided 95% Student t for df degrees of freedom (1.96 past 30)."""
    return T975.get(df, 1.96 if df > 30 else T975[1])


def scen_call(mp, ma, md, se, pw, aw, n, borrowed=False):
    """The call from the scenario statistics. Mirrors call() in components/arc-plc.js.

    Both expected payments round to $0 (whole dollars): "none". Pick: gap more
    than t x SE (t for n - 1 df), at least MIN_GAP, and the leader paid more in
    at least MIN_WINS years. Leans: gap at least LEAN_GAP and at least 1 SE,
    leader won at least MIN_WINS years. Else "close". A borrowed price spread
    (a crop with no October futures history of its own) caps the call at Leans."""
    if n < MIN_SCEN:
        return "withheld"
    if cents(mp) < 50 and cents(ma) < 50:
        return "none"
    side = "plc" if md > 0 else "arc"
    lw = pw if side == "plc" else aw
    g = abs(md)
    if g > tcrit(n - 1) * se and g >= MIN_GAP and lw >= MIN_WINS:
        v = side
    elif g >= LEAN_GAP and g >= se and lw >= MIN_WINS:
        v = "lean_" + side
    else:
        v = "close"
    if borrowed and v in ("plc", "arc"):
        v = "lean_" + v
    return v


def scen_eval(erp, bp, loan, py, parts, center, ratios, price=None, borrowed=False, alt=None):
    """Expected PLC and ARC-CO per base acre across past years.

    parts: [(weight, benchmark_yield, dy)], dy either {year: county yield /
    benchmark} or one number for every year (2026: the crop is harvested, so a
    normal crop, 1.0, or the county yield the reader typed / benchmark).
    ratios: {year: price multiplier}. Scenario year t: price = center x r_t
    (or the reader's own price), county yield = benchmark x d_t, the same real
    year for both. borrowed: the spread is not the crop's own (capped at Leans).
    alt: a second center (futures-implied); where the two centers disagree on
    the leader or on Pick, Pick becomes Leans.
    Mirrors AgArcPlc.scenarios in components/arc-plc.js."""
    def dyv(dy, t):
        return dy if isinstance(dy, (int, float)) else dy[t]
    yrs = [t for t in sorted(ratios) if all(isinstance(p[2], (int, float)) or t in p[2] for p in parts)]
    rows = []
    for t in yrs:
        p = price if price is not None else center * ratios[t]
        plc = per_base(plc_rate(erp, p, loan) * py)
        arc = sum(w * per_base(arc_rate(by, bp, by * dyv(dy, t), p, loan)) for w, by, dy in parts)
        rows.append({"t": t, "p": p, "plc": plc, "arc": arc})
    n = len(rows)
    out = {"n": n, "years": yrs, "rows": rows}
    if n == 0:
        out["verdict"] = "withheld"
        return out
    mp = sum(r["plc"] for r in rows) / n
    ma = sum(r["arc"] for r in rows) / n
    diffs = [r["plc"] - r["arc"] for r in rows]
    md = mp - ma
    sd = math.sqrt(sum((x - md) ** 2 for x in diffs) / (n - 1)) if n > 1 else 0.0
    se = sd / math.sqrt(n)
    pw = sum(1 for r in rows if cents(r["plc"]) > cents(r["arc"]))
    aw = sum(1 for r in rows if cents(r["arc"]) > cents(r["plc"]))
    out.update({"plc": mp, "arc": ma, "diff": md, "se": se, "plc_wins": pw, "arc_wins": aw, "ties": n - pw - aw,
                "t": tcrit(n - 1) if n > 1 else None})
    fixed = price is not None
    out["verdict"] = scen_call(mp, ma, md, se, pw, aw, n, borrowed and not fixed)
    if borrowed and not fixed:
        out["borrowed"] = True
    if alt is not None and not fixed and out["verdict"] in ("plc", "arc"):
        a = scen_eval(erp, bp, loan, py, parts, alt, ratios)
        out["alt"] = {k: a[k] for k in ("plc", "arc", "diff", "se", "verdict") if k in a}
        lead_a = "plc" if a.get("diff", 0) > 0 else "arc"
        if lead_a != out["verdict"] or a["verdict"] != out["verdict"]:
            out["verdict"] = "lean_" + out["verdict"]
            out["alt_down"] = True
    elif alt is not None and not fixed:
        a = scen_eval(erp, bp, loan, py, parts, alt, ratios)
        out["alt"] = {k: a[k] for k in ("plc", "arc", "diff", "se", "verdict") if k in a}
    return out


VWORD = {"plc": "Pick PLC", "arc": "Pick ARC-CO", "lean_plc": "Leans PLC", "lean_arc": "Leans ARC-CO",
         "close": "Close, either one", "none": "Neither expected to pay", "withheld": "Not enough history to say",
         "mismatch": "No typical-farm answer", "noprice": "No 2027 price yet", "pending": "Waiting on USDA's closing 2025 price"}
NOCALL = ("withheld", "mismatch", "noprice", "pending")   # no answer, for the reason in the words


def vclass(v):
    """CSS class for a verdict: pick, lean, close and no-answer each look different."""
    return "ap-w-" + {"plc": "plc", "arc": "arc", "lean_plc": "lean-plc", "lean_arc": "lean-arc", "close": "close",
                      "none": "neither"}.get(v, "withheld")


def verdict_line(v):
    """One plain sentence, numbers beside the verdict."""
    w_ = v["verdict"]
    if w_ == "withheld":
        return f"Not enough history to say: {v['n']} past years with both a price and a county yield; we need {MIN_SCEN}."
    if w_ == "mismatch":
        return "No typical-farm answer: FSA&rsquo;s county average PLC yield is above this benchmark yield. Run your own PLC yield."
    if w_ == "noprice":
        return ("No 2027 price yet: there is no 2027 price outlook for this crop we can check against past years, so no call. "
                "Type a price in the calculator to see the dollars.")
    if w_ == "pending":
        return "Waiting on USDA&rsquo;s closing 2025 season-average price, which sets the 2027 reference price."
    nums = (f"expected {usd_pay(v['plc'])} PLC vs {usd_pay(v['arc'])} ARC-CO per base acre; PLC paid more in {v['plc_wins']} of "
            f"{v['n']} past-year scenarios, ARC-CO in {v['arc_wins']}")
    if w_ == "none":
        return f"Neither expected to pay ({nums})."
    if w_ == "close":
        return f"Close, either one: about {usd(abs(v['diff']))} apart, too small or too uncertain to call ({nums})."
    note = (" Capped at Leans: this crop borrows corn&rsquo;s price swings." if v.get("borrowed") and w_.startswith("lean_") else "") + \
           (" Not a clear pick: a futures-based price start gives a different call." if v.get("alt_down") else "")
    if w_.startswith("lean_"):
        side = w_[5:]
        other = "arc" if side == "plc" else "plc"
        ow = v["arc_wins"] if side == "plc" else v["plc_wins"]
        tail_ = (f"{WORD[other]} did not pay more in any of the {v['n']} years, but the gap is within the normal swing" if ow == 0 else
                 f"but the gap is within the normal swing, so {WORD[other]} is a defensible choice")
        return f"Leans {WORD[side]}: expected about {usd(abs(v['diff']))} more per base acre; {tail_} ({nums}).{note}"
    return f"Pick {WORD[w_]} ({nums}).{note}"


# ---------------------------------------------------------------- formatting
def esc(s):
    return html.escape("" if s is None else str(s), quote=True)


def usd(v, k=2):
    return f"${rnd(v, k):,.{k}f}"


def usd_pay(v):
    """An expected payment: '$0' when it rounds to nothing, else dollars and cents."""
    return "$0" if cents(v) == 0 else usd(v)


def c2(c, scale=100):
    return f"${c / scale:.2f}" if scale == 100 else f"${c / scale:.4f}"


def bu(v, k=1):
    return f"{rnd(v, k):.{k}f}"


def nice_date(iso, year=True):
    d = dt.date.fromisoformat(iso)
    return f"{MON[d.month - 1]} {d.day}" + (f", {d.year}" if year else "")


def slugify(s):
    import unicodedata
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.lower().replace("&", " and ").replace("'", "").replace("’", "")
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s or "x"


WORD = {"plc": "PLC", "arc": "ARC-CO"}


def ranges_text(runs, scale=100):
    out = []
    n = len(runs)
    for i, r in enumerate(runs):
        lo, hi = c2(r["lo"], scale), c2(r["hi"], scale)
        if r["w"] == "none":
            out.append(f"Neither pays at {lo} or higher." if i == n - 1 else (f"Neither pays at {lo}." if r["lo"] == r["hi"] else f"Neither pays from {lo} to {hi}."))
            continue
        who = "Both pay the same" if r["w"] == "same" else f"{WORD[r['w']]} pays more"
        if i == 0:
            out.append(f"{who} at {hi} or lower.")
        elif i == n - 1:
            out.append(f"{who} at {lo} or higher.")
        else:
            out.append(f"{who} at {lo}." if r["lo"] == r["hi"] else f"{who} from {lo} to {hi}.")
    return out


# ---------------------------------------------------------------- FSA files
def load_xlsx(path):
    """openpyxl, tolerating the date-typed document properties FSA's files carry."""
    import openpyxl
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except TypeError:
        zin = zipfile.ZipFile(path)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zout:
            for it in zin.infolist():
                data = zin.read(it.filename)
                if it.filename == "docProps/core.xml":
                    data = re.sub(rb"<dcterms:(created|modified)[^>]*>.*?</dcterms:\1>", b"", data, flags=re.S)
                zout.writestr(it, data)
        buf.seek(0)
        return openpyxl.load_workbook(buf, read_only=True, data_only=True)


def _num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def parse_arcco(path):
    """One FSA 'ARC-CO trend adjusted yields, benchmark yields and guarantee
    revenues' workbook -> {py, as_of, file, rows: {(fips, sub, crop, d): {...}}}."""
    wb = load_xlsx(path)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and str(r[0] or "").strip() == "ST_Cty")
    hdr = [str(h or "").strip() for h in rows[hi]]
    low = [h.lower() for h in hdr]

    def col(pat, exclude=None):
        for i, h in enumerate(low):
            if re.search(pat, h) and not (exclude and re.search(exclude, h)):
                return i
        return None
    years = [(int(m.group(1)), i) for i, h in enumerate(low) if (m := re.match(r"(\d{4}) trend adjusted", h))]
    ib = col(r"bench ?mark(?! price)(?! revenue)", r"price|revenue")
    m = re.match(r"(\d{4})", hdr[ib] if ib is not None else "")
    py = int(m.group(1)) if m else None
    idx = {"by": ib, "bp": col(r"bench ?mark price"), "br": col(r"benchmark revenue"), "g": col(r"guarantee revenue"),
           "max": col(r"maximum payment rate"), "ay": col(r"actual yield"), "np": col(r"national price"),
           "ar": col(r"actual revenue"), "pay": col(r"arc-co payment rate")}
    ic, isub, icrop, ides = low.index("st_cty"), col(r"sub county"), low.index("crop name"), col(r"yield designation")
    icn, isn = col(r"county name"), col(r"state name")
    asof = None
    for src in (ws.title, os.path.basename(path)):
        mm = re.search(r"(20\d\d)[-_ ](\d\d)[-_ ](\d\d)", src)
        if mm:
            asof = f"{mm.group(1)}-{mm.group(2)}-{mm.group(3)}"
            break
    out = {}
    for r in rows[hi + 1:]:
        if not r or not r[ic]:
            continue
        fips = str(r[ic]).strip().zfill(5)
        crop = str(r[icrop] or "").strip()
        d = DESIG.get(str(r[ides] or "").strip().lower())
        if not d or not re.fullmatch(r"\d{5}", fips):
            continue
        sub = str(r[isub] or "").strip() if isub is not None else ""
        rec = {k: _num(r[i]) if i is not None else None for k, i in idx.items()}
        rec["yrs"] = {y: _num(r[i]) for y, i in years}
        rec["county"] = str(r[icn] or "").strip() if icn is not None else ""
        rec["state"] = str(r[isn] or "").strip() if isn is not None else ""
        out[(fips, sub, crop, d)] = rec
    return {"py": py, "as_of": asof, "file": os.path.relpath(path), "rows": out, "years": [y for y, _ in years]}


def load_fsa(root="."):
    files = {}
    for p in sorted(glob.glob(os.path.join(root, FSA_DIR, "arcco_*.xls*"))):
        try:
            f = parse_arcco(p)
        except Exception as e:  # pragma: no cover
            print(f"[arc-plc] WARN could not read {p}: {type(e).__name__}: {e}", file=sys.stderr)
            continue
        if f["py"] and (f["py"] not in files or (f["as_of"] or "") > (files[f["py"]]["as_of"] or "")):
            files[f["py"]] = f
    return files


def load_plc_county(root="."):
    """FSA average PLC yield by county (enrolled base), newest program year file."""
    best = None
    for p in sorted(glob.glob(os.path.join(root, FSA_DIR, "plc_county_yields_py*.xls*"))):
        m = re.search(r"py(\d{4})", p)
        if m and (best is None or int(m.group(1)) > best[0]):
            best = (int(m.group(1)), p)
    if not best:
        return None
    wb = load_xlsx(best[1])
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and "PLC Yield" in [str(x or "").strip() for x in r])
    hdr = [str(h or "").strip() for h in rows[hi]]
    ist, ico, icr, iy = hdr.index("State"), hdr.index("County"), hdr.index("Crop Name"), hdr.index("PLC Yield")
    want = {"CORN": "corn", "SOYBEANS": "soybeans", "WHEAT": "wheat", **XC.PLC_NAMES}
    out, base = {}, {}
    for r in rows[hi + 1:]:
        if not r or r[ist] is None:
            continue
        crop = want.get(str(r[icr] or "").strip().upper())
        y = _num(r[iy])
        if crop and y:
            f = str(r[ist]).zfill(2) + str(r[ico]).zfill(3)
            out.setdefault(f, {})[crop] = round(y, 1)
            b = _num(r[hdr.index("ENROLLED BASE (PLC+ARC_CO+ARC_IC)")]) if "ENROLLED BASE (PLC+ARC_CO+ARC_IC)" in hdr else None
            if b:
                base.setdefault(f, {})[crop] = b
    return {"py": best[0], "file": os.path.relpath(best[1]), "yields": out, "base": base}


def load_arc_tables(root="."):
    """FSA's national ARC-CO tables (data/fsa/<PY>_ARC_CO.xlsx, Table 4) -> {py: {commodity: (actual ARC-CO price, F or P, loan rate)}}.
    The 2026 table only feeds arc_plc_crops; here the settled years check FSA's county payment rates."""
    out = {}
    for p in sorted(glob.glob(os.path.join(root, FSA_DIR, "20[0-9][0-9]_ARC_CO*.xls*"))):
        py = int(os.path.basename(p)[:4])
        rows = list(load_xlsx(p).worksheets[0].iter_rows(values_only=True))
        hdr = next((r for r in rows if r and any("National Loan Rate" in str(x or "") for x in r)), None)
        if not hdr:
            continue
        i_n = next(i for i, h in enumerate(hdr) if re.search(r"Actual ARC-CO Price", str(h or "")))
        i_l = next(i for i, h in enumerate(hdr) if "National Loan Rate" in str(h or ""))
        t = {}
        for r in rows:
            nm = XC._name((r or [None])[0])
            if nm and _num(r[i_n]) is not None and _num(r[i_l]) is not None:
                t[nm] = (_num(r[i_n]), str(r[i_n + 1] or "").strip(), _num(r[i_l]))
        if t:
            out[py] = {"rows": t, "file": os.path.relpath(p, root)}
    return out


def check_fsa_pay(fsa, arc_tabs):
    """Every settled county row in FSA's ARC-CO files: the payment rate recomputed from
    the row's own benchmark yield, benchmark price, actual yield and national price
    (guarantee 86% and cap 10% through 2024, 90% and 12% from 2025) must equal FSA's
    printed rate to the cent, and the row's national price must equal FSA's national
    table for that year. Refuses to publish otherwise. -> {py: rows checked}."""
    names = {v["fsa"]: k for k, v in CROPS.items()}
    names.update({v["fsa"]: v["table"] for v in XC.XCROPS.values()})
    seen = {}
    for py, f in fsa.items():
        tab = (arc_tabs.get(py) or {}).get("rows", {})
        g_pct, cap = (0.86, 0.10) if py <= 2024 else (0.90, 0.12)
        for key, r in f["rows"].items():
            if r["pay"] is None or r["ay"] is None or r["by"] is None or r["bp"] is None or r["np"] is None:
                continue
            br = rnd(r["by"] * r["bp"])
            comp = rnd(min(max(0.0, rnd(g_pct * br) - rnd(r["ay"] * r["np"])), rnd(cap * br)))
            if abs(comp - r["pay"]) > 0.011:
                raise SystemExit(f"[arc-plc] FSA {py} ARC-CO payment rate for {key} is {r['pay']}, recomputed {comp}: refusing to publish two answers")
            tn = names.get(key[2])
            tn = {"corn": "Corn", "soybeans": "Soybeans", "wheat": "Wheat"}.get(tn, tn)
            if tab and tn in tab and abs(tab[tn][0] - r["np"]) > 0.00005:
                raise SystemExit(f"[arc-plc] FSA {py} county file price for {key[2]} is {r['np']}, its national table says {tab[tn][0]}: refusing")
            seen[py] = seen.get(py, 0) + 1
    return seen


def load_fsa_tables(root="."):
    """FSA's national tables: projected 2026/27 MYA (USDA's projection) and
    the official 2026 ERP. -> {date, proj: {crop: price}, erp: {crop: price}} or None."""
    out = {"proj": {}, "erp": {}, "date": None}
    want = {v["fsa"]: k for k, v in CROPS.items()}
    for fn, key in (("2026_MYA", "proj"), ("2026_ERP", "erp")):
        ps = glob.glob(os.path.join(root, FSA_DIR, fn + "*.xls*"))
        if not ps:
            continue
        rows = list(load_xlsx(sorted(ps)[-1]).worksheets[0].iter_rows(values_only=True))
        for r in rows[:4]:
            for v in r or ():
                m = re.match(r"([A-Z][a-z]+) (\d{1,2}), (20\d\d)", str(v or ""))
                if m and not out["date"]:
                    out["date"] = dt.datetime.strptime(m.group(0), "%B %d, %Y").date().isoformat()
        hdr = next((r for r in rows if r and r[0] == "Commodity"), None)
        for r in rows:
            k = want.get(str((r or [None])[0] or "").strip())
            if not k:
                continue
            if key == "proj":
                i = next(i for i, h in enumerate(hdr) if h and "Projected" in str(h) and "2026/27 MYA" in str(h))
                if str(r[i + 1] or "").strip() == "P" and _num(r[i]):
                    out["proj"][k] = _num(r[i])
            else:
                out["erp"][k] = _num(r[-1])
    return out if (out["proj"] or out["erp"]) else None


def wasde_text(path):
    """Plain text of a WASDE PDF: pypdf if installed, else pdftotext."""
    try:
        import pypdf
        return "\n".join(pg.extract_text() or "" for pg in pypdf.PdfReader(path).pages)
    except ImportError:
        import subprocess
        return subprocess.run(["pdftotext", path, "-"], capture_output=True, text=True, check=True).stdout


def parse_wasde(text):
    """-> {date, number, prices: {crop: projected season-average farm price, this month}}.

    Each U.S. table's 'Avg. Farm Price ($/bu)' row ends with the current
    month's projection: the wheat table, then CORN under the feed grain
    table, then SOYBEANS under the soybean table."""
    out = {"prices": {}}
    m = re.search(r"WASDE\s*-\s*(\d+)\s+Approved by the World Agricultural Outlook Board\s+([A-Z][a-z]+ \d{1,2}, 20\d\d)", text)
    if m:
        out["number"] = int(m.group(1))
        out["date"] = dt.datetime.strptime(m.group(2), "%B %d, %Y").date().isoformat()
    for k, start, sub in (("wheat", r"U\.S\. Wheat Supply and Use", None),
                          ("corn", r"U\.S\. Feed Grain and Corn Supply and Use", r"\nCORN\s*\n"),
                          ("soybeans", r"U\.S\. Soybeans and Products Supply and Use", r"\nSOYBEANS\s*\n")):
        a = re.search(start, text)
        if not a:
            continue
        pos = a.end()
        if sub:
            b = re.compile(sub).search(text, pos)
            if not b:
                continue
            pos = b.end()
        line = re.compile(r"Avg\. Farm Price \(\$/bu\)[^\n]*").search(text, pos)
        if line:
            nums = re.findall(r"\d+\.\d+", line.group(0))
            if nums:
                out["prices"][k] = float(nums[-1])
    return out


def wasde_key(p):
    """Sort key for wasdeMMYY.pdf: (year, month), so wasde0127 comes after wasde1026."""
    m = re.search(r"wasde(\d\d)(\d\d)", os.path.basename(p))
    return (int(m.group(2)), int(m.group(1))) if m else (0, 0)


def load_wasde(root="."):
    """The newest WASDE PDF in data/wasde-pdf/ (wasdeMMYY.pdf)."""
    ps = glob.glob(os.path.join(root, "data/wasde-pdf/wasde*.pdf"))
    if not ps:
        return None
    p = max(ps, key=wasde_key)
    w = parse_wasde(wasde_text(p))
    w["file"] = os.path.relpath(p, root)
    return w if w.get("prices") and w.get("date") else None


def load_rma_hist(root="."):
    """RMA projected (February) and harvest (October) prices for corn and
    soybeans by crop year, from the one table the site already publishes
    (harvest-price-tracker.html, mirrored from RMA price discovery)."""
    p = os.path.join(root, RMA_HIST_PAGE)
    if not os.path.exists(p):
        return {}
    s = open(p, encoding="utf-8").read()
    m = re.search(r'<table class="hp-hist">(.*?)</table>', s, re.S)
    if not m:
        return {}
    out = {"corn": {}, "soybeans": {}}
    for r in re.findall(r"<tr>(.*?)</tr>", m.group(1), re.S):
        cells = [re.sub(r"[^\d.]", "", html.unescape(re.sub(r"<[^>]+>", "", c)).split("$")[-1]) for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", r, re.S)]
        if len(cells) != 5 or not re.fullmatch(r"\d{4}", cells[0]):
            continue
        y = int(cells[0])
        for k, (a, b) in (("corn", (cells[1], cells[2])), ("soybeans", (cells[3], cells[4]))):
            if a and b:
                out[k][y] = (float(a), float(b))
    return out


def snapshot_prices(root="."):
    """Freeze the futures the scenarios are centered on, so a build is reproducible."""
    pj = _read(root, "data/prices.json")
    hp = _read(root, "data/harvest-prices.json")
    snap = {"taken": pj.get("fetched"), "harvest": {}, "futures": {}}
    lab = {"Corn": "corn", "Soybeans": "soybeans"}
    for c in hp.get("commodities", []):
        k = lab.get(c.get("label"))
        h = c.get("harvest") or {}
        if k and h.get("running_avg"):
            snap["harvest"][k] = {"contract": c.get("contract"), "price": h["running_avg"], "days": h.get("days_counted"),
                                  "of": h.get("days_total"), "status": h.get("status"), "updated": hp.get("updated")}
    for k, key in (("corn", "corn-dec27"), ("soybeans", "beans-nov27")):
        q = (pj.get("quotes") or {}).get(key)
        if q and q.get("close"):
            snap["futures"][k] = {"key": key, "price": round(q["close"] / 100, 4), "date": q.get("close_date")}
    bad = snapshot_problem(snap, BUILD_DATE or dt.datetime.now(dt.timezone.utc).date(), SNAP_FRESH_DAYS)
    if bad:
        raise SystemExit(f"[arc-plc] --snapshot-prices refused, {SNAPSHOT} left as it was: data/prices.json {bad}")
    with open(os.path.join(root, SNAPSHOT), "w") as f:
        json.dump(snap, f, indent=1)
        f.write("\n")
    return snap


def snapshot_problem(snap, today, max_days=SNAP_MAX_DAYS):
    """Why a futures snapshot must not be published on date today, or None when it is fine."""
    if not snap:
        return f"has no futures snapshot ({SNAPSHOT}); run --snapshot-prices"
    for k, key in (("corn", "corn-dec27"), ("soybeans", "beans-nov27")):
        fu = (snap.get("futures") or {}).get(k) or {}
        if not fu.get("price") or fu["price"] <= 0 or not fu.get("date"):
            return f"has no {key} close"
        age = (today - dt.date.fromisoformat(fu["date"][:10])).days
        if age > max_days:
            return f"{key} close is from {fu['date'][:10]}, {age} days before {today.isoformat()} (limit {max_days})"
    return None


def signup_copy(today):
    """Signup wording for date today. Each program year's window is upcoming, open or closed."""
    a, b = YEAR_INFO[2026], YEAR_INFO[2027]
    st = {py: ("upcoming" if today < dt.date.fromisoformat(YEAR_INFO[py]["open"]) else
               "open" if today <= dt.date.fromisoformat(YEAR_INFO[py]["close"]) else "closed") for py in YEARS}
    c6, o7, c7 = nice_date(a["close"]), nice_date(b["open"]), nice_date(b["close"])
    sign = ("Sign", "Every owner and operator on the farm signs.")
    bring = ("Bring", "Bring the farm&rsquo;s FSA-156EZ. It lists the base acres and PLC yields.")
    if st[2026] != "closed":
        items = [(nice_date(a["close"], False), f"2026 signup ends {c6}."),
                 (f"{nice_date(b['open'], False)} to {nice_date(b['close'], False)}",
                  f"2027 signup runs {o7} to {c7}." if st[2027] == "upcoming" else f"2027 signup is open through {c7}."),
                 ("Skip it", "Do nothing for 2026 and that farm gets no 2026 payment."), sign, bring]
    elif st[2027] != "closed":
        items = [(nice_date(b["close"], False), f"2027 signup is open through {c7}."), ("Closed", f"2026 signup closed {c6}."), sign, bring]
    else:
        items = [("Closed", f"2026 signup closed {c6}; 2027 signup closed {c7}."), bring]
    if st[2026] != "closed":
        return {"state": st, "items": items, "box6": "2026 signup closes", "box7": "2027 signup", "desc": "2026 signup ends Dec 11.",
                "quick": f"The 2026 signup closes {c6}; the 2027 signup runs {o7} to {c7}.",
                "llms": f"Signup: 2026 closes {c6}; 2027 runs {o7} to {c7}.", "faq": "The 2026 deadline comes first."}
    if st[2027] != "closed":
        return {"state": st, "items": items, "box6": "2026 signup closed", "box7": "2027 signup", "desc": "2027 signup ends Mar 15.",
                "quick": f"The 2026 signup closed {c6}. The 2027 signup is open through {c7}.",
                "llms": f"Signup: 2026 closed {c6}; 2027 runs {o7} to {c7}.", "faq": f"The 2026 signup has closed; the 2027 signup is open through {c7}."}
    return {"state": st, "items": items, "box6": "2026 signup closed", "box7": "2027 signup closed", "desc": "2027 signup closed Mar 15.",
            "quick": f"The 2026 signup closed {c6}, and the 2027 signup closed {c7}.",
            "llms": f"Signup: 2026 closed {c6}; 2027 closed {c7}.", "faq": "Both signups have closed."}


def build_date():
    return BUILD_DATE or dt.datetime.now(dt.timezone.utc).date()


# ---------------------------------------------------------------- loading
def _read(root, rel):
    with open(os.path.join(root, rel), encoding="utf-8") as f:
        return json.load(f)


def load_inputs(root="."):
    inp = {"mya": {}, "mya_meta": {}}
    for crop, c in CROPS.items():
        d = _read(root, f"data/nass/{c['mya']}.json")
        if FINAL_SCOPE not in d.get("scope", ""):
            raise SystemExit(f"{c['mya']}.json scope does not say '{FINAL_SCOPE}'; refusing a forecast price")
        inp["mya"][crop] = {int(k): float(v) for k, v in d["values"].items()}
        inp["mya_meta"][crop] = {"file": f"data/nass/{c['mya']}.json", "updated": d.get("updated")}
    inp["counties"] = _read(root, "farmland-atlas/data/counties-index.json")
    inp["fsa"] = load_fsa(root)
    inp["arc_tabs"] = load_arc_tables(root)
    inp["plc_county"] = load_plc_county(root)
    inp["fsa_tables"] = load_fsa_tables(root)
    inp["rma_hist"] = load_rma_hist(root)
    inp["wasde"] = load_wasde(root)
    inp["xtab"] = XC.load_tables(load_xlsx, root)
    inp["wasde_x"] = XC.load_wasde_x(wasde_text, root, inp["wasde"])
    inp["snap"] = _read(root, SNAPSHOT) if os.path.exists(os.path.join(root, SNAPSHOT)) else None
    inp["geo"] = _read(root, GEO_SRC) if os.path.exists(os.path.join(root, GEO_SRC)) else None
    return inp


# ---------------------------------------------------------------- compute
def year_params(inp, py):
    out = {}
    win = list(range(py - 6, py - 1))
    for crop, c in CROPS.items():
        mya = inp["mya"][crop]
        miss = [y for y in win if y not in mya]
        if miss:
            raise SystemExit(f"{crop}: no final MYA for {miss}; the {py} ERP cannot be computed")
        ys = [mya[y] for y in win]
        e = erp_calc(ys, c["statutory"], PARAMS["erp_pct"], PARAMS["erp_cap_pct"])
        bp, used, blo, bhi = benchmark_price(ys, e["erp"])
        fk, fl = YEAR_INFO[py]["fut"][crop]
        out[crop] = {"label": c["label"], "statutory": c["statutory"], "loan": c["loan"], "prior_statutory": c["prior_statutory"],
                     "mya": {str(y): v for y, v in zip(win, ys)}, "erp": e,
                     "bp": {"value": bp, "used": {str(y): v for y, v in zip(win, used)}, "dropped_low": blo, "dropped_high": bhi},
                     "futures": {"key": fk, "label": fl}, "my": YEAR_INFO[py]["my"][crop],
                     "grid": {"lo": c["grid"][0], "hi": c["grid"][1], "step": c["grid"][2]}}
    return {"window": win, **{k: YEAR_INFO[py][k] for k in ("open", "close", "pay_after")}, "crops": out}


def backtest(inp):
    """What each program paid nationally per base acre, year by year, from
    FSA's own files: PLC rate from the ERP and final MYA; ARC-CO from the
    county payment rates in FSA's arcco files, weighted by enrolled base."""
    plc = inp["plc_county"]
    if not plc:
        return None
    nat_py = {}
    for k in CROPS:
        num_ = sum(plc["yields"][f][k] * plc["base"][f][k] for f in plc["base"] if k in plc["base"][f] and k in plc["yields"].get(f, {}))
        den = sum(plc["base"][f][k] for f in plc["base"] if k in plc["base"][f] and k in plc["yields"].get(f, {}))
        nat_py[k] = num_ / den if den else None
    by_crop = {v["fsa"]: k for k, v in CROPS.items()}
    arc = {}
    for py_, f in inp["fsa"].items():
        per = {}
        for (fips, sub, crop, d), r in f["rows"].items():
            k = by_crop.get(crop)
            if k and r["pay"] is not None:
                per.setdefault((k, fips), []).append(r["pay"])
        for k in CROPS:
            ws = [(sum(v) / len(v), plc["base"].get(fips, {}).get(k)) for (kk, fips), v in per.items() if kk == k]
            ws = [(a, b) for a, b in ws if b]
            if ws:
                arc[(py_, k)] = (per_base(sum(a * b for a, b in ws) / sum(b for _a, b in ws)), len(ws))
    rows = []
    for y in BACKTEST_YEARS:
        for k in CROPS:
            mya = inp["mya"][k]
            win = [mya.get(t) for t in range(y - 6, y - 1)]
            if None in win or y not in mya:
                continue
            if y <= 2024:
                stat, loan = LAW2018[k]
                erp = erp_calc(win, stat, 0.85, 1.15)["erp"]
            else:
                stat, loan = CROPS[k]["statutory"], LAW2018[k][1]
                erp = erp_calc(win, stat, 0.88, 1.15)["erp"]
            rate = rnd(plc_rate(erp, mya[y], loan))
            a = arc.get((y, k))
            rows.append({"y": y, "k": k, "erp": erp, "mya": mya[y], "rate": rate,
                         "plc": per_base(rate * nat_py[k]) if nat_py[k] else None,
                         "arc": a[0] if a else None, "arc_n": a[1] if a else 0})
    return {"rows": rows, "nat_py": nat_py, "plc_py": plc["py"]}


def sd1(vals):
    """Sample standard deviation (n - 1)."""
    vals = list(vals)
    m = sum(vals) / len(vals)
    return math.sqrt(sum((x - m) ** 2 for x in vals) / (len(vals) - 1))


def oct_spread(inp, k):
    """Final MYA / October futures average (RMA harvest price) by year, for
    corn or soybeans: (raw ratios, their mean, ratios / mean). None without the history."""
    mya = inp["mya"][k]
    hist = (inp.get("rma_hist") or {}).get(k) or {}
    f = {t: mya[t] / hist[t][1] for t in SCEN_YEARS if t in mya and t in hist}
    if len(f) < MIN_SCEN:
        return None
    m = sum(f.values()) / len(f)
    return f, m, {t: round(v / m, 4) for t, v in f.items()}


def rescale_to(own, target_sd):
    """A crop's own price changes, scaled to mean 1, then stretched or shrunk
    so their spread (sample SD) equals target_sd. The mean stays 1."""
    m = sum(own.values()) / len(own)
    n_ = {t: v / m for t, v in own.items()}
    s_ = sd1(n_.values())
    return {t: round(1 + (v - 1) * target_sd / s_, 4) for t, v in n_.items()}


def price_scen(inp, k, py):
    """Price scenarios for crop k, program year py: {center, ratios {year: multiplier}, kind, ...}.

    2026, every crop: the center is USDA&rsquo;s projected 2026/27 season-average
      price (the October WASDE). The crop is mostly harvested, so the county
      yield is a normal crop (fix: every scenario year at the benchmark) and only
      the price moves.
    2026 spread, corn and soybeans: each past year&rsquo;s final MYA / that year&rsquo;s
      October futures average (RMA harvest price), scaled to mean 1: how far the
      season-average price ended from October futures. alt: the futures-implied
      center, this October&rsquo;s Dec 2026 / Nov 2026 average so far x the mean of
      final MYA / October futures; where it and USDA&rsquo;s center disagree on the
      leader or on Pick, Pick becomes Leans.
    2026 spread, wheat (and any crop with no October futures history of its own):
      its own MYA changes scaled to mean 1, rescaled to corn&rsquo;s October spread
      (sample SD); borrowed, so the call is capped at Leans.
    2027 corn and soybeans: Dec 2027 / Nov 2027 futures today, times the
      2011-2025 average of final MYA / February futures (RMA projected price),
      spread by each past year&rsquo;s MYA change, scaled to mean 1; county yields
      move the way they did that year.
    2027 wheat: no 2027 price we can calibrate (no wheat futures history on
      file), so no call (kind noprice); the 2026 price is not reused as a 2027 one."""
    mya = inp["mya"][k]
    yoy = {t: mya[t] / mya[t - 1] for t in SCEN_YEARS if t in mya and (t - 1) in mya}
    hist = (inp.get("rma_hist") or {}).get(k) or {}
    snap = inp.get("snap") or {}
    tabs = inp.get("fsa_tables") or {}
    usda = tabs.get("proj", {}).get(k)
    base = {"usda": usda, "usda_date": nice_date(tabs["date"]) if tabs.get("date") else None}
    wz = inp.get("wasde") or {}
    wp = wz.get("prices", {}).get(k)
    if wp:
        base.update({"usda": wp, "usda_date": nice_date(wz["date"]), "usda_src": f"USDA WASDE, {nice_date(wz['date'])}"})
    lab26 = f"USDA&rsquo;s projected 2026/27 season-average price (WASDE, {nice_date(wz['date'])})" if wp else ""
    if py == 2026 and wp:
        os_ = oct_spread(inp, k)
        if os_:
            f, m, rat = os_
            out = {**base, "kind": "oct", "center": wp, "ratios": rat, "fix": True, "center_label": lab26, "oct_mean": round(m, 4)}
            hv = (snap.get("harvest") or {}).get(k)
            if hv and hv.get("price"):
                out.update({"alt": round(hv["price"] * m, 4), "alt_fut": hv["price"], "alt_days": hv.get("days"), "alt_of": hv.get("of"),
                            "alt_contract": hv.get("contract"), "alt_date": hv.get("updated")})
            return out
        corn = oct_spread(inp, "corn")
        if corn and len(yoy) >= MIN_SCEN:
            return {**base, "kind": "yoy_scaled", "center": wp, "ratios": rescale_to(yoy, sd1(corn[2].values())), "fix": True,
                    "borrowed": True, "center_label": lab26}
        return None
    if py == 2027 and k in snap.get("futures", {}) and len(hist) >= 10:
        fu = snap["futures"][k]
        gy = [t for t in sorted(hist) if t in mya]
        g = sum(mya[t] / hist[t][0] for t in gy) / len(gy)
        m = sum(yoy.values()) / len(yoy)
        return {**base, "kind": "fut", "center": round(fu["price"] * g, 4), "fut": fu["price"], "fut_date": nice_date(fu["date"]),
                "gap": round(g, 4), "gap_years": [gy[0], gy[-1], len(gy)],
                "ratios": {t: round(v / m, 4) for t, v in yoy.items()},
                "center_label": (f"{'Dec 2027 corn' if k == 'corn' else 'Nov 2027 soybean'} futures {usd(fu['price'])} ({nice_date(fu['date'])}) "
                                 f"&times; {g:.3f}, the {gy[0]} to {gy[-1]} average of final MYA / February futures")}
    if py == 2027:
        return {**base, "kind": "noprice", "center": None, "ratios": {},
                "why": "There is no 2027 price outlook for this crop we can check against past years, so no 2027 call. "
                       "Type a 2027 price in the calculator to see the dollars."}
    return None


def compute(inp):
    years = {str(py): year_params(inp, py) for py in YEARS}
    fsa = inp["fsa"]
    primary = fsa.get(2026)
    history_pys = sorted(fsa)
    add_other_crops(inp, years, primary)
    by_crop = {v["fsa"]: k for k, v in CROPS.items()}
    # FSA's benchmark price must equal ours for the same year
    if primary:
        for (fips, sub, crop, d), rec in primary["rows"].items():
            k = by_crop.get(crop)
            if k and rec["bp"] is not None and abs(rec["bp"] - years["2026"]["crops"][k]["bp"]["value"]) > 0.005:
                raise SystemExit(f"FSA 2026 benchmark price for {crop} is {rec['bp']}, ours "
                                 f"{years['2026']['crops'][k]['bp']['value']}: refusing to publish two answers")
    by_crop.update({v["fsa"]: k for k, v in XC.XCROPS.items() if k in years["2026"]["crops"]})
    idx = {r[0]: r for r in inp["counties"]}
    counties = {}
    plc = inp["plc_county"]
    # OFFICIAL mode: rows of FSA's 2026 file. PENDING mode (no 2026 file yet):
    # the county x crop x practice rows FSA's older files list, with no
    # benchmark (by = None); pages then print only FSA's history and the
    # county average PLC yield, never a current benchmark, trigger or crossover in bushels.
    keys = sorted(primary["rows"]) if primary else sorted({k for hp in history_pys for k in fsa[hp]["rows"]})
    for key in keys:
        fips, sub, crop, d = key
        k = by_crop.get(crop)
        rec = primary["rows"][key] if primary else None
        if not k or (primary and rec["by"] is None):
            continue
        st = FIPS_ST.get(fips[:2])
        if not st:
            continue
        if True:
            cname = rec["county"] if rec else next(fsa[hp]["rows"][key]["county"] for hp in history_pys if key in fsa[hp]["rows"])
            c = counties.setdefault(fips, {"f": fips, "st": st, "n": cname, "k": {}})
            ent = {"d": d, "by": rec["by"] if rec else None, "yrs": {str(y): v for y, v in rec["yrs"].items()} if rec else {}}
            if sub:
                ent["sub"] = sub
            hist = []
            for hp in history_pys:
                h = fsa[hp]["rows"].get((fips, sub, crop, d))
                if h and h["by"] is not None:
                    hist.append({"py": hp, "by": h["by"], "bp": h["bp"], "ay": h["ay"], "np": h["np"], "pay": h["pay"]})
            ent["hist"] = hist
            # county yield / benchmark for each past year, from FSA's own files:
            # the five window years of each program year, against that year's benchmark
            dy = {}
            for hp in history_pys:
                h = fsa[hp]["rows"].get((fips, sub, crop, d))
                if not h or not h["by"]:
                    continue
                for t, v in h["yrs"].items():
                    if v is not None and t in SCEN_YEARS and (t not in dy or hp > dy[t][1]):
                        dy[t] = (v / h["by"], hp)
            # 2021 on: FSA's actual county yield for that program year (no 80% plug, no trend
            # adjustment) against that year's own benchmark, whose window leaves that year out
            act = []
            for hp in history_pys:
                h = fsa[hp]["rows"].get((fips, sub, crop, d))
                if h and h["by"] and h["ay"] is not None and hp in SCEN_YEARS:
                    dy[hp] = (h["ay"] / h["by"], 9999)
                    act.append(hp)
            ent["dy"] = {t: round(v, 4) for t, (v, _hp) in sorted(dy.items())}
            ent["dact"] = act
            ent["dsrc"] = "county"
            c["k"].setdefault(k, []).append(ent)
    states = {}
    for fips, c in counties.items():
        r = idx.get(fips)
        if r:
            name, _st = r[1].rsplit(", ", 1)
            parts = r[2].strip("/").split("/")
            c["n"], c["s"], ss = name, parts[2], parts[1]
        else:
            nm = c["n"] + (" County" if not re.search(r"(County|Parish|Borough|city|City)$", c["n"]) else "")
            c["n"], c["s"], ss = nm, slugify(nm), slugify(STATE_NAMES[c["st"]])
        S = states.setdefault(c["st"], {"n": STATE_NAMES[c["st"]], "slug": ss, "c": []})
        S["c"].append(c)
    for S in states.values():
        S["c"].sort(key=lambda x: x["n"])
        dedupe_slugs(S, idx)
    # a county with fewer than MIN_SCEN years of its own takes its state's
    # average for that crop and practice, year by year (MIN_STATE_N counties)
    pool = {}
    for c in counties.values():
        for k, es in c["k"].items():
            for e in es:
                if len(e["dy"]) >= MIN_SCEN:
                    for t, v in e["dy"].items():
                        pool.setdefault((c["st"], k, e["d"], t), []).append(v)
    for c in counties.values():
        for k, es in c["k"].items():
            for e in es:
                if len(e["dy"]) < MIN_SCEN:
                    sdy = {t: round(sum(v) / len(v), 4) for t in SCEN_YEARS
                           if len(v := pool.get((c["st"], k, e["d"], t), [])) >= MIN_STATE_N}
                    if len(sdy) >= MIN_SCEN:
                        e["dy"], e["dsrc"] = sdy, "state"
    # scenario prices: USDA's projected 2026/27 MYA (FSA's table) x each past year's MYA change
    tabs = inp.get("fsa_tables")
    if tabs:
        for k, v in tabs["erp"].items():
            if abs(v - years["2026"]["crops"][k]["erp"]["erp"]) > 0.005:
                raise SystemExit(f"FSA's 2026 ERP for {k} is {v}, ours {years['2026']['crops'][k]['erp']['erp']}: refusing two answers")
    corn_oct = oct_spread(inp, "corn")
    for y in years:
        for k, cd in years[y]["crops"].items():
            if k in CROPS:
                cd["scen"] = price_scen(inp, k, int(y))
                continue
            sx = XC.price_scen_x(k, cd, inp["xtab"], inp.get("wasde_x"), SCEN_YEARS, MIN_SCEN)
            if not sx:
                cd["scen"] = None
            elif y == "2027":
                cd["scen"] = {**{kk: sx[kk] for kk in ("usda", "usda_src") if kk in sx}, "kind": "noprice", "center": None, "ratios": {},
                              "why": "There is no 2027 price outlook for this crop we can check against past years, so no 2027 call. "
                                     "Type a 2027 price in the calculator to see the dollars."}
            elif corn_oct:
                own = sx["ratios"]
                # too few years of its own price changes: borrow corn's October spread; enough: its own, rescaled to corn's
                rat = rescale_to(own, sd1(corn_oct[2].values())) if len(own) >= MIN_SCEN else dict(corn_oct[2])
                cd["scen"] = {**sx, "kind": "yoy_scaled" if len(own) >= MIN_SCEN else "corn_oct", "ratios": rat, "fix": True, "borrowed": True,
                              "own_years": len(own), "why": ""}
            else:
                cd["scen"] = sx
    arc_tabs = inp.get("arc_tabs") or {}
    paid_checked = check_fsa_pay(fsa, arc_tabs)
    plc_all = inp["plc_county"]["yields"] if inp["plc_county"] else {}
    for c in counties.values():
        c["verdicts"], c["levels"] = {}, {}
        for k, es in c["k"].items():
            pyv = plc_all.get(c["f"], {}).get(k)
            for i, e in enumerate(es):
                if not e["by"] or not pyv:
                    continue
                for y in years:
                    cd = years[y]["crops"].get(k)
                    lv = {}
                    for nm, f_ in LEVELS:
                        py_l = rnd(pyv * f_, 1)
                        if not cd:
                            v = {"verdict": "pending", "n": 0}
                        elif py_l > e["by"]:
                            # FSA's county average PLC yield covers all practices; above this
                            # practice's benchmark it does not describe a typical farm here
                            v = {"verdict": "mismatch", "py": py_l, "by": e["by"], "n": 0}
                        else:
                            v = typical_eval(cd, py_l, e)
                        if v is None:
                            continue
                        v["py"] = py_l
                        lv[nm] = v
                    if "same" in lv:
                        c["levels"][(k, i, y)] = lv
                        c["verdicts"][(k, i, y)] = lv["same"]
    plc = inp["plc_county"]
    for c in counties.values():
        c["lead"] = lead_crop(c, plc)
    updated = data_date(inp, primary)
    return {
        "updated": updated,
        "years": years,
        "default_year": {"until": YEAR_INFO[2026]["close"], "then": 2027, "before": 2026},
        "params": PARAMS,
        "fsa": ({"py": primary["py"], "as_of": primary["as_of"], "file": primary["file"], "window": primary["years"],
                 "history": [{"py": hp, "as_of": fsa[hp]["as_of"], "file": fsa[hp]["file"]} for hp in history_pys]}
                if primary else None),
        "plc_county": ({"py": plc["py"], "file": plc["file"]} if plc else None),
        "_backtest": backtest(inp),
        "_paid_checked": paid_checked,
        "_dysrc": {t: ("act" if any(t == hp and any(r["ay"] is not None for r in fsa[hp]["rows"].values()) for hp in history_pys)
                       else max((hp for hp in history_pys if t in fsa[hp]["years"]), default=None)) for t in SCEN_YEARS},
        "_arc_tabs": {py: t["file"] for py, t in arc_tabs.items()},
        "states": states,
        "_plc_yields": plc["yields"] if plc else {},
        "_plc_base": plc["base"] if plc else {},
        "sources": {k: {"name": v[0], "url": v[1]} for k, v in SRC.items()},
        "links": {"fsa": FSA_PAGE, "fsa_data": FSA_DATA, "office": OFFICE},
        "crops": [{"k": k, "label": (CROPS.get(k) or XC.XCROPS[k])["label"]} for k in ALL if k in years["2026"]["crops"]],
        "hubs": [{"k": k, "slug": XC.XCROPS[k]["slug"], "label": XC.XCROPS[k]["label"]} for k in XC.XCROPS if k in years["2026"]["crops"]],
        "_hubs": [k for k in XC.XCROPS if k in years["2026"]["crops"]],
        "_xtab_date": (inp.get("xtab") or {}).get("date"),
        "_geo": inp.get("geo"),
    }


def typical_eval(cd, py, e):
    """One typical-farm scenario run for crop data cd at PLC yield py on benchmark entry e."""
    sc = cd.get("scen")
    if not sc:
        return None
    if sc["kind"] == "noprice":
        return {"verdict": "noprice", "n": 0}
    dy = 1.0 if sc.get("fix") else e["dy"]
    return scen_eval(cd["erp"]["erp"], cd["bp"]["value"], cd["loan"], py, [(1.0, e["by"], dy)], sc["center"], sc["ratios"],
                     borrowed=bool(sc.get("borrowed")), alt=sc.get("alt"))


def add_other_crops(inp, years, primary):
    """Every other covered crop from FSA's national tables (scripts/arc_plc_crops.py),
    checked against FSA's own 2026 figures and county file before anything is used."""
    for py in YEARS:
        crops, pending = XC.year_params_x(inp.get("xtab"), py)
        if py == 2026:
            XC.check_against_fsa(crops, primary)
        for n, p1, p0, _l in OTHER_STATUTORY:   # the 2018 law's statutory price, for the 'before 2025' line [ERS]
            k = n.lower()
            if k in crops and abs(crops[k]["statutory"] - p1) < 1e-9:
                crops[k]["prior_statutory"] = p0
        years[str(py)]["crops"].update(crops)
        years[str(py)]["pending"] = pending
    for y in years.values():
        for k in CROPS:
            y["crops"][k].update({"unit": "bu", "dp": 2, "tick": 0.01})


def dedupe_slugs(S, idx):
    """Two FSA county codes can name the same place (Virginia's former Bedford city,
    51515, and Bedford County, 51019). The one not in the county index gets its own
    name and slug, so no page overwrites another."""
    seen = {}
    for c in S["c"]:
        seen.setdefault(c["s"], []).append(c)
    for s_, cs in seen.items():
        if len(cs) < 2:
            continue
        for c in cs:
            if c["f"] in idx:
                continue
            base = re.sub(r" County$", "", c["n"])
            c["n"] = f"{base} city (former)" if c["st"] == "VA" and c["f"][2:] >= "500" else f"{base} ({c['f']})"
            c["s"] = slugify(c["n"])


def lead_crop(c, plc):
    """The crop a county page leads with: the most FSA enrolled base acres in the county
    (FSA's PLC county yield file), among crops with a 2026 benchmark here."""
    ks = [k for k in ALL if any(e["by"] for e in c["k"].get(k, []))] or [k for k in ALL if k in c["k"]]
    base = ((plc or {}).get("base") or {}).get(c["f"], {})
    if not ks:
        return None
    return max(ks, key=lambda k: (base.get(k) or 0, -ALL.index(k)))


def lead_year():
    """The program year county pages lead with: 2026 until its signup closes, then 2027."""
    return "2027" if signup_copy(build_date())["state"][2026] == "closed" else "2026"


def data_date(inp, primary):
    """One date for every 'updated' line, dateModified and sitemap lastmod: the newest
    input (NASS prices, FSA files and tables, WASDE, futures snapshot) or the day the
    signup wording last changed, whichever is later."""
    ds = [m["updated"] or "" for m in inp["mya_meta"].values()] + [(primary or {}).get("as_of") or ""]
    ds.append((inp.get("xtab") or {}).get("date") or "")
    ds.append((inp.get("wasde") or {}).get("date") or "")
    for fu in ((inp.get("snap") or {}).get("futures") or {}).values():
        ds.append((fu.get("date") or "")[:10])
    today = build_date()
    for py in YEARS:
        for d_ in (dt.date.fromisoformat(YEAR_INFO[py]["open"]), dt.date.fromisoformat(YEAR_INFO[py]["close"]) + dt.timedelta(days=1)):
            if d_ <= today:
                ds.append(d_.isoformat())
    return max(x[:10] for x in ds if x)


def geo_json(D):
    """Use my location, all on the phone: nav.json (every state's bounding box and its
    counties with a page) and geo/<ST>.json (each county's shape from the Farmland Atlas
    file, with its page slug or null). -> {path: object}."""
    g = D.get("_geo")
    if not g:
        return {}
    pages = {c["f"]: c["s"] for S in D["states"].values() for c in S["c"]}
    per, box = {}, {}
    for f in g["features"]:
        st = f["properties"]["st"]
        geom = f["geometry"]
        polys = geom["coordinates"] if geom["type"] == "MultiPolygon" else [geom["coordinates"]]
        xs = [pt[0] for poly in polys for ring in poly for pt in ring]
        ys = [pt[1] for poly in polys for ring in poly for pt in ring]
        bb = [min(xs), min(ys), max(xs), max(ys)]
        per.setdefault(st, []).append([f["id"], f["properties"]["name"], pages.get(f["id"]), bb, polys])
        b0 = box.get(st)
        box[st] = bb if not b0 else [min(b0[0], bb[0]), min(b0[1], bb[1]), max(b0[2], bb[2]), max(b0[3], bb[3])]
    nav = {"s": {}}
    for code, S in sorted(D["states"].items()):
        nav["s"][code] = {"n": S["n"], "slug": S["slug"], "c": [[c["f"], c["n"], c["s"]] for c in S["c"]]}
    for st, bb in box.items():
        nav["s"].setdefault(st, {"n": STATE_NAMES.get(st, st), "slug": None, "c": []})["b"] = bb
    out = {f"{OUT_STATE_JSON}/nav.json": nav}
    for st, fs in per.items():
        out[f"{OUT_STATE_JSON}/geo/{st}.json"] = {"st": st, "c": fs}
    return out


def in_ring(x, y, ring):
    """Even-odd ray cast. Mirrors inRing in components/arc-plc-simple.js."""
    inside = False
    j = len(ring) - 1
    for i in range(len(ring)):
        xi, yi = ring[i][0], ring[i][1]
        xj, yj = ring[j][0], ring[j][1]
        if (yi > y) != (yj > y) and x < (xj - xi) * (y - yi) / (yj - yi) + xi:
            inside = not inside
        j = i
    return inside


def locate(nav, geo_by_st, lon, lat):
    """-> (state, county record) for a point, or (None, None). Holes count by even-odd within each polygon."""
    for st, s_ in nav["s"].items():
        b = s_.get("b")
        if not b or not (b[0] <= lon <= b[2] and b[1] <= lat <= b[3]):
            continue
        for rec in (geo_by_st.get(st) or {}).get("c", []):
            bb = rec[3]
            if not (bb[0] <= lon <= bb[2] and bb[1] <= lat <= bb[3]):
                continue
            for poly in rec[4]:
                if sum(1 for ring in poly if in_ring(lon, lat, ring)) % 2 == 1:
                    return st, rec
    return None, None


def calc_json(D):
    """The calculator's files: a small main file (years, params, the state
    list) and one file per state with its counties, loaded when a state is
    picked, so a phone never downloads the whole country."""
    per = {}
    for code, S in D["states"].items():
        per[code] = {"st": code, "n": S["n"], "c": [
            {"f": c["f"], "n": c["n"], "s": c["s"],
             "k": {k: [{kk: v for kk, v in e.items() if kk in ("d", "by", "sub", "dy", "dsrc")} for e in es if e["by"]] for k, es in c["k"].items()
                   if any(e["by"] for e in es)},
             **({"plc": D["_plc_yields"][c["f"]]} if c["f"] in D["_plc_yields"] else {})}
            for c in S["c"]]}
    out = {k: v for k, v in D.items() if k not in ("states", "_plc_yields", "_backtest") and not k.startswith("_")}
    out["states"] = {code: {"n": S["n"], "slug": S["slug"], "nc": len(S["c"])} for code, S in D["states"].items()}
    return out, per


# ---------------------------------------------------------------- page parts
def chrome():
    try:
        import inject_static_nav as NAV
        return (f'<div id="site-header">{NAV.block("header")}</div>',
                f'<div id="site-footer">{NAV.block("footer")}</div>', NAV.FONT_BLOCK)
    except Exception:  # pragma: no cover
        return ('<div id="site-header"></div>', '<div id="site-footer"></div>', "")


def head(title, desc, path, jsonld, indexable=True, fonts="", og_title=None, og_desc=None):
    robots = "index,follow" if indexable else "noindex,follow"
    ogt, ogd = og_title or title, og_desc or desc
    ld = "\n".join('  <script type="application/ld+json">' + json.dumps(j, ensure_ascii=False).replace("</", "<\\/") + "</script>"
                   for j in jsonld)
    return f"""<!DOCTYPE html>
<html lang="en" data-theme="dark">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0, viewport-fit=cover">
  <meta name="theme-color" content="#0a0c0d">
  <script>/* agsist-theme-early: saved theme, else the device's, set before first paint so pages do not flash the wrong color. loader.js applies the same rule. */try{{var _t=localStorage.getItem('agsist-theme');if(_t!=='light'&&_t!=='dark')_t=window.matchMedia&&matchMedia('(prefers-color-scheme: light)').matches?'light':'dark';document.documentElement.setAttribute('data-theme',_t);}}catch(e){{}}</script><script>/* agsist-ga-guard 2026-10-09: Google Analytics loads only when the browser sends no Global Privacy Control signal and the off switch on /privacy is not set, and only once the page is shown (a page loaded ahead in the background is not a visit). dataLayer and gtag always exist, so page code that calls them never throws. */(function(w,d,n){{var off=false,v,i,s;w.dataLayer=w.dataLayer||[];if(typeof w.gtag!=='function'){{w.gtag=function(){{w.dataLayer.push(arguments);}};}}try{{off=w.localStorage.getItem('agsist-ga-off')==='1';}}catch(e){{}}if(n.globalPrivacyControl===true){{off=true;}}w.agsistGaOff=off;w.gtag('set','allow_google_signals',false);w.gtag('set','allow_ad_personalization_signals',false);if(off){{return;}}i=function(){{s=d.createElement('script');s.async=true;s.src='https://www.googletagmanager.com/gtag/js?id=G-6KXCTD5Z9H';(d.head||d.documentElement).appendChild(s);}};if(d.prerendering){{d.addEventListener('prerenderingchange',i,{{once:true}});}}else{{i();}}}})(window,document,navigator);</script>
  <script>window.dataLayer=window.dataLayer||[];function gtag(){{dataLayer.push(arguments);}}gtag('js',new Date());gtag('config','G-6KXCTD5Z9H');</script>
  {fonts}
  <link rel="preload" href="/components/styles.css?v={STYLES_V}" as="style">
  <link rel="stylesheet" href="/components/styles.css?v={STYLES_V}">
  <link rel="stylesheet" href="/components/asof.css?v={ASOF_V}">
  <link rel="stylesheet" href="/components/arc-plc.css?v={CALC_V}">
  <link rel="stylesheet" href="/components/arc-plc-simple.css?v={CALC_V}">
  <link rel="icon" type="image/x-icon" href="/img/favicon.ico">
  <link rel="icon" type="image/png" sizes="32x32" href="/img/favicon-32.png">
  <link rel="icon" type="image/png" sizes="16x16" href="/img/favicon-16.png">
  <link rel="apple-touch-icon" href="/img/apple-touch-icon.png">
  <link rel="manifest" href="/manifest.json">
  <title>{esc(title)}</title>
  <meta name="description" content="{esc(desc)}">
  <meta name="robots" content="{robots}">
  <link rel="canonical" href="{SITE}{path}">
  <meta property="og:type" content="website">
  <meta property="og:site_name" content="AGSIST">
  <meta property="og:locale" content="en_US">
  <meta property="og:title" content="{esc(ogt)}">
  <meta property="og:description" content="{esc(ogd)}">
  <meta property="og:url" content="{SITE}{path}">
  <meta property="og:image" content="{SITE}/img/og/agsist.jpg">
  <meta property="og:image:alt" content="{esc(ogt)}">
  <meta property="og:image:width" content="1200">
  <meta property="og:image:height" content="630">
  <meta name="twitter:card" content="summary_large_image">
  <meta name="twitter:site" content="@agsist">
  <meta name="twitter:title" content="{esc(ogt)}">
  <meta name="twitter:description" content="{esc(ogd)}">
  <meta name="twitter:image" content="{SITE}/img/og/agsist.jpg">
{ld}
</head>"""


def tail(ftr, calc=False, math=False):
    s = f"""{ftr}
<script src="/components/asof.js?v={ASOF_V}" defer></script>
<script src="/components/loader.js?v={LOADER_V}" defer></script>
<script src="/components/arc-plc-simple.js?v={CALC_V}" defer></script>"""
    if math:
        s += f'\n<script src="/{OUT_FORM_JS}?v={CALC_V}" defer></script>'
    if calc:
        s += f'\n<script src="/components/arc-plc.js?v={CALC_V}" defer></script>'
    if math:
        s += f'\n<script src="/components/arc-plc-math.js?v={CALC_V}" defer></script>'
    return s + "\n</body>\n</html>\n"


def asof(iso, prefix="Data updated"):
    day = (iso or "")[:10]
    if not day:
        return ""
    return (f'<time class="asof" data-asof="" data-asof-prefix="{esc(prefix)}" datetime="{day}">'
            f'{esc(prefix)} {nice_date(day)}</time>')


def src_link(k):
    n, u = SRC[k]
    return f'<a href="{esc(u)}" rel="noopener">{esc(n)}</a>'


def deadlines_html():
    """The deadline list (deadlines_box); kept under the old name for arc_plc_crops."""
    return deadlines_box()


def form_html(D, crop="corn", st="", fips="", embed=False):
    chips = ('<label class="ap-crop-l" for="ap-crop">Crop</label><select id="ap-crop" class="ap-crop-s">' + "".join(
        f'<option value="{x["k"]}"{" selected" if x["k"] == crop else ""}>{esc(x["label"])}</option>' for x in D.get("crops") or
        [{"k": k, "label": v["label"]} for k, v in CROPS.items()]) + '</select>')
    ychips = "".join(f'<button type="button" class="chip" data-year="{y}" aria-pressed="false">{y}</button>' for y in YEARS)
    return f"""<section class="ap-calc" id="calculator" data-arcplc data-src="/{OUT_JSON}" data-state-src="/{OUT_STATE_JSON}/" data-crop="{esc(crop)}" data-state="{esc(st)}" data-fips="{esc(fips)}"{' data-embed="1"' if embed else ''} aria-labelledby="ap-calc-h">
  <h2 id="ap-calc-h" class="ap-h2">Run your farm&rsquo;s numbers</h2>
  <p class="ap-hint ap-runone">Run each crop on each FSA farm number separately. The election is made crop by crop, farm by farm.</p>
  <div class="ap-row"><div class="ap-crops" role="group" aria-label="Program year">{ychips}</div><div class="ap-crops ap-crop-f">{chips}</div></div>
  <div class="ap-grid">
    <div class="ap-f"><label for="ap-st">State</label><select id="ap-st"><option value="">Pick a state</option></select></div>
    <div class="ap-f"><label for="ap-co">County</label><select id="ap-co" disabled><option value="">Pick a state first</option></select></div>
    <div class="ap-f ap-f-wide" id="ap-prac-f" hidden><label for="ap-prac">Practice</label><select id="ap-prac" aria-describedby="ap-prac-hint"></select>
      <p class="ap-hint" id="ap-prac-hint">FSA has separate irrigated and non-irrigated benchmarks here. A farm with both is weighted by its historical irrigated percentage on FSA&rsquo;s records, not by what you plant this year.</p></div>
    <div class="ap-f" id="ap-share-f" hidden><label for="ap-share">Irrigated share of this farm&rsquo;s base (%)</label><input id="ap-share" type="number" inputmode="decimal" min="0" max="100" step="1" value="50"></div>
    <div class="ap-f"><label for="ap-base">Base acres for this crop</label><input id="ap-base" type="number" inputmode="decimal" min="0" step="0.1" placeholder="Blank shows per base acre" aria-describedby="ap-base-hint"><p class="ap-hint" id="ap-base-hint">From the FSA-156EZ. Base acres are not what you plant; each farm number has its own base.</p></div>
    <div class="ap-f"><label for="ap-py" id="ap-py-l">PLC payment yield (bu/acre)</label><input id="ap-py" type="number" inputmode="decimal" min="0" step="1" placeholder="From your FSA-156EZ" aria-describedby="ap-py-hint"><p class="ap-hint" id="ap-py-hint">On the FSA-156EZ for the farm. Each farm has its own.</p></div>
    <div class="ap-f"><label for="ap-by" id="ap-by-l">ARC-CO benchmark yield (bu/acre)</label><input id="ap-by" type="number" inputmode="decimal" min="0" step="0.01" placeholder="FSA official, from the county" aria-describedby="ap-by-hint"><p class="ap-hint" id="ap-by-hint">Fills in from FSA&rsquo;s official county file when we have it. You can type the number your county office gives you.</p></div>
    <div class="ap-f"><label for="ap-y" id="ap-y-l">Expected county yield (bu/acre)</label><input id="ap-y" type="number" inputmode="decimal" min="0" step="0.1" aria-describedby="ap-y-hint"><p class="ap-hint" id="ap-y-hint">County average, not your farm. A normal year often comes in a bit above this.</p></div>
    <div class="ap-f ap-f-wide"><label for="ap-p" id="ap-p-l">Expected season-average price ($/bu)</label><input id="ap-p" type="number" inputmode="decimal" min="0" step="0.01" placeholder="You set this" aria-describedby="ap-p-hint ap-fut"><p class="ap-hint" id="ap-p-hint">USDA national season-average price for the marketing year, not your local cash price.</p><p class="ap-hint ap-fut" id="ap-fut"></p></div>
    <div class="ap-f"><label for="ap-farm">FSA farm number (optional)</label><input id="ap-farm" type="text" inputmode="numeric" autocomplete="off" maxlength="12" aria-describedby="ap-farm-hint"><p class="ap-hint" id="ap-farm-hint">Goes on the printout and keeps each farm&rsquo;s numbers apart on this device.</p></div>
  </div>
  <div class="ap-out" id="ap-out" aria-busy="true"><p class="ap-need">Loading FSA and USDA numbers for the calculator.</p></div>
</section>"""


OUT_FORM_JS = "components/arc-plc-form.js"


def form_parts(D, crop="corn", st="", fips="", embed=False):
    """form_html split into its <section> tag and its inside."""
    f = form_html(D, crop, st, fips, embed)
    i = f.index(">") + 1
    return f[:i], f[i:-len("</section>")]


def form_js(D):
    """components/arc-plc-form.js: the calculator's inside, once for every county page (the county pages
    carry only its <section> tag; components/arc-plc.js puts this in before it starts). Written from
    form_html, so the main page, the embed and the county pages share one form."""
    return ("/* components/arc-plc-form.js: MACHINE-OWNED, written by scripts/build_arc_plc.py (form_js) from form_html.\n"
            "   The ARC or PLC calculator's inputs, for the county pages, which carry only the <section> it goes in. */\n"
            "window.AgArcForm = " + json.dumps(form_parts(D)[1], ensure_ascii=False).replace("</", "<\\/") + ";\n")


def form_shell(D, crop, st, fips):
    """A county page's calculator: the section tag, filled from components/arc-plc-form.js."""
    tag, _inner = form_parts(D, crop, st, fips)
    return (tag[:-1] + ' data-form="1">\n  <p class="ap-need">The calculator did not load. Check the connection and reload the page, or '
            f'<a href="/arc-plc?st={esc(st)}&amp;fips={esc(fips)}&amp;crop={esc(crop)}#calculator">open it on the main page</a>.</p>\n</section>')


# ---------------------------------------------------------------- main page
def erp_math_html(c):
    e = c["erp"]
    lo_done = hi_done = False
    parts = []
    for y, v in c["mya"].items():
        if not lo_done and v == e["dropped_low"]:
            parts.append(f'<s>{y}: {usd(v)}</s><span class="sr-only"> (dropped)</span>'); lo_done = True
        elif not hi_done and v == e["dropped_high"]:
            parts.append(f'<s>{y}: {usd(v)}</s><span class="sr-only"> (dropped)</span>'); hi_done = True
        else:
            parts.append(f"{y}: {usd(v)}")
    if e["binding"] == "statutory":
        verdict = (f'88% of the average is {usd(e["pct_value"])}, below the {usd(e["statutory"])} statutory price, '
                   f'so the statutory price holds: <b>{usd(e["erp"])}</b>.')
    elif e["binding"] == "cap":
        verdict = f'88% of the average is {usd(e["pct_value"])}, above the cap, so the cap holds: <b>{usd(e["erp"])}</b>.'
    else:
        verdict = (f'88% of that is <b>{usd(e["erp"])}</b>, above the {usd(e["statutory"])} statutory price and under '
                   f'the {usd(e["cap"])} cap (115%).')
    return (f'<p class="ap-math"><b>{esc(c["label"])}.</b> Season-average prices {" &middot; ".join(parts)} (struck: the high and low year). '
            f'Average of the middle three: {usd(e["olympic_avg"], 4)}. {verdict}</p>')


def bp_math_html(c):
    return (f'<p class="ap-math"><b>{esc(c["label"])}.</b> ' + " &middot; ".join(
        f"{y}: {usd(v)}" + (f" (raised from {usd(c['mya'][y])})" if v != c["mya"][y] else "") for y, v in c["bp"]["used"].items())
        + f'. Drop the high ({usd(c["bp"]["dropped_high"])}) and low ({usd(c["bp"]["dropped_low"])}); the middle three average <b>{usd(c["bp"]["value"])}</b>.</p>')


def worked_example(D):
    c = D["years"]["2027"]["crops"]["corn"]
    erp, bp, loan = c["erp"]["erp"], c["bp"]["value"], c["loan"]
    by, py_, base, y, p = 180.0, 150.0, 100.0, 162.0, 4.00
    pr, ar = plc_rate(erp, p, loan), arc_rate(by, bp, y, p, loan)
    br = by * bp
    g = 0.9 * br
    txt = " ".join(ranges_text(ranges(erp, loan, py_, by, bp, y)))
    return f"""<ol class="ap-steps">
  <li><b>PLC.</b> Rate = {usd(erp)} &minus; {usd(p)} = <b>{usd(pr)}/bu</b>. Per base acre: {usd(pr)} &times; {py_:.0f} bu PLC yield &times; 85% = <b>{usd(per_base(pr * py_))}</b>.</li>
  <li><b>ARC-CO.</b> Benchmark revenue = {by:.0f} bu &times; {usd(bp)} = {usd(br)}. Guarantee (90%) = {usd(g)}. Actual revenue = {y:.0f} bu &times; {usd(p)} = {usd(y * p)}.
  Shortfall {usd(g - y * p)}, capped at 12% of benchmark ({usd(0.12 * br)}): <b>{usd(ar)} per acre</b> before the 85% factor; <b>{usd(per_base(ar))}</b> per base acre.</li>
  <li><b>On 100 base acres:</b> PLC {usd(per_base(pr * py_) * base, 0)}, ARC-CO {usd(per_base(ar) * base, 0)}, before the 5.7% sequestration cut.</li>
  <li><b>Where they cross at a {y:.0f} bu county yield:</b> {esc(txt)}</li>
</ol>"""


def faq_items(D):
    y6, y7 = D["years"]["2026"]["crops"], D["years"]["2027"]["crops"]
    return [
        ("Should I pick ARC or PLC for 2026 or 2027?",
         f"It comes down to where the season-average price lands and how your county yields. PLC pays when the national price ends below "
         f"the effective reference price (2027 est.: {usd(y7['corn']['erp']['erp'])} corn, {usd(y7['soybeans']['erp']['erp'])} soybeans, "
         f"{usd(y7['wheat']['erp']['erp'])} wheat). ARC-CO pays when county revenue falls below 90% of its benchmark, so it helps more when "
         f"the county yield drops while the price holds. The calculator gives a verdict for your farm (pick, leans, or close, either one) with "
         f"the expected payment per base acre for each, figured across past years' prices and your county's yields."),
        ("What is the 2027 effective reference price for corn?",
         f"{usd(y7['corn']['erp']['erp'])} per bushel by our math from final USDA prices (est.): 88% of the 2021 to 2025 Olympic average "
         f"season-average price, above the {usd(y7['corn']['statutory'])} statutory price. Soybeans {usd(y7['soybeans']['erp']['erp'])}, "
         f"wheat {usd(y7['wheat']['erp']['erp'])}. For 2026 the figures are {usd(y6['corn']['erp']['erp'])}, {usd(y6['soybeans']['erp']['erp'])} "
         f"and {usd(y6['wheat']['erp']['erp'])}, matching published FSA-based figures. FSA publishes the official 2027 numbers."),
        ("When is the ARC/PLC deadline?",
         f"2026: September 16 to December 11, 2026. 2027: November 2, 2026 to March 15, 2027. {signup_copy(build_date())['faq']}"),
        ("What happens if I don't enroll?",
         "For 2026, a farm with no election by December 11, 2026 keeps its 2025 election but gets no 2026 payment. "
         "Keeping your old choice still takes a signed contract."),
        ("Can I sign a multi-year contract?",
         "Yes, for 2026 through 2031. It stays in force while the farm's records and producers do not change. The 2026 election "
         "itself must be unanimous among the producers on the farm; from 2027 the election can be changed crop by crop in each year's signup."),
        ("How is ARC-CO paid?",
         "ARC-CO compares county revenue (county yield times the higher of the national season-average price or the loan rate) with 90% "
         "of the county benchmark revenue. The shortfall is paid, up to 12% of benchmark revenue, on 85% of your base acres. A farm with "
         "irrigated and non-irrigated county figures is weighted by its historical irrigated percentage."),
        ("Can I buy SCO if I pick ARC?",
         "Under the 2025 law, yes: SCO is no longer tied to your ARC or PLC election. For 2026 the coverage up to 90% is sold as ECO; "
         "SCO itself goes to 90% in 2027. Confirm the details for your crop with your crop insurance agent."),
        ("Who gets the payment on rented ground?",
         "On a cash lease the tenant gets the ARC or PLC payment unless the two agree otherwise. On a share lease the landlord and tenant "
         "split it by their shares, and both sign the contract."),
        ("When do the payments come?",
         "After the marketing year ends and USDA publishes the season-average price: after October 1, 2027 for 2026 crops and after "
         "October 1, 2028 for 2027 crops."),
        ("What is the ARC/PLC payment limit?",
         "$155,000 per person or legal entity in the 2025 law, adjusted for inflation; USDA put 2025 at $160,000. The 2026 and 2027 figures "
         "are not out. The limit is per person across all crops and farms."),
    ]


TIERS = ("plc", "lean_plc", "close", "lean_arc", "arc", "none", "withheld", "noprice", "pending", "mismatch")
TIER_HEAD = {"plc": "Pick PLC", "lean_plc": "Leans PLC", "close": "Close", "lean_arc": "Leans ARC-CO", "arc": "Pick ARC-CO",
             "none": "Neither pays", "withheld": "Too few years", "noprice": "No 2027 price", "pending": "Waiting on USDA", "mismatch": "No typical farm"}


def tallies(D):
    """{(year, crop): {verdict: county and practice combinations}} for the typical farm."""
    t = {}
    for S in D["states"].values():
        for c in S["c"]:
            for (k, i, y), v in c["verdicts"].items():
                d_ = t.setdefault((y, k), {})
                d_[v["verdict"]] = d_.get(v["verdict"], 0) + 1
    return t


def tally_sentence(D):
    t = tallies(D).get(("2026", "corn"), {})
    n = sum(t.get(w, 0) for w in ("plc", "arc", "close", "lean_plc", "lean_arc", "none"))
    if not n:
        return ""
    return (f'<p><b>For a typical farm, 2026 corn</b>, across {n:,} county and practice combinations: pick ARC-CO {t.get("arc", 0):,}, leans ARC-CO '
            f'{t.get("lean_arc", 0):,}, close {t.get("close", 0):,}, leans PLC {t.get("lean_plc", 0):,}, pick PLC {t.get("plc", 0):,}'
            + (f', neither expected to pay {t["none"]:,}' if t.get("none") else "") + '. Your own PLC yield '
            f'moves the answer; the calculator below gives yours, with the expected payments. <a href="#method">How this is figured</a>.</p>')


def dy_source_text(D):
    """Which FSA file each scenario year's county yield comes from (E11)."""
    m = D.get("_dysrc") or {}
    act = [t for t in sorted(m) if m[t] == "act"]
    win = [(t, m[t]) for t in sorted(m) if m[t] not in ("act", None)]
    out = []
    if win:
        out.append(", ".join(f"{t} from the program year {hp} file" for t, hp in win))
        out[-1] = (f"{win[0][0]} to {win[-1][0]}: the trend-adjusted yield (county yield or 80% of T-yield) in the newest FSA file that lists it ("
                   + out[-1] + "), divided by that file&rsquo;s benchmark")
    if act:
        out.append(f"{act[0]} to {act[-1]}: FSA&rsquo;s actual county yield in that year&rsquo;s own program year file, with no 80% plug and no trend "
                   f"adjustment, divided by that year&rsquo;s benchmark, whose five-year window leaves the year out")
    return "; ".join(out)


def method_html(D):
    """How the verdict is figured, the scenario inputs, the county tally, and the back-test."""
    y6, y7 = D["years"]["2026"]["crops"], D["years"]["2027"]["crops"]
    if not all(y6[k].get("scen") for k in CROPS) or not all(y7[k].get("scen") for k in ("corn", "soybeans")):
        return ('<h2 id="method">How the ARC or PLC verdict is figured</h2><p>Not available: the price inputs for the scenarios are not loaded.</p>')
    cols = [(y, k) for y in ("2026", "2027") for k in CROPS if D["years"][y]["crops"][k]["scen"].get("center")]
    yrs = sorted({int(t) for y, k in cols for t in D["years"][y]["crops"][k]["scen"]["ratios"]})
    def cell(y, k, t):
        sc = D["years"][y]["crops"][k]["scen"]
        r = sc["ratios"].get(t, sc["ratios"].get(str(t)))
        return f'<td class="num">{usd(sc["center"] * r) if r is not None else "n/a"}</td>'
    rat = "".join(f'<tr><th scope="row">{t}</th>' + "".join(cell(y, k, t) for y, k in cols) + "</tr>" for t in yrs)
    tally = tallies(D)
    tl = "".join(f'<tr><th scope="row">{y} {D["years"][y]["crops"][k]["label"].lower()}</th>' + "".join(f'<td class="num">{tally.get((y, k), {}).get(w, 0):,}</td>' for w in TIERS) + "</tr>"
                 for y in ("2026", "2027") for k in CROPS)
    bt = D.get("_backtest")
    btab = ""
    if bt:
        have = sorted({r["y"] for r in bt["rows"] if r["arc"] is not None})
        miss = [y for y in BACKTEST_YEARS if y not in have]
        rows = "".join(
            f'<tr><td>{r["y"]}</td><td>{y6[r["k"]]["label"]}</td><td class="num">{usd(r["erp"])}</td><td class="num">{usd(r["mya"])}</td>'
            f'<td class="num">{usd(r["rate"]) if r["rate"] else "none"}</td><td class="num">{usd(r["plc"]) if r["plc"] else "none"}</td>'
            f'<td class="num">{(usd(r["arc"]) if r["arc"] else "none") if r["arc"] is not None else "not in our files"}</td></tr>' for r in bt["rows"])
        btab = (f'<h3 id="backtest">What each program paid, nationally, per base acre</h3>'
                f'<div class="ap-scroll"><table class="tbl ap-t"><thead><tr><th>Year</th><th>Crop</th><th class="num">ERP</th><th class="num">Final MYA</th>'
                f'<th class="num">PLC rate</th><th class="num">PLC $/base ac</th><th class="num">ARC-CO $/base ac</th></tr></thead><tbody>{rows}</tbody></table></div>'
                f'<p class="ap-small">PLC: ERP minus the final season-average price (NASS), never below the loan rate, times FSA&rsquo;s {bt["plc_py"]} national average PLC yield '
                f'(corn {bt["nat_py"]["corn"]:.1f}, soybeans {bt["nat_py"]["soybeans"]:.1f}, wheat {bt["nat_py"]["wheat"]:.1f} bu, weighted by enrolled base) times 85%. 2019 to 2024 use the '
                f'2018 law (statutory $3.70, $8.40, $5.50; 85% escalator); 2025 the 2025 law. ARC-CO: FSA&rsquo;s county payment rates from its program year files '
                f'({", ".join(str(y) for y in have)}), irrigated and non-irrigated averaged equally where FSA splits a county, weighted by each county&rsquo;s enrolled base, times 85%. '
                f'Every county rate in those files was recomputed from the file&rsquo;s own yields and prices and matched FSA&rsquo;s to the cent '
                f'({sum((D.get("_paid_checked") or {}).values()):,} rows).'
                + (f' Not in our files: {", ".join(str(y) for y in miss)}.' if miss else "") + ' National averages hide wide county differences.</p>')
    c6, s6, c7, s7, w6 = y6["corn"]["scen"], y6["soybeans"]["scen"], y7["corn"]["scen"], y7["soybeans"]["scen"], y6["wheat"]["scen"]
    corn_sd = sd1(v for v in c6["ratios"].values()) if c6.get("ratios") else None
    def altp(sc):
        return (f'{usd(sc["alt"])} ({esc(sc.get("alt_contract") or "")} October average so far, {usd(sc["alt_fut"])} over {sc.get("alt_days")} of {sc.get("alt_of")} trading days, '
                f'&times; {sc["oct_mean"]:.3f}, the {yrs[0]} to {yrs[-1]} average of final MYA / October futures)') if sc.get("alt") else "not available"
    return f"""<h2 id="method">How the ARC or PLC verdict is figured</h2>
  <p>The verdict compares what each program is expected to pay per base acre across past years ({yrs[0]} to {yrs[-1]}, each year with a final USDA price), not one guess.</p>
  <ul class="ap-list">
    <li><b>2026 price, every crop.</b> Start: USDA&rsquo;s projected 2026/27 season-average price ({c6.get('usda_src', 'USDA WASDE')}): {usd(c6['center'])} corn,
    {usd(s6['center'])} soybeans, {usd(w6['center'])} wheat. Corn and soybean spread: how far each past year&rsquo;s final season-average price (NASS) ended from that
    October&rsquo;s futures average (RMA&rsquo;s harvest price, on our <a href="/harvest-price-tracker#history">price tracker</a>), scaled to average 1: the uncertainty left in October.</li>
    <li><b>Second start, as a check.</b> Corn {altp(c6)}; soybeans {altp(s6)}. Where USDA&rsquo;s start and this one disagree on which program leads, or on a clear pick,
    a clear pick drops to Leans. A study of corn found WASDE and futures-based forecasts about equally accurate ({src_link('HOFF15')}).</li>
    <li><b>2026 spread for wheat and the smaller crops.</b> None of them has its own October futures history on file. Wheat uses its own year-to-year price changes,
    scaled to average 1 and then to the size of corn&rsquo;s October swings (standard deviation {corn_sd * 100:.1f}%). The smaller crops have too few years in FSA&rsquo;s
    price tables, so they use corn&rsquo;s October swings. Because the spread is borrowed, their call stops at Leans.</li>
    <li><b>2026 county yield.</b> The 2026 crop is mostly harvested, so every scenario year uses a normal crop: the county&rsquo;s FSA benchmark yield. Only the price moves.
    In the calculator, type the county yield if you know it and the 2026 call uses it.</li>
    <li><b>2027 price, corn and soybeans.</b> Start: {c7['center_label']}: {usd(c7['center'])} corn; soybeans {s7['center_label']}: {usd(s7['center'])}. Spread: each past year&rsquo;s change in the
    final season-average price, scaled to average 1, because 2027 is a whole year out.</li>
    <li><b>2027, wheat and every other crop.</b> No 2027 price we can check against past years (no futures history on file for them), so no 2027 call. Reusing the 2026 price
    would only repeat the 2026 answer. Type a 2027 price in the calculator to see the dollars.</li>
    <li><b>2027 county yield.</b> For the same year as the price, the county&rsquo;s yield against its benchmark, from FSA&rsquo;s own files. {dy_source_text(D)}.
    A county with fewer than {MIN_SCEN} years uses its state&rsquo;s average for that crop and practice (at least {MIN_STATE_N} counties a year).</li>
    <li><b>The call.</b> With at least {MIN_SCEN} years: <b>Pick</b> when the expected gap is more than t times its standard error (t from Student&rsquo;s t for the number of years,
    2.26 for 10 and 2.23 for 11), at least {usd(MIN_GAP)} per base acre, and the leading program paid more in at least {MIN_WINS} of the years; <b>Leans</b> when the gap is at least
    {usd(LEAN_GAP)} and at least one standard error, and the leader paid more in at least {MIN_WINS} years; otherwise <b>close, either one</b>. When both expected payments round
    to $0: <b>neither expected to pay</b>. Expected dollars are always shown beside the call.</li>
    <li><b>Typical farm</b> (county pages): FSA&rsquo;s county average PLC yield and FSA&rsquo;s official benchmark. That PLC yield covers all practices, irrigated and not, so the
    irrigated answer is labeled as using it. Where it is above a practice&rsquo;s benchmark, no typical answer is given for that practice. The county page&rsquo;s Lower and Higher
    buttons use 85% and 115% of it.</li>
  </ul>
  <h3>Scenario prices ({len(yrs)} years)</h3>
  <div class="ap-scroll"><table class="tbl ap-t"><thead><tr><th>Year</th>{''.join(f'<th class="num">{y} {D["years"][y]["crops"][k]["label"].lower()}</th>' for y, k in cols)}</tr></thead><tbody>{rat}</tbody></table></div>
  <p class="ap-small">Season-average price used in each scenario year, $/bu.</p>
  <h3>Typical-farm verdicts, all counties</h3>
  <div class="ap-scroll"><table class="tbl ap-t"><thead><tr><th>Year, crop</th>{''.join(f'<th class="num">{TIER_HEAD[w]}</th>' for w in TIERS)}</tr></thead><tbody>{tl}</tbody></table></div>
  <p class="ap-small">County and practice combinations. Your own PLC yield moves the answer; run it.</p>
  {btab}"""


def build_main(D, ch):
    hdr, ftr, fonts = ch
    y6, y7 = D["years"]["2026"]["crops"], D["years"]["2027"]["crops"]
    c7, s7, w7 = y7["corn"], y7["soybeans"], y7["wheat"]
    faq = faq_items(D)
    fsa = D["fsa"]
    n_c = sum(len(S["c"]) for S in D["states"].values())
    title = "ARC vs PLC 2026 and 2027 Calculator | AGSIST"
    desc = (f"Est. 2027 effective reference prices: corn {usd(c7['erp']['erp'])}, soybeans {usd(s7['erp']['erp'])}, wheat "
            f"{usd(w7['erp']['erp'])}. " + ("FSA official county benchmarks. " if fsa else "Bring your county's FSA benchmark. ") + signup_copy(build_date())["desc"])
    og_title = f"ARC or PLC: est. 2027 corn ERP {usd(c7['erp']['erp'])}, soybeans {usd(s7['erp']['erp'])}"
    jsonld = [
        {"@context": "https://schema.org", "@type": "WebApplication", "@id": f"{SITE}/arc-plc#app",
         "name": "ARC or PLC: 2026 and 2027 decision calculator", "url": f"{SITE}/arc-plc",
         "applicationCategory": "FinanceApplication", "operatingSystem": "Any (web browser)",
         "isAccessibleForFree": True, "offers": {"@type": "Offer", "price": "0", "priceCurrency": "USD"},
         "description": desc, "creator": {"@type": "Organization", "name": "AGSIST", "url": SITE}, "dateModified": D["updated"][:10]},
        {"@context": "https://schema.org", "@type": "FAQPage",
         "mainEntity": [{"@type": "Question", "name": q, "acceptedAnswer": {"@type": "Answer", "text": a}} for q, a in faq]},
        {"@context": "https://schema.org", "@type": "Dataset", "@id": f"{SITE}/arc-plc#data",
         "name": "PLC effective reference prices and ARC-CO benchmark prices 2026 and 2027, with FSA county benchmark yields",
         "description": ("Effective reference prices and ARC-CO benchmark prices computed from USDA NASS final season-average prices under "
                         "7 CFR 1412.3 as amended for the 2025 law, plus FSA's official ARC-CO county benchmark yields"
                         + (f" for program year {fsa['py']} ({n_c} counties)." if fsa else ".")),
         "url": f"{SITE}/arc-plc", "isAccessibleForFree": True, "license": "https://creativecommons.org/licenses/by/4.0/",
         "creator": {"@type": "Organization", "name": "AGSIST", "url": SITE},
         "isBasedOn": ["https://quickstats.nass.usda.gov/", FSA_DATA, SRC["FR26"][1]],
         "distribution": {"@type": "DataDownload", "encodingFormat": "application/json", "contentUrl": f"{SITE}/{OUT_JSON}"},
         "temporalCoverage": "2020/2027", "dateModified": D["updated"][:10]},
        {"@context": "https://schema.org", "@type": "BreadcrumbList", "itemListElement": [
            {"@type": "ListItem", "position": 1, "name": "Home", "item": f"{SITE}/"},
            {"@type": "ListItem", "position": 2, "name": "Tools", "item": f"{SITE}/tools"},
            {"@type": "ListItem", "position": 3, "name": "ARC or PLC", "item": f"{SITE}/arc-plc"}]},
        {"@context": "https://schema.org", "@type": "WebPage", "@id": f"{SITE}/arc-plc", "name": title, "url": f"{SITE}/arc-plc", "description": desc,
         "author": AUTHOR, "dateModified": D["updated"][:10], "inLanguage": "en-US"},
    ]
    erp_rows = "".join(
        f'<tr><th scope="row">{esc(y7[k]["label"])}</th><td class="num">{usd(y6[k]["erp"]["erp"])}</td><td class="num"><b>{usd(y7[k]["erp"]["erp"])}</b></td>'
        f'<td class="num">{usd(y6[k]["bp"]["value"])}</td><td class="num">{usd(y7[k]["bp"]["value"])}</td><td class="num">{usd(y7[k]["statutory"])}</td>'
        f'<td class="num">{usd(y7[k]["prior_statutory"])}</td><td class="num">{usd(y7[k]["loan"])}</td></tr>' for k in CROPS)
    faq_html = "".join(f"<details><summary>{esc(q)}</summary><p>{esc(a)}</p></details>" for q, a in faq)
    state_links = " &middot; ".join(f'<a href="/arc-plc/{S["slug"]}">{esc(S["n"])}</a>'
                                    for _k, S in sorted(D["states"].items(), key=lambda kv: kv[1]["n"]))
    srcs = "".join(f"<li>{src_link(k)}</li>" for k in SRC) + "".join(
        f'<li><a href="{esc(u)}" rel="noopener">{esc(n)}</a></li>' for n, u in XC.TABLE_URLS.values())
    fsa_line = (f"FSA&rsquo;s official ARC-CO benchmark yields for program year {fsa['py']} (FSA file as of {nice_date(fsa['as_of'])}) cover "
                f"{n_c:,} counties." if fsa else "FSA&rsquo;s official county benchmark file is not loaded yet; type the benchmark your county office gives you.")
    body = f"""
<body>
{hdr}
<main class="ap-wrap" id="main">
  {bc_html([("Home", "/"), ("Tools", "/tools"), ("ARC or PLC", None)])}
  <p class="page-kicker">Farm program &middot; 2026 and 2027 crop years</p>
  <h1>ARC or PLC for 2026 and 2027: decision calculator</h1>
  {byline_html(D)}
  <h2 class="ap-find-h">Find your county&rsquo;s short answer</h2>
  {picker_html(D)}
  {deadlines_html()}
  <div class="ap-quick" id="quick-answer">
    <p><b>Quick answer.</b> By our math from final USDA prices, the 2027 PLC effective reference prices are <b>{usd(c7['erp']['erp'])}</b> for corn,
    <b>{usd(s7['erp']['erp'])}</b> for soybeans and <b>{usd(w7['erp']['erp'])}</b> for wheat (2026: {usd(y6['corn']['erp']['erp'])},
    {usd(y6['soybeans']['erp']['erp'])}, {usd(y6['wheat']['erp']['erp'])}). PLC pays if the national season-average price ends below that.
    ARC-CO pays when county revenue falls below 90% of its benchmark. {signup_copy(build_date())['quick']}</p>
    {tally_sentence(D)}
    <p class="ap-small">{fsa_line} {asof(D['updated'])}</p>
  </div>
  <p class="page-lede">Pick the year, crop and county, type your PLC yield, and set a price. The table shows which program pays more per base acre across
  a range of prices and county yields. It does not forecast the price. You set it.</p>

  {form_html(D)}

  <div class="ap-faq-sections">
  <h2 id="dates">When is the ARC/PLC signup for 2026 and 2027?</h2>
  <ul class="ap-list">
    <li><b>2026:</b> Sept 16 to Dec 11, 2026. The 2026 election must be unanimous among the producers on the farm.</li>
    <li><b>2027:</b> Nov 2, 2026 to Mar 15, 2027. From 2027 you can change the election crop by crop in each year&rsquo;s signup.</li>
    <li><b>Multi-year contract, 2026 through 2031:</b> one signature covers the years, and it stays in force while the farm&rsquo;s records and producers do not change.
    What it saves is paperwork; the yearly chance to switch programs from 2027 on still applies through the annual election. If the operator, owners or farm records change, it ends and you enroll yearly.
    Some advisers prefer signing each year so nothing is assumed.</li>
    <li><b>If you don&rsquo;t sign:</b> for 2026, a farm with no election by Dec 11, 2026 keeps its 2025 election and gets no 2026 payment.</li>
    <li><b>When payments come:</b> after Oct 1, 2027 for 2026 crops; after Oct 1, 2028 for 2027 crops.</li>
  </ul>

  </div>
  <h2 id="erp">2026 and 2027 effective reference prices</h2>
  <div class="ap-scroll"><table class="tbl ap-t">
    <thead><tr><th>Crop</th><th class="num">2026 ERP</th><th class="num">2027 ERP (est.)</th><th class="num">2026 ARC price</th><th class="num">2027 ARC price (est.)</th><th class="num">Statutory now</th><th class="num">Before 2025</th><th class="num">Loan rate</th></tr></thead>
    <tbody>{erp_rows}</tbody>
  </table></div>
  <p class="ap-small">$ per bushel. ERP = PLC effective reference price. ARC price = ARC-CO benchmark price. Est. = our math from final USDA prices; FSA publishes the
  official 2027 figures. 2026 figures match those published by FSA-based sources{" and FSA&rsquo;s 2026 county file" if fsa else ""}.</p>
  <details class="ap-det"><summary>Show the math for 2027</summary>
  <p>The effective reference price is the lesser of 115% of the statutory price, or the greater of the statutory price and 88% of the
  Olympic average season-average price for the five most recent crop years (2021 to 2025 for 2027; 2020 to 2024 for 2026).</p>
  {''.join(erp_math_html(y7[k]) for k in CROPS)}
  <p>The ARC-CO benchmark price uses the same five years, with any year below the 2027 effective reference price raised to it:</p>
  {''.join(bp_math_html(y7[k]) for k in CROPS)}
  </details>

  {XC.other_crops_table(D)}

  <div class="ap-faq-sections">
  <h2 id="changed">What changed for 2026 and 2027?</h2>
  <p class="ap-small">As of {nice_date(D['updated'][:10])}. The 2025 law (One Big Beautiful Bill Act, Pub. L. 119-21) and FSA&rsquo;s January 2026 rule:</p>
  <ul class="ap-list">
    <li><b>Higher reference prices:</b> corn $3.70 to $4.10, soybeans $8.40 to $10.00, wheat $5.50 to $6.35, for 2025 through 2030.</li>
    <li><b>Effective reference price:</b> 88% of the Olympic average (was 85%), capped at 115% of statutory.</li>
    <li><b>ARC-CO:</b> guarantee 90% of benchmark revenue (was 86%); payments up to 12% of benchmark (was 10%).</li>
    <li><b>New base acres:</b> up to 30 million, from 2019 to 2023 plantings, cut 3.69% on the new acres only. New base takes the farm&rsquo;s existing PLC yield for that crop, or the county average PLC yield if the farm has none.</li>
    <li><b>2025 only:</b> farms got the higher of ARC-CO or PLC automatically. That does not carry into 2026 or 2027.</li>
    <li><b>SCO:</b> no longer tied to the ARC/PLC election; for 2026 the band to 90% is sold as ECO, and SCO goes to 90% in 2027. Ask your crop insurance agent.</li>
    <li><b>Payment limit:</b> $155,000 per person, indexed (2025: $160,000). Actively engaged members of LLCs and S corporations can each carry a limit.</li>
    <li><b>Signup moved:</b> the 2026 window opened Sept 16, 2026, later than the usual spring signup.</li>
  </ul>

  <h2 id="how">How do ARC-CO and PLC pay?</h2>
  <ul class="ap-list">
    <li><b>PLC rate</b> = effective reference price &minus; the higher of the national season-average price or the loan rate. Paid on 85% of base acres &times; your PLC payment yield.</li>
    <li><b>ARC-CO rate</b> = 90% of benchmark revenue &minus; actual county revenue, never more than 12% of benchmark revenue. Paid on 85% of base acres.</li>
    <li><b>Benchmark revenue</b> = FSA county benchmark yield &times; benchmark price. Actual revenue = actual county yield &times; the higher of the season-average price or the loan rate.</li>
  </ul>
  <h3>A worked example (made-up corn farm, 2027 prices)</h3>
  <p>County benchmark yield 180 bu, PLC yield 150 bu, 100 base acres, a county yield of 162 bu and a $4.00 season-average price.</p>
  {worked_example(D)}

  <h2 id="checklist">What should I check before I sign?</h2>
  <ul class="ap-list">
    <li><b>Each crop, each farm number.</b> You elect ARC-CO or PLC crop by crop on each FSA farm. Run them one at a time.</li>
    <li><b>Rented ground.</b> On a cash lease the tenant gets the payment unless agreed otherwise. On a share lease landlord and tenant split it by share and both sign.</li>
    <li><b>10 base acres.</b> A farm with 10 or fewer total base acres gets no payment, unless the producer is socially disadvantaged, limited resource, a veteran or a beginning farmer.</li>
    <li><b>Income rule.</b> The $900,000 average adjusted gross income limit still applies to ARC and PLC. The new exception for those with 75% of gross income from farming covers certain conservation and disaster programs, not ARC or PLC.</li>
    <li><b>Payment limit.</b> Per person across all crops and farms, not per farm. Sequestration (5.7% lately) comes off before the limit.</li>
    <li><b>ARC-IC</b> (individual coverage, whole farm) also exists. This tool does not model it.</li>
  </ul>

  <h2 id="benchmark-yield">Where do I find my county&rsquo;s ARC-CO benchmark yield?</h2>
  <p>{fsa_line} FSA&rsquo;s benchmark is the Olympic average of the county&rsquo;s five most recent yields (2020 to 2024 for 2026), with low years raised to
  80% of the county T-yield and each year trend-adjusted. FSA has not posted 2027 benchmarks; building them needs RMA&rsquo;s 2027 trend factors, which are not out,
  so the 2027 view uses FSA&rsquo;s official 2026 benchmark, labeled. Counties FSA does not list get no page here: ask the county office for
  &ldquo;the ARC-CO benchmark yield for my county, crop and practice&rdquo; or check <a href="{FSA_DATA}" rel="noopener">FSA&rsquo;s program data</a>.</p>
  <p class="ap-states">{state_links}</p>
  </div>

  {method_html(D)}

  <h2 id="faq">Questions farmers ask</h2>
  <div class="ap-faq">{faq_html}</div>

  <p class="ap-disc"><b>This is an estimate, not a USDA determination.</b> Payments depend on FSA&rsquo;s official yields, prices and your farm records.
  Read FSA&rsquo;s <a href="{FSA_PAGE}" rel="noopener">ARC/PLC page</a>, its <a href="{FSA_DATA}" rel="noopener">program data</a>,
  or <a href="{OFFICE}" rel="noopener">find your county FSA office</a>. Dollars are before the 5.7% sequestration cut.
  Run an extension office or co-op site? <a href="/developer#arc-plc">Embed the calculator</a>.</p>
  <details class="ap-det"><summary>Sources</summary><ul class="ap-src">{srcs}<li><a href="https://quickstats.nass.usda.gov/" rel="noopener">USDA NASS Quick Stats</a> (season-average prices)</li></ul></details>
  <p class="ap-small">Related: <a href="/farm-bill">Farm bill tracker</a> &middot; <a href="/breakeven">Break-even calculator</a> &middot; <a href="/harvest-price-tracker">Harvest price and pre-sell</a></p>
</main>
"""
    return head(title, desc, "/arc-plc", jsonld, True, fonts, og_title=og_title) + body + tail(ftr, calc=True), {
        "title": title, "desc": desc, "og_title": og_title, "faq": [q for q, _ in faq]}


# ---------------------------------------------------------------- state and county pages
RATIOS = (0.6, 0.7, 0.8, 0.9, 1.0)


def county_metrics(by, crop_d):
    bp, erp = crop_d["bp"]["value"], crop_d["erp"]["erp"]
    br = by * bp
    return {"br": br, "g": 0.9 * br, "max": 0.12 * br, "max_base": per_base(0.12 * br),
            "trig_y": 0.9 * br / erp, "cap_y": 0.78 * br / erp}


def trigger_c(crop_d):
    """Highest cent at which ARC-CO pays when county yield = benchmark: price < 90% x benchmark price."""
    return int(math.floor(round(0.9 * crop_d["bp"]["value"] * 100, 6) - 1e-6))


def crossover_table(crop_d, by=None):
    """PLC vs ARC-CO at county yield = benchmark, by PLC yield as % of the
    official benchmark. Depends only on that share, so it holds in any county.
    Each row prints every band (ranges_text), worked at the PLC yield the row
    shows, rounded as printed."""
    erp, loan, bp = crop_d["erp"]["erp"], crop_d["loan"], crop_d["bp"]["value"]
    sc = XC.scale(crop_d)
    base_by = by or 100.0
    rows = []
    for r in RATIOS:
        py_ = rnd(base_by * r, 1) if by else base_by * r
        runs = ranges(erp, loan, py_, base_by, bp, base_by, scale=sc)
        lab = f"{int(r * 100)}%" + (f"<br><span class=\"mut\">{XC.yf(crop_d, py_, 1)}</span>" if by else "")
        rows.append(f'<tr><td>{lab}</td><td>{esc(" ".join(ranges_text(runs, sc)))}</td></tr>')
    return ('<div class="ap-scroll"><table class="tbl ap-t ap-x"><thead><tr><th>PLC yield, % of your official benchmark</th>'
            '<th>Which pays more, by season-average price</th></tr></thead><tbody>' + "".join(rows) + "</tbody></table></div>")


def ent_label(e):
    return DLABEL[e["d"]] + (f", {e['sub']}" if e.get("sub") else "")


ASK = ("Ask the county FSA office for &ldquo;the 2026 ARC-CO benchmark yield for my county, crop and practice (irrigated or "
       "non-irrigated)&rdquo;, or look it up in <a href=\"{u}\" rel=\"noopener\">FSA&rsquo;s program data</a> (file: ARC-CO benchmark yields "
       "and revenues, program year 2026). Then type it into the calculator.").format(u=FSA_DATA)


def pending_note(D):
    return (f'<div class="ap-quick ap-pending"><p><b>FSA&rsquo;s official 2026 benchmark yields are not loaded here yet.</b> '
            f'Until they are, this page shows only FSA&rsquo;s own published figures for the county (its history and the county average PLC yield) and '
            f'no current benchmark, trigger yield or break-even in bushels. {ASK}</p></div>')


def hist_table(e, cd=None):
    hist = [h for h in e.get("hist", []) if h["py"] != 2026]
    if not hist:
        return ""
    cd = cd or {"unit": "bu", "dp": 2}
    u = XC.UNIT[cd["unit"]]["short"]
    rows = "".join(
        f'<tr><td>{h["py"]}</td><td class="num">{h["by"]:,.2f}</td><td class="num">{XC.pf(cd, h["bp"]) if h["bp"] else "n/a"}</td>'
        f'<td class="num">{f"{h["ay"]:.2f}" if h["ay"] is not None else "not yet"}</td>'
        f'<td class="num">{("none" if h["pay"] == 0 else usd(h["pay"])) if h["pay"] is not None else "not yet"}</td></tr>' for h in hist)
    return (f'<h3>FSA history, {esc(ent_label(e).lower())}</h3><div class="ap-scroll"><table class="tbl ap-t"><thead><tr><th>Year</th>'
            f'<th class="num">Benchmark</th><th class="num">Price</th><th class="num">County yield</th><th class="num">ARC-CO $/ac</th></tr></thead>'
            f'<tbody>{rows}</tbody></table></div><p class="ap-small">Program year; benchmark yield in {u}/acre; benchmark price $/{u}; actual county '
            f'yield; ARC-CO payment rate per acre before the 85% factor. All from FSA&rsquo;s files. Older program years use older windows and prices, '
            f'so they are history, not this year&rsquo;s benchmark.</p>')


def state_page(st, S, D, ch, nav_states):
    hdr, ftr, fonts = ch
    name, sl = S["n"], S["slug"]
    y6 = D["years"]["2026"]["crops"]
    fsa = D["fsa"]
    official = bool(fsa)
    indexable = official and len(S["c"]) >= MIN_STATE_COUNTIES
    path = f"/arc-plc/{sl}"
    plcy = D["_plc_yields"]
    plc_py = (D["plc_county"] or {}).get("py")
    if official:
        title = next((t for t in (f"{name} ARC-CO Benchmark Yields 2026, Official | AGSIST",
                                  f"{name} ARC-CO Benchmark Yields 2026 | AGSIST",
                                  f"{name} ARC-CO Benchmarks 2026") if len(t) <= 60), f"{name} ARC-CO 2026")
        desc = (f"FSA official 2026 ARC-CO benchmark yields for {len(S['c'])} {name} counties by crop and practice, with benchmark revenue "
                f"and the PLC vs ARC break-even.")
    else:
        title = next((t for t in (f"{name} ARC or PLC 2026 and 2027 by County | AGSIST", f"{name} ARC or PLC by County | AGSIST")
                      if len(t) <= 60), f"{name} ARC or PLC")
        desc = f"{name} counties: FSA average PLC yields and ARC-CO history, and how to get your official 2026 benchmark."
    rows = []
    for c in S["c"]:
        if not official:
            pv = plcy.get(c["f"], {})
            rows.append(f'<tr><th scope="row"><a href="/arc-plc/{sl}/{c["s"]}">{esc(re.sub(r" (County|Parish)$", "", c["n"]))}</a></th>'
                        + "".join(f'<td class="num">{f"{pv[k]:.1f}" if pv.get(k) else "n/a"}</td>' for k in CROPS) + "</tr>")
            continue
        first = True
        nrow = sum(len(c["k"].get(k, [])) for k in ALL if k in y6)
        for k in ALL:
            if k not in y6:
                continue
            for e in c["k"].get(k, []):
                link = f'<a href="/arc-plc/{sl}/{c["s"]}">{esc(re.sub(r" (County|Parish)$", "", c["n"]))}</a>' if first else ""
                py_avg = plcy.get(c["f"], {}).get(k)
                if official:
                    # kept short: Texas lists about 1,500 county, crop and practice rows (the
                    # maximum payment is 10.2% of revenue per base acre; each county page shows it)
                    m = county_metrics(e["by"], y6[k])
                    th = (f'<th scope="row"{f" rowspan={nrow}" if nrow > 1 else ""}>{link}</th>') if first else ""
                    rows.append(f'<tr>{th}<td>{esc(y6[k]["label"])}</td>'
                                f'<td>{PRAC[e["d"]]}{", " + esc(e["sub"]) if e.get("sub") else ""}</td><td>{XC.yf(y6[k], e["by"])}</td><td>{usd(m["br"], 0)}</td></tr>')
                    first = False
                    continue
                else:
                    cells = f'<td class="num">{f"{py_avg:.1f}" if py_avg else "n/a"}</td>'
                rows.append(f'<tr><th scope="row">{link}</th><td>{esc(y6[k]["label"])}</td><td>{esc(ent_label(e))}</td>{cells}</tr>')
                first = False
    head_cells = ('<th class="num">Yield</th><th class="num">Revenue</th>' if official
                  else f'<th class="num">Avg PLC yield{f" ({plc_py})" if plc_py else ""}</th>')
    jsonld = [{"@context": "https://schema.org", "@type": "BreadcrumbList", "itemListElement": [
        {"@type": "ListItem", "position": 1, "name": "Home", "item": f"{SITE}/"},
        {"@type": "ListItem", "position": 2, "name": "ARC or PLC", "item": f"{SITE}/arc-plc"},
        {"@type": "ListItem", "position": 3, "name": name, "item": f"{SITE}{path}"}]},
        {"@context": "https://schema.org", "@type": "WebPage", "@id": f"{SITE}{path}", "name": title, "url": f"{SITE}{path}", "description": desc,
         "author": AUTHOR, "dateModified": D["updated"][:10], "inLanguage": "en-US"}]
    if official:
        jsonld.insert(0, {"@context": "https://schema.org", "@type": "Dataset", "@id": f"{SITE}{path}#dataset",
                          "name": f"{name} 2026 ARC-CO benchmark yields by county (FSA official)",
                          "description": (f"FSA's official program year 2026 ARC-CO benchmark yields for {len(S['c'])} {name} counties by crop and "
                                          f"practice, with benchmark revenue at the 2026 benchmark price, from FSA's county file as of {fsa['as_of']}."),
                          "url": f"{SITE}{path}", "isAccessibleForFree": True, "creator": {"@type": "Organization", "name": "AGSIST", "url": SITE},
                          "isBasedOn": FSA_DATA, "dateModified": D["updated"][:10],
                          "spatialCoverage": {"@type": "Place", "name": f"{name}, United States"},
                          "variableMeasured": ["ARC-CO benchmark yield (bu/acre)", "ARC-CO benchmark revenue ($/acre)"]})
    others = " &middot; ".join(f'<a href="/arc-plc/{x["slug"]}">{esc(x["n"])}</a>' for x in nav_states if x["slug"] != sl)
    rel = ['<a href="/arc-plc">ARC vs PLC calculator</a>']
    for d_, lab in (("rent", "cash rent by county"), ("yield", "yield by year")):
        if os.path.exists(os.path.join(ROOT, d_, f"{sl}.html")):
            rel.append(f'<a href="/{d_}/{sl}">{esc(name)} {lab}</a>')
    if os.path.exists(os.path.join(ROOT, "farmland-atlas", sl, "index.html")):
        rel.append(f'<a href="/farmland-atlas/{sl}/">{esc(name)} Farmland Atlas</a>')
    lede = (f"FSA&rsquo;s official program year 2026 ARC-CO benchmark yields for {len(S['c'])} {esc(name)} counties, by crop and by practice where "
            f"FSA splits irrigated and non-irrigated. FSA has not posted 2027 benchmarks; the 2027 benchmark shifts the window one year. "
            f"{asof(fsa['as_of'], 'FSA file as of')}" if official else
            f"The {len(S['c'])} {esc(name)} counties FSA lists for corn, soybeans or wheat, with FSA&rsquo;s county average PLC yield. {asof(D['updated'])}")
    table_h = "Official benchmark yield and revenue by county" if official else "County average PLC yields (FSA)"
    table_note = ("Benchmark yield per acre (FSA, 2026), in bushels or, for peanuts, rice, pulses, seed cotton and most oilseeds, pounds. Practice: All = all practices, "
                  "Irr. = irrigated, Non-irr. = non-irrigated. Revenue = yield &times; 2026 benchmark price, $/acre. The most ARC-CO can pay is 12% of revenue on 85% of base, "
                  "10.2% of revenue per base acre; each county page works it out."
                  if official else f"Average PLC payment yield on enrolled base in the county by crop, bu/acre, FSA program year {plc_py}. Your farm&rsquo;s own is on the FSA-156EZ.")
    body = f"""
<body>
{hdr}
<main class="ap-wrap" id="main">
  {bc_html([("Home", "/"), ("ARC or PLC", "/arc-plc"), (name, None)])}
  <p class="page-kicker">ARC-CO and PLC &middot; {"FSA official" if official else "by county"}</p>
  <h1>{esc(name)} {"ARC-CO benchmark yields for 2026, by county" if official else "ARC or PLC by county, 2026 and 2027"}</h1>
  {byline_html(D)}
  <p class="page-lede">{lede}</p>
  <h2 class="ap-find-h">Find your county&rsquo;s short answer</h2>
  {picker_html(D, st)}
  {deadlines_html()}
  {'' if official else pending_note(D)}
  <div class="ap-quick"><p><b>Prices that apply in every county.</b> 2026 effective reference price: corn {usd(y6['corn']['erp']['erp'])}, soybeans {usd(y6['soybeans']['erp']['erp'])},
  wheat {usd(y6['wheat']['erp']['erp'])}. 2026 ARC-CO benchmark price: corn {usd(y6['corn']['bp']['value'])}, soybeans {usd(y6['soybeans']['bp']['value'])}, wheat {usd(y6['wheat']['bp']['value'])}.
  When the county yield comes in at its benchmark, ARC-CO pays when the season-average price ends at {c2(trigger_c(y6['corn']))} or lower for corn,
  {c2(trigger_c(y6['soybeans']))} for soybeans and {c2(trigger_c(y6['wheat']))} for wheat.</p>{XC.state_crops_line(D, S)}</div>
  <h2 id="counties">{table_h}</h2>
  <p class="ap-small">{table_note}</p>
  <div class="ap-scroll"><table class="tbl ap-t ap-st">
    <thead><tr><th>County</th>{('<th>Crop</th><th>Practice</th>' + head_cells) if official else ''.join(f'<th class="num">{CROPS[k]["label"]}</th>' for k in CROPS)}</tr></thead>
    <tbody>{''.join(rows)}</tbody></table></div>
  <p class="ap-small">A county or crop not listed: FSA&rsquo;s files have nothing for it. {ASK}</p>
  <h2 id="crossover">PLC vs ARC-CO break-even at the benchmark yield</h2>
  <p>When the county yield comes in at its benchmark, which program pays more depends on the season-average price and on your PLC yield as a share of the
  county&rsquo;s official benchmark. Expressed that way, the break-even is the same in every {esc(name)} county. 2026 prices, corn:</p>
  {crossover_table(y6['corn'])}
  <p>Soybeans:</p>
  {crossover_table(y6['soybeans'])}
  <p>Wheat:</p>
  {crossover_table(y6['wheat'])}
  <p class="ap-small">Per base acre, both programs on 85% of base. Your PLC yield is on the FSA-156EZ. A county yield below benchmark favors ARC-CO.</p>
  <p class="ap-disc">An estimate, not a USDA determination. <a href="{FSA_PAGE}" rel="noopener">FSA ARC/PLC</a> &middot; <a href="{FSA_DATA}" rel="noopener">FSA program data</a> &middot; <a href="{OFFICE}" rel="noopener">Find your county office</a></p>
  <p class="ap-small">Related: {' &middot; '.join(rel)}</p>
  <h2>Other states</h2>
  <p class="ap-states">{others}</p>
</main>
"""
    page = head(title, desc, path, jsonld, indexable, fonts) + body + tail(ftr)
    return page, {"path": path, "indexable": indexable, "title": title, "desc": desc}


SHORT = {"plc": "Pick PLC", "arc": "Pick ARC-CO", "lean_plc": "Leans PLC", "lean_arc": "Leans ARC-CO",
         "close": "Close, either one", "none": "Neither expected to pay"}


def why_text(v, cd, by, py):
    """2-3 plain sentences, every claim with its number."""
    erp, loan, center = cd["erp"]["erp"], cd["loan"], cd["scen"]["center"]
    max_plc = per_base(plc_rate(erp, 0, loan) * py)
    max_arc = per_base(0.12 * by * cd["bp"]["value"])
    short = sum(1 for r in v["rows"] if r["arc"] > 0)
    return (f"PLC pays on any season-average price below {XC.pf(cd, erp)}, up to {usd(max_plc)} per base acre at the {XC.pf(cd, loan)} loan rate, so it "
            f"protects against a deep price drop. ARC-CO is capped at {usd(max_arc)} per base acre but also pays when the county&rsquo;s "
            f"yield is short; it paid something in {short} of the {v['n']} past-year scenarios. With prices centered at {XC.pf(cd, center)}, "
            f"PLC is expected to pay {usd_pay(v['plc'])} and ARC-CO {usd_pay(v['arc'])} per base acre.")


def default_entries(es):
    """Non-irrigated (or all practices) first, irrigated after."""
    order = {"non": 0, "all": 0, "irr": 1}
    return sorted(range(len(es)), key=lambda i: (order[es[i]["d"]], es[i].get("sub", "")))


def crop_order(c, D):
    """The county's crops with a benchmark, the lead crop (most FSA base acres) first, then by base acres."""
    base = ((D.get("_plc_base") or {}).get(c["f"]) or {})
    ks = [k for k in ALL if k in c["k"] and k in D["years"]["2026"]["crops"]]
    return sorted(ks, key=lambda k: (k != c.get("lead"), -(base.get(k) or 0), ALL.index(k)))


def typical_html(c, k, D, cname):
    """The typical-farm verdict block for one crop, and the headline verdict (lead year, default practice)."""
    es = c["k"].get(k) or []
    LY = lead_year()
    OY = "2027" if LY == "2026" else "2026"
    cdl = D["years"][LY]["crops"].get(k) or D["years"]["2026"]["crops"][k]
    cd6 = D["years"]["2026"]["crops"][k]
    pyv = D["_plc_yields"].get(c["f"], {}).get(k)
    plc_py = (D["plc_county"] or {}).get("py")
    why = ((cd6.get("scen") or {}).get("why") or "")
    lines, head_v = [], None
    for i in default_entries(es):
        e = es[i]
        if not e["by"]:
            continue
        vl, vo = c["verdicts"].get((k, i, LY)), c["verdicts"].get((k, i, OY))
        lab = esc(ent_label(e))
        if not vl:
            vl, vo, LY_, OY_ = vo, None, OY, LY
        else:
            LY_, OY_ = LY, OY
        if not vl:
            continue
        if vl["verdict"] == "mismatch":
            lines.append(f'<p class="ap-small"><b>{lab}:</b> no typical-farm verdict. FSA&rsquo;s average PLC yield for the county ({XC.yf(cd6, pyv, 1)}) covers '
                         f'all practices and is above this benchmark ({XC.yf(cd6, e["by"])}), so it does not describe a typical farm here. Run your own PLC yield below.</p>')
            continue
        if head_v is None and vl["verdict"] in SHORT:
            head_v = (vl, e, LY_)
        src = "this county&rsquo;s own FSA yields" if e["dsrc"] == "county" else f"the {esc(STATE_NAMES[c['st']])} average for this crop and practice (the county has fewer than {MIN_SCEN} years)"
        note7 = " The 2027 view uses FSA&rsquo;s 2026 benchmark; FSA posts 2027 later."
        lines.append(f'<div class="ap-verdict {vclass(vl["verdict"])}"><p><b>{lab}, {LY_}: {verdict_line(vl)}</b></p>'
                     + (f'<p>{why_text(vl, D["years"][LY_]["crops"][k], e["by"], pyv)}</p>' if vl["verdict"] not in NOCALL else (f'<p class="ap-small">{why}</p>' if why else ""))
                     + (f'<p class="ap-small"><b>{OY_}:</b> {verdict_line(vo)}{note7 if OY_ == "2027" else ""}</p>' if vo else "")
                     + f'<p class="ap-small">For a typical farm: FSA&rsquo;s county average PLC yield {XC.yf(cd6, pyv, 1)} (program year {plc_py}) and the official '
                       f'{XC.yf(cd6, e["by"])} benchmark.' + (' That PLC yield is FSA&rsquo;s average over all practices, irrigated and not; an irrigated farm&rsquo;s own is often higher.' if e["d"] == "irr" else "")
                       + f' County yields from {src}. Run your own numbers in the calculator below.</p></div>')
    return "".join(lines), head_v


def quick_county(c, D):
    """Quick answer: each crop's typical-farm verdict for the lead year, default practice, numbers beside it."""
    LY = lead_year()
    items = []
    for k in crop_order(c, D):
        es = c["k"].get(k) or []
        for i in default_entries(es):
            v = c["verdicts"].get((k, i, LY))
            if not v or v["verdict"] in ("mismatch",):
                continue
            if v["verdict"] == "withheld" and k not in CROPS:
                break   # these crops are withheld everywhere; their sections say why
            lab = D["years"]["2026"]["crops"][k]["label"] + ("" if es[i]["d"] == "all" else ", " + DLABEL[es[i]["d"]].lower())
            if v["verdict"] in NOCALL:
                items.append(f"<li><b>{esc(lab)}:</b> {VWORD[v['verdict']].lower()}.</li>")
            else:
                items.append(f"<li><b>{esc(lab)}:</b> {VWORD[v['verdict']]} (est. {usd_pay(v['plc'])} PLC vs {usd_pay(v['arc'])} ARC-CO per base acre; "
                             f"PLC more in {v['plc_wins']} of {v['n']} past-year scenarios, ARC-CO in {v['arc_wins']}).</li>")
            break
    if not items:
        return ""
    return (f'<div class="ap-quick" id="quick-answer"><p><b>Quick answer for {LY}, for a typical farm in this county.</b> Run your own numbers; '
            'your PLC yield and base can change it.</p><ul class="ap-list">' + "".join(items) + '</ul>'
            '<p class="ap-small">Typical farm = FSA&rsquo;s county average PLC yield and FSA&rsquo;s official benchmark. Prices: USDA&rsquo;s October WASDE projection, '
            'spread by how far the season-average price ended from October futures in each past year; county yields from the same past years. '
            '<a href="/arc-plc#method">How this is figured</a>.</p></div>')


# ---------------------------------------------------------------- the simple page (county answer cards, finder, counter sheet)
AUTHOR = {"@type": "Person", "name": "Sigurd Lindquist", "url": f"{SITE}/about", "jobTitle": "Certified Crop Adviser (CCA)"}
SHEET_SUFFIX = "-sheet"   # /arc-plc/<state>/<county>-sheet: old counter sheet URLs, now tiny stubs that forward to /arc-plc/sheet
OUT_SHEET = f"{OUT_DIR}/sheet.html"   # /arc-plc/sheet?c=<state>/<county>: the one counter sheet page (noindex), built in the browser
ICON_SVG = {
    "pick": '<svg class="ap-ic" viewBox="0 0 16 16" aria-hidden="true"><circle cx="8" cy="8" r="6.5" fill="currentColor" stroke="currentColor" stroke-width="1.5"/></svg>',
    "lean": ('<svg class="ap-ic" viewBox="0 0 16 16" aria-hidden="true"><circle cx="8" cy="8" r="6.5" fill="none" stroke="currentColor" stroke-width="1.5"/>'
             '<path d="M8 1.5a6.5 6.5 0 0 1 0 13z" fill="currentColor"/></svg>'),
    "close": '<svg class="ap-ic" viewBox="0 0 16 16" aria-hidden="true"><circle cx="8" cy="8" r="6.5" fill="none" stroke="currentColor" stroke-width="1.5"/></svg>',
    "ask": ('<svg class="ap-ic" viewBox="0 0 16 16" aria-hidden="true"><rect x="1.5" y="1.5" width="13" height="13" fill="none" stroke="currentColor" '
            'stroke-width="1.5" stroke-dasharray="2.5 1.8"/></svg>'),
}
# the county pages draw each tier mark from one sprite (icon_sprite) instead of repeating the SVG on every card
ICON = {k: f'<svg class="ap-ic" viewBox="0 0 16 16" aria-hidden="true"><use href="#api-{k}"/></svg>' for k in ICON_SVG}
TIER_WORD = {"pick": "Clear pick", "lean": "Leans", "close": "Close", "ask": "No answer yet"}


def icon_sprite():
    return ('<svg width="0" height="0" style="position:absolute" aria-hidden="true" focusable="false">'
            + "".join(re.sub(r'^<svg class="ap-ic" viewBox="0 0 16 16" aria-hidden="true">(.*)</svg>$', f'<symbol id="api-{k}" viewBox="0 0 16 16">\\1</symbol>', v)
                      for k, v in ICON_SVG.items()) + "</svg>")
PIN = ('<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true">'
       '<circle cx="12" cy="12" r="7"/><circle cx="12" cy="12" r="2.5" fill="currentColor"/><path d="M12 2v3M12 19v3M2 12h3M19 12h3"/></svg>')


def shape(v):
    w = v["verdict"]
    return "pick" if w in ("plc", "arc") else "lean" if w.startswith("lean_") else "close" if w in ("close", "none") else "ask"


def tier_html(v):
    s_ = shape(v)
    return f'<span class="ap-tier ap-tier-{s_}">{ICON[s_]}{TIER_WORD[s_]}</span>'


def key_html():
    return '<p class="ap-key">' + "".join(f'<span>{ICON[s_]}{TIER_WORD[s_]}</span>' for s_ in ("pick", "lean", "close", "ask")) + "</p>"


def big_word(v):
    w = v["verdict"]
    if w in ("plc", "arc"):
        return WORD[w]
    if w.startswith("lean_"):
        return WORD[w[5:]]
    return {"close": "Either one", "none": "Neither expected to pay", "mismatch": "Ask FSA"}.get(w, "No answer yet")


def short_word(v):
    """A few words: the 2027 line and the counter sheet."""
    w = v["verdict"]
    if w in ("plc", "arc"):
        return WORD[w]
    if w.startswith("lean_"):
        return "Leans " + WORD[w[5:]]
    return {"close": "Either one", "none": "Neither pays", "mismatch": "Ask FSA", "pending": "Wait for USDA",
            "noprice": "No 2027 price yet", "withheld": "Not enough history"}[w]


def low1(t):
    """Lower-case the first word unless it is a program name."""
    return t if t.startswith(("ARC", "PLC")) else t[0].lower() + t[1:]


def whole(x):
    return int(rnd(abs(x), 0))


def avg_verb(lc):
    return "average" if lc.endswith("s") else "averages"


def gap_line(v, lc, y, cd, by=None):
    """The card's one plain sentence, whole dollars per base acre."""
    w = v["verdict"]
    if w in ("plc", "arc"):
        return f"{WORD[w]} comes out about ${whole(v['diff'])} more per base acre."
    if w.startswith("lean_"):
        side = w[5:]
        other = "arc" if side == "plc" else "plc"
        ow = v["arc_wins"] if side == "plc" else v["plc_wins"]
        t = f"{WORD[side]} comes out about ${whole(v['diff'])} more per base acre, but that swings from year to year."
        if ow == 0:
            t += f" {WORD[other]} did not come out ahead in any of the {v['n']} price years we ran."
        if v.get("borrowed"):
            t += f" {lc[0].upper() + lc[1:]} {'borrow' if lc.endswith('s') else 'borrows'} corn&rsquo;s price swings, so {'they stop' if lc.endswith('s') else 'it stops'} at Leans."
        if v.get("alt_down"):
            t += " A futures-based price start gives a different answer, so it stops at Leans."
        return t
    if w == "close":
        if cents(abs(v["diff"])) < 100:
            return "They come out less than $1 apart per base acre."
        if abs(v["diff"]) < MIN_GAP:
            return f"They come out about ${whole(v['diff'])} apart per base acre. Too small to matter."
        lead = "plc" if v["diff"] > 0 else "arc"
        return f"{WORD[lead]} comes out about ${whole(v['diff'])} more per base acre, but it swings more than that from year to year."
    if w == "none":
        cen = (cd.get("scen") or {}).get("center")
        return (f"Neither one is likely to pay on {lc} for {y} at USDA&rsquo;s price outlook" + (f" ({XC.pf(cd, cen)})" if cen else "") + ".")
    if w == "mismatch":
        return "FSA&rsquo;s county average PLC yield is above the county benchmark yield, so there is no typical-farm answer. Ask your FSA office."
    if w == "noprice":
        return f"There is no {y} price outlook for {lc} we can check against past years, so no {y} answer yet."
    if w == "pending":
        return "Waiting on USDA&rsquo;s closing 2025 price, which sets the 2027 reference price."
    return "Not enough years of price and county yield history to call it."


def erp_line(cd, lc, y):
    return f"PLC only pays if {lc} {avg_verb(lc)} under {XC.say_price(cd, cd['erp']['erp'])} for {y}."


def year_line(v, y):
    w = v["verdict"]
    t = {"mismatch": "ask your FSA office", "pending": "waiting on USDA&rsquo;s closing 2025 price", "noprice": f"no {y} price yet",
         "withheld": "not enough history to say", "none": "neither expected to pay", "close": "close, either one"}.get(w)
    return f"<b>{y}:</b> " + (t if t else low1(short_word(v)))


def bc_html(items):
    return ('<nav class="ap-bc" aria-label="Breadcrumb">' + " &rsaquo; ".join(
        f'<a href="{h}">{esc(lab)}</a>' if h else f'<span aria-current="page">{esc(lab)}</span>' for lab, h in items) + "</nav>")


def byline_html(D):
    return (f'<p class="ap-byline">By <a href="/about">Sigurd Lindquist</a>, a Certified Crop Adviser (CCA). '
            f'{asof(D["updated"])}</p>')


def deadlines_box():
    sc = signup_copy(build_date())
    return ('<div class="ap-dead"><ul>' + "".join(f'<li><span class="d">{t}</span><span>{x}</span></li>' for t, x in sc["items"]) + "</ul></div>")


def picker_html(D, st="", fips=""):
    """State and county pickers that go to the county page, and Use my location
    (components/arc-plc-simple.js). The page carries the state list and its own county; the rest of
    the state's counties come from data/arc-plc/nav.json once the page has loaded (a full list in
    every one of 2,800 pages was 24 MB of the site)."""
    sts = sorted(D["states"].items(), key=lambda kv: kv[1]["n"])
    so = "".join(f'<option value="{k}"{" selected" if k == st else ""}>{esc(S["n"])}</option>' for k, S in sts if not st or k == st)
    sfill = " data-fill" if st else ""
    if st and st in D["states"]:
        S = D["states"][st]
        co = '<option value="">Pick a county</option>' + "".join(
            f'<option value="{c["f"]}" data-s="{c["s"]}" selected>{esc(c["n"])}</option>' for c in S["c"] if c["f"] == fips)
        dis = " data-fill"
        nos = f'<noscript><p class="ap-small"><a href="/arc-plc/{S["slug"]}">Every {esc(S["n"])} county</a></p></noscript>'
    else:
        co, dis, nos = '<option value="">Pick a state first</option>', " disabled", ""
    return f"""<div class="ap-find" data-arcfind data-nav="/{OUT_STATE_JSON}/nav.json" data-geo="/{OUT_STATE_JSON}/geo/">
    <div class="ap-find-g">
      <div class="ap-f"><label for="af-st">State</label><select id="af-st"{sfill}><option value="">Pick a state</option>{so}</select></div>
      <div class="ap-f"><label for="af-co">County</label><select id="af-co"{dis}>{co}</select></div>
      <button type="button" class="btn btn-secondary ap-loc" id="af-loc">{PIN}<span>Use my location</span></button>
    </div>
    <p class="ap-find-msg" id="af-msg" role="status" aria-live="polite"></p>
    {nos}
  </div>"""


def card_entries(c, k, D):
    """The practice a crop's card answers for (non-irrigated or all practices; irrigated when the
    county average PLC yield is above the non-irrigated benchmark, as the calculator does), and the other."""
    es = c["k"].get(k) or []
    idx = [i for i in default_entries(es) if es[i]["by"] and not es[i].get("sub")] or [i for i in default_entries(es) if es[i]["by"]]
    if not idx:
        return None, None
    pyv = D["_plc_yields"].get(c["f"], {}).get(k)
    mi = idx[0]
    irr = [i for i in idx if es[i]["d"] == "irr"]
    non = [i for i in idx if es[i]["d"] != "irr"]
    if irr and non and pyv and pyv > es[non[0]]["by"]:
        mi = irr[0]
    oi = next((i for i in idx if es[i]["d"] != es[mi]["d"]), None)
    return mi, oi


def crop_split(c, D):
    """(main crops, smaller crops, crops with a benchmark but no enrolled base): by FSA enrolled base acres in the county."""
    base = (D.get("_plc_base") or {}).get(c["f"], {})
    ks = [k for k in ALL if k in D["years"]["2026"]["crops"] and any(e["by"] for e in c["k"].get(k, []))]
    with_b = sorted([k for k in ks if (base.get(k) or 0) > 0], key=lambda k: (-base[k], ALL.index(k)))
    tot = sum(base[k] for k in with_b)
    main = [k for k in with_b if base[k] >= 0.05 * tot][:4] or with_b[:1]
    return main, [k for k in with_b if k not in main], [k for k in ks if k not in with_b]


def level_tag(c, k, i, e, LY, OY, nm):
    """data-* for one level block: what the test re-derives with components/arc-plc.js."""
    parts = []
    for y in (LY, OY):
        v = c["levels"][(k, i, y)][nm]
        parts.append(f'{y}:{v["verdict"]}:{cents(v.get("plc") or 0)}:{cents(v.get("arc") or 0)}')
    return f' data-k="{k}" data-e="{e["d"]}{":" + esc(e["sub"]) if e.get("sub") else ""}" data-py="{c["levels"][(k, i, LY)][nm]["py"]}" data-t="{"|".join(parts)}"'


def crop_card(c, k, D, LY, OY):
    cd = D["years"][LY]["crops"].get(k) or D["years"]["2026"]["crops"][k]
    lab = cd["label"]
    lc = CROP_LC[k]
    mi, oi = card_entries(c, k, D)
    es = c["k"][k]
    e = es[mi]
    head = f'<h3>{esc(lab)}{" &middot; " + DLABEL[e["d"]].lower() if es and any(x["d"] != e["d"] for x in es if x["by"]) else ""}</h3>'
    if (k, mi, LY) not in c["levels"]:
        why = ("FSA has no county average PLC yield for " + esc(lc) + " here, so there is no typical farm to figure. Run your own numbers below.")
        return (f'<article class="ap-crop" id="crop-{k}">{head}<div class="ap-ans"><span class="ap-big">No answer yet</span>{tier_html({"verdict": "withheld"})}</div>'
                f'<p>{why}</p></article>')
    blocks = []
    for nm, f_ in LEVELS:
        v = c["levels"][(k, mi, LY)][nm]
        v7 = c["levels"][(k, mi, OY)][nm]
        g = gap_line(v, lc, LY, cd)
        if v["verdict"] == "mismatch" and c["levels"][(k, mi, LY)]["same"]["verdict"] != "mismatch":
            g = ("A PLC yield that high is above the county benchmark yield, so this page cannot figure it. Ask your FSA office.")
        oth = ""
        if oi is not None and (k, oi, LY) in c["levels"]:
            vo = c["levels"][(k, oi, LY)][nm]
            note = " Figured with the county&rsquo;s all-practice PLC yield; an irrigated farm&rsquo;s own is often higher." if es[oi]["d"] == "irr" else ""
            oth = f'<p class="ap-oth"><b>{DLABEL[es[oi]["d"]]} ground:</b> {low1(short_word(vo)) if vo["verdict"] not in ("mismatch",) else "ask your FSA office"}.{note}</p>'
        if e["d"] == "irr":
            oth = '<p class="ap-oth">Figured with the county&rsquo;s all-practice PLC yield; an irrigated farm&rsquo;s own is often higher.</p>' + oth
        erp = "" if v["verdict"] == "mismatch" else f'<p>{erp_line(cd, lc, LY)}</p>'
        blocks.append(f'<div class="ap-lv"{level_tag(c, k, mi, e, LY, OY, nm)} data-lv="{nm}"{"" if nm == "same" else " hidden"}>'
                      f'<div class="ap-ans"><span class="ap-big">{big_word(v)}</span>{tier_html(v)}</div>'
                      f'<p>{g}</p>{erp}{oth}<p class="ap-y2">{year_line(v7, OY)}</p></div>')
    return f'<article class="ap-crop" id="crop-{k}">{head}{"".join(blocks)}</article>'


def sheet_shell(D):
    """/arc-plc/sheet: the one counter sheet page (noindex). One letter page, black and white, for the
    FSA office. components/arc-plc-sheet.js reads ?c=<state>/<county> and fills it from
    data/arc-plc/math/<ST>.json, QR code included (components/qr.js)."""
    cfg = {"states": {S["slug"]: st for st, S in sorted(D["states"].items())}, "math": f"/{OUT_MATH}/",
           "icons": ICON_SVG, "key": key_html().replace('class="ap-key"', 'class="ap-key sh-key"').replace("<svg class=\"ap-ic\" viewBox=\"0 0 16 16\" aria-hidden=\"true\"><use href=\"#api-", "@@")}
    for k, v in ICON_SVG.items():
        cfg["key"] = cfg["key"].replace(f'@@{k}"/></svg>', v)
    js = json.dumps(cfg, separators=(",", ":"), ensure_ascii=False).replace("</", "<\\/")
    return f"""<!DOCTYPE html>
<html lang="en" data-theme="light">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="robots" content="noindex,follow">
<link rel="canonical" href="{SITE}/arc-plc">
<title>ARC or PLC counter sheet | AGSIST</title>
<link rel="icon" type="image/x-icon" href="/img/favicon.ico">
<link rel="stylesheet" href="/components/styles.css?v={STYLES_V}">
<link rel="stylesheet" href="/components/states.css?v=2">
<link rel="stylesheet" href="/components/arc-plc-simple.css?v={CALC_V}">
</head>
<body class="ap-sheet-body">
<p class="ap-sheet-bar"><a id="sh-back" href="/arc-plc">Find your county</a> <button type="button" class="btn btn-primary" onclick="window.print()">Print this sheet</button></p>
<main class="ap-sheet sh-msg" id="main"><noscript><p>This sheet is built in the browser. Turn on JavaScript, or open your county&rsquo;s page from <a href="/arc-plc">ARC or PLC</a>.</p></noscript></main>
<script type="application/json" id="sh-data">{js}</script>
<script src="/components/states.js?v=1"></script>
<script src="/components/arc-plc-math.js?v={CALC_V}"></script>
<script src="/components/qr.js?v={CALC_V}"></script>
<script src="/components/arc-plc-sheet.js?v={CALC_V}"></script>
</body>
</html>
"""


def sheet_href(S, c):
    return f"/arc-plc/sheet?c={S['slug']}/{c['s']}"


def sheet_stub(S, c):
    """An old /arc-plc/<state>/<county>-sheet URL (live 2026-10-09 to 10-10): a few hundred bytes that forward to the shared sheet."""
    to = sheet_href(S, c)
    return (f'<!DOCTYPE html>\n<html lang="en"><head><meta charset="utf-8"><meta name="robots" content="noindex,follow">'
            f'<link rel="canonical" href="{SITE}/arc-plc/{S["slug"]}/{c["s"]}"><meta http-equiv="refresh" content="0; url={to}">'
            f'<title>{esc(c["n"])} counter sheet</title><script>location.replace("{to}"+location.hash)</script></head>'
            f'<body><p><a href="{to}">{esc(c["n"])}, {esc(S["n"])}: counter sheet for the FSA office</a></p></body></html>\n')


OUT_MATH = f"{OUT_STATE_JSON}/math"   # data/arc-plc/math/<ST>.json: "Show the math" and the counter sheet, loaded on demand


def _cents2(x):
    """An amount as whole cents the way usd() prints it (half up after rnd), 0 when usd_pay() prints $0."""
    return 0 if cents(x) == 0 else int(round(rnd(x, 2) * 100))


def vpack(v):
    """One typical-farm verdict, compact, for components/arc-plc-math.js:
    [verdict] or [verdict, n] (withheld) or
    [verdict, PLC cents, ARC-CO cents, PLC wins, ARC-CO wins, n, |gap| cents, years ARC-CO paid, flags 1 borrowed 2 futures start
    disagrees 4 gap under a dollar (cents()), |gap| in whole dollars (whole(), rounded once from the unrounded gap)]."""
    if not v:
        return None
    w_ = v["verdict"]
    if w_ in ("pending", "noprice", "mismatch"):
        return [w_]
    if w_ == "withheld":
        return [w_, v.get("n", 0)]
    return [w_, _cents2(v["plc"]), _cents2(v["arc"]), v["plc_wins"], v["arc_wins"], v["n"], int(round(rnd(abs(v["diff"]), 2) * 100)),
            sum(1 for r in v["rows"] if r["arc"] > 0), (1 if v.get("borrowed") else 0) | (2 if v.get("alt_down") else 0) | (4 if cents(abs(v["diff"])) < 100 else 0),
            whole(v["diff"])]


def _n(v):
    """A JSON number without a trailing .0 (2,037.0 prints the same from 2037)."""
    return int(v) if isinstance(v, float) and v.is_integer() else v


def math_globals(D, ks):
    """What every county's math shares: crop program numbers and the sentences that name no county."""
    LY = lead_year()
    OY = "2027" if LY == "2026" else "2026"
    y6, y7 = D["years"]["2026"]["crops"], D["years"]["2027"]["crops"]
    crops = {}
    for k in ks:
        cd = y6[k]
        sc_ = cd.get("scen") or {}
        x = {"lab": cd["label"], "lc": CROP_LC[k], "u": cd.get("unit", "bu"), "dp": cd.get("dp", 2), "sc": XC.scale(cd),
             "y": {y: {"erp": D["years"][y]["crops"][k]["erp"]["erp"], "loan": D["years"][y]["crops"][k]["loan"],
                       "bp": D["years"][y]["crops"][k]["bp"]["value"], "cen": (D["years"][y]["crops"][k].get("scen") or {}).get("center")}
                   for y in ("2026", "2027") if k in D["years"][y]["crops"]},
             "src": (sc_.get("usda_src") or "USDA") if sc_.get("center") else None,
             "bor": bool(sc_.get("borrowed")),
             "y7": "pend" if k not in y7 else ("nop" if (y7[k].get("scen") or {}).get("kind") == "noprice" else "ok"),
             "why": sc_.get("why") or "",
             "hub": (f'{XC.XCROPS[k].get("note", "")} <a href="{XC.hub_path(k)}">All {esc(CROP_LC[k])} counties</a>.' if k in XC.XCROPS else "")}
        if k in ("corn", "soybeans") and sc_.get("alt"):
            x["alt"] = (f"{CROP_LC[k].capitalize()}: USDA {usd(sc_['center'])}; futures-based {usd(sc_['alt'])} ({esc(sc_.get('alt_contract') or '')} October average so far, "
                        f"{usd(sc_['alt_fut'])} over {sc_.get('alt_days')} of {sc_.get('alt_of')} trading days, &times; {sc_['oct_mean']:.3f}, the {min(sc_['ratios'])} to {max(sc_['ratios'])} average of final price / October futures).")
        if k in ("corn", "soybeans") and (y7.get(k) or {}).get("scen"):
            x["c7"] = f"{CROP_LC[k]} {usd(y7[k]['scen']['center'])} ({y7[k]['scen']['center_label']})"
        crops[k] = x
    fsa = D["fsa"] or {}
    wz = (y6["corn"].get("scen") or {})
    snap_d = wz.get("alt_date")
    fut7 = (y7["corn"].get("scen") or {}).get("fut_date")
    src = [f"FSA 2026 ARC-CO benchmark yields (county file of {nice_date(fsa['as_of'])})" if fsa.get("as_of") else "",
           f"FSA county average PLC yields, program year {(D['plc_county'] or {}).get('py')}",
           f"FSA price tables of {nice_date(D['_xtab_date'])}" if D.get("_xtab_date") else "",
           wz.get("usda_src", ""),
           f"CME futures: Dec 2026 corn and Nov 2026 soybean October average through {nice_date(snap_d)}" if snap_d else "",
           f"Dec 2027 corn and Nov 2027 soybean close of {fut7}" if fut7 else "",
           "USDA NASS season-average prices for closed marketing years",
           "FSA ARC-CO county files, program years " + ", ".join(str(h["py"]) for h in (fsa.get("history") or [])) + " (county yields and payments)"]
    sc = signup_copy(build_date())
    return {"ly": LY, "oy": OY, "plc_py": (D["plc_county"] or {}).get("py"), "min_scen": MIN_SCEN,
            "gap": usd(MIN_GAP), "lean": usd(LEAN_GAP), "wins": MIN_WINS,
            "srcs": "; ".join(x for x in src if x), "asof": asof(D["updated"]), "ask": ASK,
            "wz": next((cd["scen"]["usda_src"] for cd in y6.values() if (cd.get("scen") or {}).get("usda_src", "").startswith("USDA WASDE")), "USDA"),
            "dl": [x for _t, x in sc["items"]], "upd": nice_date(D["updated"][:10]), "c": crops}


def math_json(D):
    """data/arc-plc/math/<ST>.json for every state: what components/arc-plc-math.js needs to
    build a county's "Show the math" and /arc-plc/sheet builds the counter sheet from. Numbers
    the build figures (verdicts, payments) are carried as figured; the page only lays them out."""
    if not D["fsa"]:
        return {}
    LY = lead_year()
    OY = "2027" if LY == "2026" else "2026"
    y6 = D["years"]["2026"]["crops"]
    win = list((D["fsa"] or {}).get("window") or [])
    out = {}
    for st, S in sorted(D["states"].items()):
        recs, ks_all = {}, []
        for c in S["c"]:
            main, small, nob = crop_split(c, D)
            order = crop_order(c, D)
            plcy = D["_plc_yields"].get(c["f"], {})
            es_out, v_out, al = {}, {}, {}
            for k, es in c["k"].items():
                rows = []
                for e in es:
                    r = {"d": e["d"], "by": _n(e["by"])}
                    if e.get("dsrc") != "county":
                        r["ds"] = e.get("dsrc") or "?"
                    if e.get("sub"):
                        r["sub"] = e["sub"]
                    vals = [v for v in (e.get("yrs") or {}).values() if v is not None]
                    if e["by"] is not None and len(vals) == 5:
                        if [int(y) for y in e["yrs"]] != win:
                            raise SystemExit(f"[arc-plc] {c['n']} {k}: trend-yield years {list(e['yrs'])} are not FSA's window {win}")
                        # the five trend-adjusted yields as %g prints them, and which two the Olympic average drops
                        # (the first value equal to the low, then to the high: the same marks the old pages carried)
                        _avg, lo, hi = olympic(vals)
                        x, lo_d, hi_d = 0, False, False
                        for j, v in enumerate(e["yrs"].values()):
                            if not lo_d and v == lo:
                                x |= 1 << j; lo_d = True
                            elif not hi_d and v == hi:
                                x |= 1 << j; hi_d = True
                        r["ys"], r["x"] = [_n(float(f"{v:g}")) for v in e["yrs"].values()], x
                    hist = [h for h in e.get("hist", []) if h["py"] != 2026]
                    if hist:
                        # program year, benchmark, price, county yield, payment rate, for each year in turn
                        r["h"] = [_n(x_) for h in hist for x_ in (h["py"], h["by"], h["bp"], h["ay"], h["pay"])]
                    rows.append(r)
                es_out[k] = rows
                v_out[k] = [({LY: vpack(c["verdicts"].get((k, i, LY))), OY: vpack(c["verdicts"].get((k, i, OY)))}
                             if (c["verdicts"].get((k, i, LY)) or c["verdicts"].get((k, i, OY))) else None) for i in range(len(es))]
            for k in ("corn", "soybeans"):
                if k in main + small and (y6[k].get("scen") or {}).get("alt"):
                    mi, _ = card_entries(c, k, D)
                    v = c["levels"].get((k, mi, "2026"), {}).get("same", {})
                    a = v.get("alt")
                    al[k] = [v.get("verdict"), a["verdict"] if a else None]
            rec = {"s": c["s"], "n": c["n"], "o": order, "m": main, "sm": small, "nb": nob,
                   "py": {k: _n(plcy[k]) for k in c["k"] if plcy.get(k)}, "e": es_out, "v": v_out,
                   "ci": {k: card_entries(c, k, D)[0] for k in main + small}}
            if al:
                rec["al"] = al
            recs[c["f"]] = rec
            for k in order + main + small + nob:
                if k not in ks_all:
                    ks_all.append(k)
        g = math_globals(D, [k for k in ALL if k in ks_all])
        g.update({"st": st, "n": S["n"], "slug": S["slug"], "win": win})
        out[f"{OUT_MATH}/{st}.json"] = json.dumps({"updated": D["updated"], "g": g, "c": recs}, separators=(",", ":"), ensure_ascii=False) + "\n"
    return out


MATH_BOX = ('<div class="ap-mathb"><noscript><p class="ap-small">The working is built in the browser from '
            '<a href="/' + OUT_MATH + '/{st}.json">this state&rsquo;s data file</a>. <a href="/arc-plc#method">How this is figured</a>.</p></noscript></div>')


def key_numbers_html(c, D, cname):
    """Key numbers, in the page for readers and crawlers: per crop and practice, FSA's 2026 benchmark yield, the
    price PLC pays under (the effective reference price) for both years, FSA's county average PLC yield, and the
    ARC-CO payment rate FSA paid in each settled year. Every figure is FSA's or the build's; a missing one says why."""
    y6, y7 = D["years"]["2026"]["crops"], D["years"]["2027"]["crops"]
    plcy = D["_plc_yields"].get(c["f"], {})
    hy = [h["py"] for h in ((D["fsa"] or {}).get("history") or []) if h["py"] != 2026]
    rows = []
    for k in crop_order(c, D):
        cd = y6[k]
        es = c["k"].get(k) or []
        for e in es:
            lab = cd["label"] + ("" if e["d"] == "all" and not e.get("sub") else ", " + ent_label(e).lower())
            hist = {h["py"]: h for h in e.get("hist", [])}
            pays = "".join(f'<td>{("none" if hist[y]["pay"] == 0 else usd(hist[y]["pay"])) if y in hist and hist[y]["pay"] is not None else ("not yet" if y in hist else "not in file")}</td>' for y in hy)
            e7 = f"{XC.pf(y7[k], y7[k]['erp']['erp'])}{XC.per(cd)}" if k in y7 else "not set yet"
            rows.append(f'<tr><th scope="row">{esc(lab)}</th><td>{XC.yf(cd, e["by"]) if e["by"] else "not loaded"}</td>'
                        f'<td>{XC.pf(cd, cd["erp"]["erp"])}{XC.per(cd)}</td><td>{e7}</td>'
                        f'<td>{XC.yf(cd, plcy[k], 1) if plcy.get(k) else "none from FSA"}</td>{pays}</tr>')
    if not rows:
        return ""
    return (f'<h2 id="key-numbers">Key numbers</h2><p class="ap-small">FSA&rsquo;s figures for '
            f'{esc(cname)}. PLC pays when the season-average price ends under the effective reference price. ARC-CO paid is FSA&rsquo;s rate per '
            f'acre for the county, before the 85% factor.</p><div class="ap-scroll"><table class="tbl ap-t ap-keyn"><thead><tr><th scope="col" rowspan="2">Crop</th><th scope="col" rowspan="2">Benchmark<br>yield 2026</th>'
            f'<th scope="colgroup" colspan="2">PLC pays under</th><th scope="col" rowspan="2">County avg<br>PLC yield</th>'
            + (f'<th scope="colgroup" colspan="{len(hy)}">ARC-CO paid per acre</th>' if hy else "") + '</tr><tr><th scope="col">2026</th><th scope="col">2027</th>'
            + "".join(f'<th scope="col">{y}</th>' for y in hy) + f'</tr></thead><tbody>{"".join(rows)}</tbody></table></div>')


def county_title(cname, st, lab, LY, by_txt):
    """'{County}, {ST} ARC or PLC {year}: {Crop} Benchmark {yield}', shortened until it fits 60 characters.
    Always carries 'ARC or PLC'; the verdict goes in the description, not here."""
    short = re.sub(r" County$", " Co.", cname)
    for t in (f"{cname}, {st} ARC or PLC {LY}: {lab} Benchmark {by_txt}",
              f"{cname}, {st} ARC or PLC {LY}: {lab} Benchmark",
              f"{short}, {st} ARC or PLC {LY}: {lab} Benchmark {by_txt}",
              f"{short}, {st} ARC or PLC {LY}: {lab} Benchmark",
              f"{short}, {st} ARC or PLC {LY}: {lab}",
              f"{cname}, {st} ARC or PLC {LY} and {int(LY) + 1 if LY == '2026' else 2026}",
              f"{short}, {st} ARC or PLC {LY}"):
        if len(t) <= 60:
            return t
    return f"{short} {st} ARC or PLC {LY}"[:60]


def county_page(st, S, c, D, ch):
    hdr, ftr, fonts = ch
    name, sl = S["n"], S["slug"]
    cname, cslug = c["n"], c["s"]
    path = f"/arc-plc/{sl}/{cslug}"
    y6 = D["years"]["2026"]["crops"]
    fsa = D["fsa"]
    official = bool(fsa)
    plcy = D["_plc_yields"].get(c["f"], {})
    plc_py = (D["plc_county"] or {}).get("py")
    LY = lead_year()
    order = crop_order(c, D)
    lead = order[0] if order else None
    head_v = typical_html(c, lead, D, cname)[1] if official and lead else None
    if official and lead:
        le = next((c["k"][lead][i] for i in default_entries(c["k"][lead]) if c["k"][lead][i]["by"]), None)
        cdl = y6[lead]
        lab = XC.XCROPS[lead].get("short", cdl["label"]) if lead in XC.XCROPS else cdl["label"]
        if le and le["d"] == "irr":
            lab = "Irrigated " + lab.lower()
        by_txt = (f"{le['by']:.1f} bu" if cdl["unit"] == "bu" else f"{le['by']:,.0f} lb") if le else ""
        title = county_title(cname, st, lab, LY, by_txt)
        bm = f"FSA {LY if LY == '2026' else '2026'} ARC-CO benchmark {XC.yf(cdl, le['by'], 1)}" if le else "FSA ARC-CO benchmarks"
        if head_v:
            vv = head_v[0]
            desc = (f"{cname}, {name} {lab.lower()}: {bm}. Typical farm {head_v[2]}: {VWORD[vv['verdict']]}, est. {usd_pay(vv['plc'])} PLC vs "
                    f"{usd_pay(vv['arc'])} ARC-CO per base acre. Run your numbers.")
            if len(desc) > 155:
                desc = (f"{cname}, {st} {lab.lower()}: {bm}. {head_v[2]}: {VWORD[vv['verdict']]}, est. {usd_pay(vv['plc'])} PLC vs "
                        f"{usd_pay(vv['arc'])} ARC-CO per base acre.")
            if len(desc) > 155:
                desc = f"{cname}, {st}: {VWORD[vv['verdict']]} for {lab.lower()} in {head_v[2]}, est. {usd_pay(vv['plc'])} PLC vs {usd_pay(vv['arc'])} ARC-CO per base acre."
        else:
            bits = [f"{CROP_LC[k]}{'' if e['d'] == 'all' else ' ' + DLABEL[e['d']].lower()} {XC.yf(y6[k], e['by'], 1)}" for k in order for e in c["k"].get(k, []) if e["by"]]
            desc = f"{cname}, {name}: FSA official 2026 ARC-CO benchmark yields: {', '.join(bits)}. Revenue, guarantee and PLC vs ARC break-even."
            if len(desc) > 155:
                desc = f"{cname}, {st}: FSA official 2026 ARC-CO benchmark yields by crop and practice, revenue, guarantee and PLC vs ARC break-even."
            if len(desc) > 155:
                desc = f"{cname}, {st}: FSA 2026 ARC-CO benchmark yields by crop, revenue and the PLC vs ARC break-even."
    elif official:
        title = county_title(cname, st, "ARC-CO", LY, "")
        desc = f"{cname}, {st}: FSA official 2026 ARC-CO benchmark yields by crop and practice, revenue, guarantee and PLC vs ARC break-even."
    else:
        title = next((t for t in (f"{cname}, {st} ARC or PLC 2026 and 2027 | AGSIST", f"{cname}, {st} ARC or PLC | AGSIST")
                      if len(t) <= 60), f"{cname}, {st} ARC or PLC")
        desc = f"{cname}, {name}: FSA county average PLC yields and ARC-CO history, and how to get the official 2026 benchmark."
    blocks, faqs, ds_vars = [], [], []
    first = lead or next(k for k in ALL if k in c["k"])
    for k in (order if official else [k for k in ALL if k in c["k"]]):
        es = c["k"].get(k)
        if not es:
            continue
        cd = y6[k]
        u = XC.UNIT[cd["unit"]]["short"]
        th, hv = typical_html(c, k, D, cname) if official else ("", None)
        hub = (f' <a href="{XC.hub_path(k)}">All {esc(CROP_LC[k])} counties</a>.' if k in XC.XCROPS else "")
        blocks.append(f'<h2 id="{k}">{esc(cd["label"])} {hv[2] if hv else LY} in {esc(cname)}: {VWORD[hv[0]["verdict"]] if hv else "ARC-CO benchmark and break-even"}</h2>')
        if hub:
            blocks.append(f'<p class="ap-small">{XC.XCROPS[k].get("note", "")}{hub}</p>')
        if th:
            blocks.append(th)
            if hv:
                faqs.append((f"Should I pick ARC or PLC for {CROP_LC[k]} in {cname} for {hv[2]}?",
                             re.sub(r"<[^>]+>", "", html.unescape(f"For a typical farm ({ent_label(hv[1]).lower()}): {verdict_line(hv[0])} "
                                                                  f"{why_text(hv[0], D['years'][hv[2]]['crops'][k], hv[1]['by'], plcy[k])} Run your own PLC yield and base acres; the answer can change."))))
        if plcy.get(k) and plc_py:
            blocks.append(f'<p>Average PLC yield on enrolled {CROP_LC[k]} base in this county: <b>{XC.yf(cd, plcy[k], 1)}</b> (FSA, program year {plc_py}). '
                          f'Your farm&rsquo;s own is on the FSA-156EZ.</p>')
        for e in es:
            if e["by"] is None:
                blocks.append(f'<h3>{esc(ent_label(e))}</h3><p class="ap-small">Official 2026 benchmark: not loaded here yet. Ask the county office.</p>{hist_table(e, cd)}')
                continue
            if official:
                # the per-practice working (benchmark, Olympic average, break-even, history) is built on open
                # by components/arc-plc-math.js from data/arc-plc/math/<ST>.json
                ds_vars.append(f"{CROP_LC[k]} {ent_label(e).lower()} ARC-CO benchmark yield ({u}/acre)")
                continue
            m = county_metrics(e["by"], cd)
            yrs = e["yrs"]
            vals = [v for v in yrs.values() if v is not None]
            ys = ""
            if len(vals) == 5:
                _avg, lo, hi = olympic(vals)
                lo_d = hi_d = False
                parts = []
                for y, v in yrs.items():
                    if not lo_d and v == lo:
                        parts.append(f'<s>{y}: {v:g}</s><span class="sr-only"> (dropped)</span>'); lo_d = True
                    elif not hi_d and v == hi:
                        parts.append(f'<s>{y}: {v:g}</s><span class="sr-only"> (dropped)</span>'); hi_d = True
                    else:
                        parts.append(f"{y}: {v:g}")
                ys = (f'<p class="ap-math">FSA trend-adjusted county yields (county yield or 80% of T-yield): {" &middot; ".join(parts)} '
                      f'(struck: high and low). Middle three average: <b>{XC.yf(cd, e["by"])}</b>.</p>')
            blocks.append(f"""
  <h3>{esc(ent_label(e))}: official 2026 ARC-CO benchmark</h3>
  <div class="ap-kv">
    <div><span>Benchmark yield (FSA)</span><b>{XC.yf(cd, e['by'])}</b></div>
    <div><span>Benchmark price 2026</span><b>{XC.pf(cd, cd['bp']['value'])}</b></div>
    <div><span>Benchmark revenue</span><b>{usd(m['br'])}/ac</b></div>
    <div><span>Guarantee (90%)</span><b>{usd(m['g'])}/ac</b></div>
    <div><span>Max payment (12%)</span><b>{usd(m['max'])}/ac</b></div>
    <div><span>Max per base acre (&times;85%)</span><b>{usd(m['max_base'])}</b></div>
  </div>
  {ys}
  <p>At a season-average price equal to the 2026 effective reference price ({XC.pf(cd, cd['erp']['erp'])}{XC.per(cd)}), PLC pays nothing and ARC-CO pays when the
  county yield comes in below <b>{XC.yf(cd, m['trig_y'], 1)}</b>, reaching its cap below <b>{XC.yf(cd, m['cap_y'], 1)}</b>.</p>
  <h3>PLC vs ARC-CO break-even at a {XC.yf(cd, e['by'])} county yield, 2026 prices</h3>
  {crossover_table(cd, e['by'])}
  {hist_table(e, cd)}""")
            ds_vars.append(f"{CROP_LC[k]} {ent_label(e).lower()} ARC-CO benchmark yield ({u}/acre)")
        if official and es[0]["by"]:
            faqs.append((f"What is the ARC-CO benchmark yield for {CROP_LC[k]} in {cname}?",
                         f"FSA's official 2026 benchmark yield for {CROP_LC[k]} in {cname}, {name} is "
                         + "; ".join(f"{XC.yf(cd, e['by'])} ({ent_label(e).lower()})" for e in es if e["by"])
                         + f". The 2026 benchmark price is {XC.pf(cd, cd['bp']['value'])} per {XC.word(cd)}. FSA has not posted the 2027 benchmark."))
        elif plcy.get(k) and plc_py:
            faqs.append((f"What is the average PLC yield for {CROP_LC[k]} in {cname}?",
                         f"FSA lists {plcy[k]:,.1f} {XC.word(cd)}s per acre as the average PLC yield on enrolled {CROP_LC[k]} base in {cname}, {name} for program year "
                         f"{plc_py}. Each farm has its own PLC yield, shown on its FSA-156EZ."))
        if not official:
            blocks.append(f"<p>Break-even by PLC yield as a share of the official benchmark (2026 prices):</p>{crossover_table(cd)}")
    faq_html = "".join(f"<details><summary>{esc(q)}</summary><p>{esc(a)}</p></details>" for q, a in faqs)
    OY = "2027" if LY == "2026" else "2026"
    jsonld = [{"@context": "https://schema.org", "@type": "BreadcrumbList", "itemListElement": [
        {"@type": "ListItem", "position": 1, "name": "Home", "item": f"{SITE}/"},
        {"@type": "ListItem", "position": 2, "name": "ARC or PLC", "item": f"{SITE}/arc-plc"},
        {"@type": "ListItem", "position": 3, "name": name, "item": f"{SITE}/arc-plc/{sl}"},
        {"@type": "ListItem", "position": 4, "name": cname, "item": f"{SITE}{path}"}]},
        {"@context": "https://schema.org", "@type": "WebPage", "@id": f"{SITE}{path}", "name": title, "url": f"{SITE}{path}",
         "description": desc, "author": AUTHOR, "publisher": {"@type": "Organization", "name": "AGSIST", "url": SITE},
         "dateModified": D["updated"][:10], "inLanguage": "en-US"}]
    if faqs:
        jsonld.insert(0, {"@context": "https://schema.org", "@type": "FAQPage",
                          "mainEntity": [{"@type": "Question", "name": q, "acceptedAnswer": {"@type": "Answer", "text": a}} for q, a in faqs]})
    if official:
        jsonld.insert(0, {"@context": "https://schema.org", "@type": "Dataset", "@id": f"{SITE}{path}#dataset",
                          "name": f"{cname}, {name} 2026 ARC-CO benchmark yields (FSA official)",
                          "description": (f"FSA's official program year 2026 ARC-CO benchmark yields for {cname}, {name} by crop and practice, with the "
                                          f"five trend-adjusted yields behind each, benchmark revenue, guarantee, maximum payment and FSA's history."),
                          "url": f"{SITE}{path}", "isAccessibleForFree": True, "license": "https://creativecommons.org/licenses/by/4.0/",
                          "creator": {"@type": "Organization", "name": "AGSIST", "url": SITE}, "author": AUTHOR,
                          "isBasedOn": FSA_DATA, "dateModified": D["updated"][:10],
                          "distribution": {"@type": "DataDownload", "encodingFormat": "application/json", "contentUrl": f"{SITE}/{OUT_STATE_JSON}/{st}.json"},
                          "spatialCoverage": {"@type": "Place", "name": f"{cname}, {name}, United States"}, "variableMeasured": ds_vars})
    atlas = os.path.exists(os.path.join(ROOT, "farmland-atlas", sl, f"{cslug}.html"))
    rel = [f'<a href="/arc-plc/{sl}">All {esc(name)} counties</a>', '<a href="/arc-plc">ARC or PLC calculator and method</a>']
    if atlas:
        rel.append(f'<a href="/farmland-atlas/{sl}/{cslug}">{esc(cname)} in the Farmland Atlas</a>')
    missing = [CROP_LC[k] for k in CROPS if k not in c["k"]]
    miss_txt = f'<p class="ap-small">Nothing from FSA here for {" or ".join(missing)}. {ASK}</p>' if missing else ""
    share = f"{SITE}{path}"
    sc_ = signup_copy(build_date())
    if official:
        main_k, small_k, _nob = crop_split(c, D)
        cards = "".join(crop_card(c, k, D, LY, OY) for k in main_k)
        smalls = "".join(crop_card(c, k, D, LY, OY) for k in small_k)
        base = (D.get("_plc_base") or {}).get(c["f"], {})
        cap = int(math.ceil(max(base[k] for k in small_k) / 100.0) * 100) if small_k else 0
        crops_html = (f'<h2 id="crops">{"Your main crops" if small_k else "Your crops"}</h2>{key_html()}<div class="ap-cards2">{cards}</div>'
                      + (f'<h2 id="smaller">Smaller crops here</h2><p class="ap-small">Each has under {cap:,} enrolled base acres in the whole county (FSA, program year {plc_py}).</p>'
                         f'<div class="ap-cards2 ap-cards2-s">{smalls}</div>' if small_k else "")) if (main_k or small_k) else (
            '<p class="ap-need">FSA lists no enrolled base acres here for a crop with a benchmark, so there is no typical farm to figure. Run your own numbers below.</p>')
        tweak = f"""<h2 id="your-plc-yield">Is your farm&rsquo;s PLC yield lower, about the same, or higher than most farms here?</h2>
  <div class="ap-seg" role="group" aria-label="Your farm&rsquo;s PLC yield compared with the county average">
    <button type="button" data-lv="lower" aria-pressed="false">Lower</button><button type="button" data-lv="same" aria-pressed="true">About the same</button><button type="button" data-lv="higher" aria-pressed="false">Higher</button>
  </div>
  <p class="ap-small ap-seg-note" data-same="Your PLC yield is on your farm&rsquo;s FSA-156EZ. Not sure? Leave it on About the same." data-lower="Showing a PLC yield 15% under the county average." data-higher="Showing a PLC yield 15% over the county average.">Your PLC yield is on your farm&rsquo;s FSA-156EZ. Not sure? Leave it on About the same.</p>"""
        math_top = ""   # built on open by components/arc-plc-math.js
    else:
        crops_html, tweak, math_top = pending_note(D), "", ""
    body = f"""
<body>
{hdr}
<main class="ap-wrap ap-simple" id="main">
  {icon_sprite() if official else ""}{bc_html([("Home", "/"), ("ARC or PLC", "/arc-plc"), (name, f"/arc-plc/{sl}"), (cname, None)])}
  <p class="page-kicker">{esc(cname)}, {esc(name)}</p>
  <h1>ARC or PLC: the short answer<span class="sr-only"> for {esc(cname)}, {esc(name)}</span></h1>
  {byline_html(D)}
  <p class="page-lede">What pays more on your base acres for a typical farm here, crop by crop. {sc_['desc']}</p>
  {picker_html(D, st, c['f'])}
  {tweak}
  {crops_html}
  <h2 id="deadlines">Deadlines</h2>
  {deadlines_box()}
  {key_numbers_html(c, D, cname) if official else ''}
  <p class="ap-tools"><a class="btn btn-primary" href="{sheet_href(S, c)}" rel="nofollow">Counter sheet for the FSA office</a>
  <a class="btn btn-secondary" href="{share}" data-share="{share}" data-title="{esc(cname)} ARC or PLC">Share this county</a></p>
  <details class="ap-det ap-mathd" id="math"{f' data-math="/{OUT_MATH}/{st}.json" data-f="{c["f"]}"' if official else ""}><summary>Show the math</summary>
  {MATH_BOX.format(st=st) if official else math_top + chr(10) + "  " + "".join(blocks) + chr(10) + "  " + miss_txt}
  <p class="ap-small">Break-even prices are per base acre with both programs on 85% of base and 2026 prices. Your PLC yield is on the FSA-156EZ.</p>
  </details>
  <details class="ap-full" id="run"><summary>Run every number for your farm</summary>
  {form_shell(D, first, st, c['f']) if official else form_html(D, first, st, c['f'])}
  </details>
  {f'<div class="ap-faq ap-noprint"><h2>Questions about {esc(cname)}</h2>{faq_html}</div>' if faqs else ''}
  <p class="ap-disc">An estimate, not a USDA determination. Source: <a href="{FSA_DATA}" rel="noopener">FSA ARC/PLC program data</a>. <a href="{OFFICE}" rel="noopener">Find your county office</a>.</p>
  <p class="ap-small ap-noprint">Related: {' &middot; '.join(rel)}</p>
  <p class="ap-printonly ap-small">Printed from {share}</p>
</main>
<script>document.addEventListener('click',function(e){{var a=e.target.closest&&e.target.closest('[data-share]');if(!a||!navigator.share)return;e.preventDefault();navigator.share({{title:a.getAttribute('data-title'),url:a.getAttribute('data-share')}}).catch(function(){{}});}});</script>
"""
    page = head(title, desc, path, jsonld, official, fonts) + body + tail(ftr, calc=True, math=official)
    return page, {"path": path, "indexable": official, "title": title, "desc": desc, "lead": lead, "st": st}


def embed_page(D):
    return f"""<!DOCTYPE html>
<html lang="en" data-theme="light">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="robots" content="noindex"><!-- embed target; canonical tool is /arc-plc -->
<link rel="canonical" href="{SITE}/arc-plc">
<title>ARC or PLC | AGSIST</title>
<link rel="icon" type="image/x-icon" href="/img/favicon.ico">
<link rel="stylesheet" href="/components/styles.css?v={STYLES_V}">
<link rel="stylesheet" href="/components/asof.css?v={ASOF_V}">
<link rel="stylesheet" href="/components/arc-plc.css?v={CALC_V}">
<script>try{{var m=/[?&]theme=(dark|light)/.exec(location.search);document.documentElement.setAttribute('data-theme',m?m[1]:(matchMedia('(prefers-color-scheme: dark)').matches?'dark':'light'));}}catch(e){{}}</script>
<style>body{{margin:0;background:var(--bg)}}.ap-wrap{{padding:12px 12px 8px}}</style>
</head>
<body>
<main class="ap-wrap ap-embed" id="main">
{form_html(D, embed=True)}
<p class="ap-small">An estimate, not a USDA determination. Prices, FSA county benchmarks and the full math at <a href="{SITE}/arc-plc" target="_blank" rel="noopener">agsist.com/arc-plc</a>.</p>
</main>
<script src="/components/asof.js?v={ASOF_V}" defer></script>
<script src="/components/arc-plc.js?v={CALC_V}" defer></script>
<script>
(function(){{function post(){{try{{parent.postMessage({{agsist:'arc-plc',height:document.body.scrollHeight}},'*');}}catch(e){{}}}}
window.addEventListener('load',post);document.addEventListener('arcplc:render',post);
if(window.ResizeObserver){{try{{new ResizeObserver(post).observe(document.body);}}catch(e){{}}}}}})();
</script>
</body>
</html>
"""


def sitemap(urls):
    body = "".join(f"  <url><loc>{SITE}{u}</loc><lastmod>{lm}</lastmod><changefreq>monthly</changefreq><priority>{p}</priority></url>\n"
                   for u, p, lm in urls)
    return f'<?xml version="1.0" encoding="UTF-8"?>\n<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n{body}</urlset>\n'


LLMS_PREFIX = "- [ARC or PLC for 2026 and 2027](https://agsist.com/arc-plc):"
LLMS_PREFIX2 = "- [ARC-CO benchmark yields by county](https://agsist.com/arc-plc/iowa):"
LLMS_OLD = "- [ARC or PLC for 2027](https://agsist.com/arc-plc):"


def llms_text(D, root="."):
    p = os.path.join(root, "llms.txt")
    if not os.path.exists(p):
        return None
    cur = open(p, encoding="utf-8").read()
    y6, y7 = D["years"]["2026"]["crops"], D["years"]["2027"]["crops"]
    n_c = sum(len(S["c"]) for S in D["states"].values())
    l1 = (f"{LLMS_PREFIX} PLC effective reference prices computed from final USDA NASS season-average prices under the 2025 farm law: 2026 corn "
          f"{usd(y6['corn']['erp']['erp'])}, soybeans {usd(y6['soybeans']['erp']['erp'])}, wheat {usd(y6['wheat']['erp']['erp'])}; 2027 est. corn "
          f"{usd(y7['corn']['erp']['erp'])}, soybeans {usd(y7['soybeans']['erp']['erp'])}, wheat {usd(y7['wheat']['erp']['erp'])}. "
          f"{signup_copy(build_date())['llms']} A calculator compares PLC and ARC-CO per base acre by price and county yield, "
          f"with what changed under the 2025 law, SCO, rented ground and payment limits")
    l2 = (f"{LLMS_PREFIX2} FSA's official 2026 ARC-CO benchmark yields for {n_c:,} counties by crop (corn, soybeans, wheat and every other covered crop FSA lists) and practice "
          f"(irrigated, non-irrigated), with benchmark revenue, guarantee, maximum payment, FSA history and the PLC vs ARC break-even, at "
          f"/arc-plc/<state> and /arc-plc/<state>/<county> (e.g. /arc-plc/iowa/story-county)")
    lines = cur.split("\n")
    out, done = [], False
    for ln in lines:
        if ln.startswith((LLMS_PREFIX, LLMS_PREFIX2, LLMS_OLD)):
            if not done:
                out += [l1] + ([l2] if D["fsa"] else [])
                done = True
            continue
        out.append(ln)
    return "\n".join(out) if done else None


def render_all(D):
    ch = chrome()
    main_j, per = calc_json(D)
    out = {OUT_JSON: json.dumps(main_j, separators=(",", ":"), ensure_ascii=False) + "\n"}
    for code, j in per.items():
        out[f"{OUT_STATE_JSON}/{code}.json"] = json.dumps(j, separators=(",", ":"), ensure_ascii=False) + "\n"
    for rel, obj in geo_json(D).items():
        out[rel] = json.dumps(obj, separators=(",", ":"), ensure_ascii=False) + "\n"
    main, meta = build_main(D, ch)
    out[OUT_MAIN] = main
    out[OUT_EMBED] = embed_page(D)
    lm = D["updated"][:10]
    urls = []   # /arc-plc itself is listed in sitemap.xml only
    for k in D.get("_hubs", []):
        page, m = XC.hub_page(sys.modules[__name__], D, k, ch)
        out[f"{OUT_DIR}/{XC.XCROPS[k]['slug']}.html"] = page
        out[f"{OUT_STATE_JSON}/hub/{k}.json"] = json.dumps(XC.hub_json(D, k), separators=(",", ":"), ensure_ascii=False) + "\n"
        urls.append((m["path"], "0.6", lm))
    nav_states = sorted(D["states"].values(), key=lambda x: x["n"])
    stats = {"states": 0, "states_indexed": 0, "counties": 0, "meta": meta, "examples": {}, "hubs": len(D.get("_hubs", []))}
    for st, S in sorted(D["states"].items()):
        page, m = state_page(st, S, D, ch, nav_states)
        out[f"{OUT_DIR}/{S['slug']}.html"] = page
        stats["states"] += 1
        if m["indexable"]:
            stats["states_indexed"] += 1
            urls.append((m["path"], "0.6", lm))
        stats["examples"].setdefault("state", m)
        for c in S["c"]:
            page, m = county_page(st, S, c, D, ch)
            out[f"{OUT_DIR}/{S['slug']}/{c['s']}.html"] = page
            if D["fsa"]:
                out[f"{OUT_DIR}/{S['slug']}/{c['s']}{SHEET_SUFFIX}.html"] = sheet_stub(S, c)
            stats["counties"] += 1
            if m["indexable"]:
                urls.append((m["path"], "0.5", lm))
            stats["examples"].setdefault("county", m)
    out[OUT_FORM_JS] = form_js(D)
    if D["fsa"]:
        out[OUT_SHEET] = sheet_shell(D)
        out.update(math_json(D))
    if len({u for u, _p, _l in urls}) != len(urls):
        raise SystemExit("[arc-plc] two pages share one URL; refusing to write the sitemap")
    out[OUT_SITEMAP] = sitemap(urls)
    t = XC.llms_merge(llms_text(D), D, (LLMS_PREFIX, LLMS_PREFIX2))
    if t:
        out["llms.txt"] = t
    return out, stats


def existing_outputs(root):
    paths = set()
    for d_, ext in ((OUT_DIR, ".html"), (OUT_STATE_JSON, ".json")):
        for base, _dirs, files in os.walk(os.path.join(root, d_)):
            for f in files:
                if f.endswith(ext):
                    paths.add(os.path.relpath(os.path.join(base, f), root))
    return paths


def budget_problems(out):
    """The size budget for what this build writes under arc-plc/, read from scripts/check_site_budget.py
    (one set of numbers for the build and the site check): the folder's total and the cap on any one page.
    Every file in arc-plc/ is this build's output, so the total is the sum of what it is about to write.
    How it crept up (Oct 2026): 2,772 county pages at 125 KB with the math baked in, plus a sheet each."""
    import check_site_budget as SB
    cap_mb, page_kb = SB.FOLDER_MB.get(OUT_DIR, SB.DEFAULT_FOLDER_MB), SB.SET_PAGE_KB
    sizes = {rel: len(t.encode("utf-8")) for rel, t in out.items() if rel.startswith(OUT_DIR + "/")}
    total = sum(sizes.values())
    bad = []
    if total > cap_mb * 1048576:
        bad.append(f"{OUT_DIR}/ would be {total / 1048576:.1f} MB, over its {cap_mb} MB budget (scripts/check_site_budget.py)")
    fat = sorted(((n, rel) for rel, n in sizes.items() if n > page_kb * 1024), reverse=True)
    for n, rel in fat[:5]:
        bad.append(f"{rel} would be {n / 1024:.0f} KB, over the {page_kb} KB cap for a page in a generated set")
    if len(fat) > 5:
        bad.append(f"...and {len(fat) - 5} more pages over {page_kb} KB")
    return bad, total


def main_build(root=".", check=False):
    inp = load_inputs(root)
    bad = None if check else snapshot_problem(inp["snap"], build_date())
    if bad:
        raise SystemExit(f"[arc-plc] refusing to write: the futures snapshot {bad}. Nothing written.")
    D = compute(inp)
    out, stats = render_all(D)
    over, arc_bytes = budget_problems(out)
    if over:
        for o in over:
            print(f"[arc-plc] OVER BUDGET: {o}", file=sys.stderr)
        raise SystemExit("[arc-plc] refusing to write: the pages are over the site size budget. Move bulk to data loaded on demand. Nothing written.")
    stale = []
    for rel, content in out.items():
        p = os.path.join(root, rel)
        cur = open(p, encoding="utf-8").read() if os.path.exists(p) else None
        if cur != content:
            stale.append(rel)
            if not check:
                os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
                with open(p, "w", encoding="utf-8") as f:
                    f.write(content)
    gone = sorted(existing_outputs(root) - set(out))
    for rel in gone:
        stale.append(rel + " (no longer built)")
        if not check:
            os.remove(os.path.join(root, rel))
    if check:
        if stale:
            print(f"[arc-plc] STALE: {len(stale)} file(s), e.g. {stale[:5]}", file=sys.stderr)
            return 1
        print("[arc-plc] OK: outputs match the inputs.")
        return 0
    f = D["fsa"]
    print(f"[arc-plc] FSA official file: {('PY' + str(f['py']) + ' as of ' + str(f['as_of'])) if f else 'none'}; "
          f"{stats['states']} state pages ({stats['states_indexed']} indexable), {stats['counties']} county pages; "
          f"{len(stale)} file(s) written, {len(gone)} removed; {OUT_DIR}/ is {arc_bytes / 1048576:.1f} MB.")
    for y in D["years"]:
        for k, x in D["years"][y]["crops"].items():
            print(f"[arc-plc] {y} {k}: ERP {x['erp']['erp']} ({x['erp']['binding']}), benchmark price {x['bp']['value']}")
    return 0


# ---------------------------------------------------------------- selftest
def selftest():
    fails = []

    def ck(name, ok, detail=""):
        print(("ok   " if ok else "FAIL ") + name + (f"  {detail}" if detail and not ok else ""))
        if not ok:
            fails.append(name)

    corn = {2017: 3.36, 2018: 3.61, 2019: 3.56, 2020: 4.53, 2021: 6.00, 2022: 6.54, 2023: 4.55, 2024: 4.24, 2025: 4.16}
    soy = {2017: 9.33, 2018: 8.48, 2019: 8.57, 2020: 10.8, 2021: 13.3, 2022: 14.2, 2023: 12.4, 2024: 10.0, 2025: 10.5}
    wht = {2018: 5.16, 2019: 4.58, 2020: 5.05, 2021: 7.63, 2022: 8.83, 2023: 6.96, 2024: 5.52, 2025: 5.06}

    def win(d, y):
        return [d[k] for k in range(y - 6, y - 1)]

    ck("ERP 2024 corn 4.01 (85%)", erp_calc(win(corn, 2024), 3.70, 0.85)["erp"] == 4.01)
    ck("ERP 2024 soybeans 9.26", erp_calc(win(soy, 2024), 8.40, 0.85)["erp"] == 9.26)
    ck("ERP 2024 wheat 5.50 statutory", erp_calc(win(wht, 2024), 5.50, 0.85)["binding"] == "statutory")
    for y in (2025, 2026):
        ck(f"ERP {y} corn 4.42", erp_calc(win(corn, y), 4.10)["erp"] == 4.42)
        ck(f"ERP {y} soybeans 10.71", erp_calc(win(soy, y), 10.00)["erp"] == 10.71)
    ck("ERP 2026 wheat 6.35", erp_calc(win(wht, 2026), 6.35)["erp"] == 6.35)
    e = erp_calc(win(corn, 2027), 4.10)
    ck("ERP 2027 corn 4.34, 2025 dropped as low", e["erp"] == 4.34 and e["dropped_low"] == 4.16 and e["dropped_high"] == 6.54, str(e))
    ck("ERP 2027 soybeans 10.62", erp_calc(win(soy, 2027), 10.00)["erp"] == 10.62)
    ck("ERP 2027 wheat 6.35", erp_calc(win(wht, 2027), 6.35)["erp"] == 6.35)
    ck("cap 115% binds at 4.72", erp_calc([10, 10, 10, 10, 10], 4.10)["erp"] == 4.72)
    ck("BP 2023 corn 3.98", benchmark_price(win(corn, 2023), 3.70)[0] == 3.98)
    ck("BP 2024 corn 4.85", benchmark_price(win(corn, 2024), 4.01)[0] == 4.85)
    ck("BP 2024 soybeans 11.12", benchmark_price(win(soy, 2024), 9.26)[0] == 11.12)
    ck("BP 2024 wheat 6.21", benchmark_price(win(wht, 2024), 5.50)[0] == 6.21)
    ck("BP 2026 corn 5.03", benchmark_price(win(corn, 2026), 4.42)[0] == 5.03)
    ck("BP 2026 soybeans 12.17", benchmark_price(win(soy, 2026), 10.71)[0] == 12.17)
    ck("BP 2026 wheat 6.98", benchmark_price(win(wht, 2026), 6.35)[0] == 6.98)
    ck("BP 2027 corn 4.96", benchmark_price(win(corn, 2027), 4.34)[0] == 4.96)
    ck("BP 2027 soybeans 12.11", benchmark_price(win(soy, 2027), 10.62)[0] == 12.11)
    ck("BP 2027 wheat 6.98", benchmark_price(win(wht, 2027), 6.35)[0] == 6.98)
    ck("cents half up: 0.125 -> 13, 27.455 -> 2746", cents(0.125) == 13 and cents(27.455) == 2746)
    ck("PLC rate 4.34-4.00 = 0.34", abs(plc_rate(4.34, 4.00, 2.42) - 0.34) < 1e-9)
    ck("PLC rate floored at loan rate", abs(plc_rate(4.34, 2.00, 2.42) - 1.92) < 1e-9)
    ck("ARC capped at 12%", abs(arc_rate(180, 4.96, 162, 4.00, 2.42) - 107.136) < 1e-9)
    ck("ARC uncapped shortfall", abs(arc_rate(180, 4.96, 170, 4.40, 2.42) - 55.52) < 1e-9)
    runs = ranges(4.34, 2.42, 136.16, 170.2, 4.96, 170.2)
    ck("crossover runs plc/arc/none", [x["w"] for x in runs] == ["plc", "arc", "none"], str(runs))
    ck("PLC pays more at 3.59 or lower", plc_below(runs) == 359, str(runs))
    ck("text", ranges_text(runs) == ["PLC pays more at $3.59 or lower.", "ARC-CO pays more from $3.60 to $4.46.",
                                     "Neither pays at $4.47 or higher."], str(ranges_text(runs)))
    ck("break-even depends only on PLC/benchmark share", plc_below(ranges(4.34, 2.42, 80.0, 100.0, 4.96, 100.0)) == 359)
    big = ranges(4.34, 2.42, 150, 170, 4.96, 0.001)
    ck("tiny county yield cannot run the scan past 1.5 x max(ERP, BP)", big[-1]["hi"] <= int(math.ceil(4.96 * 150)), str(big[-1]))
    # scenario engine, 8 hand-worked years: ERP 4.00, BP 5.00, loan 2.00, PLC yield 100, benchmark 100, center $4.00
    R = {2015: 1.0, 2016: 0.75, 2017: 1.25, 2018: 1.25, 2019: 1.0, 2020: 0.5, 2021: 1.1, 2022: 0.9}
    Dy = {2015: 1.0, 2016: 1.0, 2017: 1.0, 2018: 0.8, 2019: 1.2, 2020: 1.0, 2021: 0.9, 2022: 1.1}
    v = scen_eval(4.0, 5.0, 2.0, 100, [(1.0, 100, Dy)], 4.0, R)
    # PLC: 2016 $3.00 -> 1.00 x 100 x .85 = 85; 2020 $2.00 -> 170; 2022 $3.60 -> 34; mean 289/8 = 36.125
    # ARC: 42.5, 51 (cap 60 x .85), 0, 42.5, 0, 51, 45.9, 45.9; mean 278.8/8 = 34.85
    ck("scenario means 36.125 PLC, 34.85 ARC", abs(v["plc"] - 36.125) < 1e-9 and abs(v["arc"] - 34.85) < 1e-9, str(v))
    ck("scenario wins 2 PLC, 4 ARC, 2 ties; SE 19.44 -> too close", v["plc_wins"] == 2 and v["arc_wins"] == 4 and abs(v["se"] - 19.4408) < 1e-3
       and v["verdict"] == "close", str(v))
    v = scen_eval(4.0, 5.0, 2.0, 100, [(1.0, 100, Dy)], 4.0, R, price=3.0)
    ck("at a fixed $3.00: PLC 85 every year, ARC 51 every year -> PLC", v["plc"] == 85 and abs(v["arc"] - 51) < 1e-9 and v["verdict"] == "plc", str(v))
    # Leans: PLC yield 60 -> PLC 51 (2016), 102 (2020), 20.4 (2022): mean 173.4/8 = 21.675 vs ARC 34.85;
    # ARC more in 4 years (2016 is a 51-51 tie), PLC in 1; gap 13.175 >= $3 but under 2 x SE (2 x 11.636)
    v = scen_eval(4.0, 5.0, 2.0, 60, [(1.0, 100, Dy)], 4.0, R)
    ck("leans ARC-CO: 21.675 vs 34.85, wins 1 vs 4, within 2 SE", abs(v["plc"] - 21.675) < 1e-9 and v["plc_wins"] == 1 and v["arc_wins"] == 4
       and abs(v["se"] - 11.636) < 1e-3 and v["verdict"] == "lean_arc", str({k: v[k] for k in v if k != "rows"}))
    R7 = {k: v_ for k, v_ in R.items() if k != 2022}
    ck("fewer than 8 years -> withheld", scen_eval(4.0, 5.0, 2.0, 100, [(1.0, 100, Dy)], 4.0, R7)["verdict"] == "withheld")
    inp0 = load_inputs(ROOT)
    if inp0.get("rma_hist"):
        h = inp0["rma_hist"]["corn"]
        ck("RMA history read from the price tracker: 2012 corn $5.68 / $7.50, 2024 $4.66 / $4.16",
           h.get(2012) == (5.68, 7.50) and h.get(2024) == (4.66, 4.16), str(h.get(2012)))
        sc = price_scen(dict(inp0, wasde={"prices": {"corn": 4.70}, "date": "2026-10-09"}), "corn", 2026)
        # 2024: final MYA 4.24 / October futures average 4.16 = 1.01923; the 2015-2025 average of
        # those ratios is 0.99139, so the spread is centered: $4.70 x 1.01923 / 0.99139 = $4.83
        ck("2026 corn scenario 2024: $4.70 x (4.24/4.16) / 0.99139 = $4.83", sc["kind"] == "oct" and round(sc["center"] * sc["ratios"][2024], 2) == 4.83, str(sc)[:200])
        ck("2026 corn spread: mean 1, 2025 in once its final price is out; yields fixed at a normal crop",
           abs(sum(sc["ratios"].values()) / len(sc["ratios"]) - 1) < 1e-3 and (2025 in sc["ratios"]) == (2025 in inp0["mya"]["corn"]) and sc.get("fix") is True)
        w26 = price_scen(dict(inp0, wasde={"prices": {"wheat": 6.30, "corn": 4.70}, "date": "2026-10-09"}), "wheat", 2026)
        cs = sd1(sc["ratios"].values())
        ck("2026 wheat: own price changes rescaled to corn's October spread, mean 1, borrowed (capped at Leans)",
           w26["borrowed"] and abs(sd1(w26["ratios"].values()) - cs) < 2e-4 and abs(sum(w26["ratios"].values()) / len(w26["ratios"]) - 1) < 1e-3,
           f"{sd1(w26['ratios'].values())} vs {cs}")
        ck("2027 wheat: no 2027 price to calibrate, no call", price_scen(inp0, "wheat", 2027)["kind"] == "noprice")
    # WASDE parser on the text layout of the report
    fake = ("WASDE - 676 Approved by the World Agricultural Outlook Board  October 9, 2026\n"
            "U.S. Wheat Supply and Use  1/\nAvg. Farm Price ($/bu)  2/ 5.52 5.06 6.40 6.30\n"
            "U.S. Feed Grain and Corn Supply and Use  1/\nFEED GRAINS\nCORN \nAvg. Farm Price ($/bu)  4/ 4.24 4.16 4.80 4.70\n"
            "U.S. Soybeans and Products Supply and Use (Domestic Measure)  1/\nSOYBEANS \nAvg. Farm Price ($/bu)  2/ 10.00 10.50 12.00 12.00\n")
    w = parse_wasde(fake)
    ck("WASDE parser: Oct 9, 2026 #676, corn 4.70, soybeans 12.00, wheat 6.30",
       w == {"prices": {"wheat": 6.30, "corn": 4.70, "soybeans": 12.00}, "number": 676, "date": "2026-10-09"}, str(w))
    if os.path.exists(os.path.join(ROOT, "data/wasde-pdf/wasde1026.pdf")):
        w = parse_wasde(wasde_text(os.path.join(ROOT, "data/wasde-pdf/wasde1026.pdf")))
        ck("WASDE Oct 2026 PDF: corn 4.70, soybeans 12.00, wheat 6.30", w.get("prices") == {"wheat": 6.30, "corn": 4.70, "soybeans": 12.00}
           and w.get("date") == "2026-10-09", str(w))
    # published PLC rates the back-test must reproduce: 2019 corn $0.14, 2019 wheat $0.92, 2020 wheat $0.45, 2025 wheat $1.29
    bt = backtest(load_inputs(ROOT))
    got = {(r["y"], r["k"]): r["rate"] for r in bt["rows"]}
    ck("back-test PLC rates 2019 corn 0.14, wheat 0.92; 2020 wheat 0.45; 2025 wheat 1.29",
       got.get((2019, "corn")) == 0.14 and got.get((2019, "wheat")) == 0.92 and got.get((2020, "wheat")) == 0.45 and got.get((2025, "wheat")) == 1.29, str(got))
    # FSA parser on a workbook laid out like FSA's (header row, designations, years)
    try:
        import openpyxl
        import tempfile
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ARCCO 2026 (2026-01-16)"
        ws.append(["ARC-CO TREND ADJUSTED YIELDS, BENCHMARK YIELDS AND GUARANTEE REVENUES FOR PROGRAM YEAR 2026"])
        ws.append([])
        ws.append([])
        ws.append(["ST_Cty", "State Name", "County Name", "Sub County", "Crop Name", "Unit", "ARC-CO Yield Designation",
                   "2020 trend adjusted (county yield or 80% of T)", "2021 trend adjusted (county yield or 80% of T)",
                   "2022 trend adjusted (county yield or 80% of T)", "2023 trend adjusted (county yield or 80% of T)",
                   "2024 trend adjusted (county yield or 80% of T)", None, "2026 Bench Mark (2020-24 olympic avg)",
                   "2026 Bench Mark Price (2020-24 olympic avg)", "2026 Benchmark Revenue", "2026 Guarantee Revenue",
                   "2026 Maximum Payment Rate", "2026 Actual Yield", "2026 National Price", "2026 Actual Revenue",
                   "2026 Formula Payment Rate", "2026 ARC-CO Payment Rate"])
        ws.append(["55017", "Wisconsin", "Chippewa", "", "Corn", "Bushel", "Nonirrigated", 170, 160, 175, 180, 165, None,
                   170.0, 5.03, 855.1, 769.59, 102.61, None, None, None, None, None])
        ws.append(["55017", "Wisconsin", "Chippewa", "", "Corn", "Bushel", "Irrigated", 200, 205, 210, 190, 195, None,
                   200.0, 5.03, 1006.0, 905.4, 120.72, None, None, None, None, None])
        tmp = os.path.join(tempfile.mkdtemp(), "arcco_2026_data_2026-01-16.xlsx")
        wb.save(tmp)
        f = parse_arcco(tmp)
        r = f["rows"].get(("55017", "", "Corn", "non"))
        ck("parser: program year, as-of, designation, benchmark, window",
           f["py"] == 2026 and f["as_of"] == "2026-01-16" and r and r["by"] == 170.0 and r["bp"] == 5.03
           and list(r["yrs"]) == [2020, 2021, 2022, 2023, 2024] and ("55017", "", "Corn", "irr") in f["rows"], str(f)[:300])
        # end to end: the official path, as when FSA's 2026 file lands in data/fsa/
        inp = load_inputs(ROOT)
        inp["fsa"] = dict(inp["fsa"])
        inp["fsa"][2026] = f
        D = compute(inp)
        S = D["states"]["WI"]
        cp, cm = county_page("WI", S, S["c"][0], D, ("", "", ""))
        sp, sm = state_page("WI", S, D, ("", "", ""), [S])
        # the benchmark and both practices are in the page (Key numbers); the working behind them (Olympic average,
        # trigger yield, break-even) is in data/arc-plc/math/WI.json, which components/arc-plc-math.js lays out on open
        # (test/arc-plc-math.test.mjs renders it against the old server-rendered math)
        mj = json.loads(math_json(D)[f"{OUT_MATH}/WI.json"])
        me = (mj["c"].get("55017") or {}).get("e", {}).get("corn", [])
        ck("official path: county page indexed, official benchmark in Key numbers, both practices, math in the state's math file",
           cm["indexable"] and 'id="key-numbers"' in cp and "170.00 bu" in cp and "200.00 bu" in cp and 'content="index,follow"' in cp
           and f'data-math="/{OUT_MATH}/WI.json" data-f="55017"' in cp and sorted((e["d"], e["by"]) for e in me) == [("irr", 200), ("non", 170)]
           and all(len(e["ys"]) == 5 and bin(e["x"]).count("1") == 2 for e in me), cm["title"])
        ck("official path: counter sheet link goes to the shared sheet, no per-county sheet page",
           'href="/arc-plc/sheet?c=wisconsin/chippewa-county"' in cp and SHEET_SUFFIX + '"' not in cp)
        ck("official path: calculator file carries both practices",
           [e["d"] for e in calc_json(D)[1]["WI"]["c"][0]["k"]["corn"]] == ["irr", "non"])
        bad = dict(f)
        bad["rows"] = {k: dict(v, bp=4.00) for k, v in f["rows"].items()}
        inp["fsa"][2026] = bad
        try:
            compute(inp)
            ck("refuses when FSA's benchmark price disagrees with ours", False)
        except SystemExit:
            ck("refuses when FSA's benchmark price disagrees with ours", True)
        inp["fsa"].pop(2026)
        D = compute(inp)
        Sx = D["states"].get("WI")
        if Sx:
            cx = next(c for c in Sx["c"] if c["f"] == "55017")
            cp, cm = county_page("WI", Sx, cx, D, ("", "", ""))
            ck("pending path: noindex, no benchmark kv, no trigger yield in bushels",
               not cm["indexable"] and "Benchmark yield (FSA)" not in cp and "county yield comes in below" not in cp
               and 'content="noindex,follow"' in cp)
    except ImportError:
        ck("openpyxl installed", False)
    ck("no em dash in range text", "—" not in " ".join(ranges_text(runs)))
    # ---- size budget (numbers from scripts/check_site_budget.py)
    import check_site_budget as SB
    small = {f"{OUT_DIR}/a.html": "x" * 1000, "data/arc-plc/math/WI.json": "x" * (SB.SET_PAGE_KB * 1024 + 10)}
    ck("budget: small pages pass, data files are not pages", budget_problems(small)[0] == [])
    fat = dict(small, **{f"{OUT_DIR}/texas.html": "x" * (SB.SET_PAGE_KB * 1024 + 1)})
    ck("budget: a page over the cap is named", any("texas.html" in b for b in budget_problems(fat)[0]))
    _cap = SB.FOLDER_MB.get(OUT_DIR, SB.DEFAULT_FOLDER_MB)
    try:
        SB.FOLDER_MB[OUT_DIR] = 0.0005
        ck("budget: arc-plc/ over its folder budget refuses", any("over its" in b for b in budget_problems(small)[0]))
    finally:
        SB.FOLDER_MB[OUT_DIR] = _cap
    stub = sheet_stub({"slug": "wisconsin", "n": "Wisconsin"}, {"s": "chippewa-county", "n": "Chippewa County"})
    ck("old sheet URL: a stub under 1 KB that forwards to the shared sheet, canonical to the county",
       len(stub.encode()) < 1024 and 'url=/arc-plc/sheet?c=wisconsin/chippewa-county"' in stub and 'location.replace("/arc-plc/sheet?c=wisconsin/chippewa-county"' in stub
       and 'rel="canonical" href="https://agsist.com/arc-plc/wisconsin/chippewa-county"' in stub and "noindex" in stub, str(len(stub)))
    vp = vpack({"verdict": "lean_arc", "plc": 0.905, "arc": 5.2049, "diff": -4.2999, "plc_wins": 0, "arc_wins": 4, "n": 11,
                "rows": [{"arc": 1}, {"arc": 0}, {"arc": 2}], "borrowed": True})
    # 0.905 -> $0.91 (half up after rnd); 5.2049 -> $5.20; |gap| 4.2999 -> $4.30, whole dollars 4; ARC-CO paid in 2 rows; borrowed flag 1
    ck("math file: a packed verdict carries the cents as printed", vp == ["lean_arc", 91, 520, 0, 4, 11, 430, 2, 1, 4], str(vp))
    ck("math file: no-call verdicts carry only the reason", vpack({"verdict": "pending", "n": 0}) == ["pending"]
       and vpack({"verdict": "withheld", "n": 6}) == ["withheld", 6])
    # ---- Use my location: point in polygon (mirrors components/arc-plc-simple.js)
    sq, hole = [[0, 0], [10, 0], [10, 10], [0, 10], [0, 0]], [[4, 4], [6, 4], [6, 6], [4, 6], [4, 4]]
    navx = {"s": {"XX": {"b": [0, 0, 10, 10]}}}
    geox = {"XX": {"c": [["99001", "Square", "square-county", [0, 0, 10, 10], [[sq, hole]]]]}}
    ck("point in polygon: inside, inside the hole, outside", locate(navx, geox, 2, 2)[1][0] == "99001"
       and locate(navx, geox, 5, 5) == (None, None) and locate(navx, geox, 11, 5) == (None, None))
    gp = os.path.join(ROOT, GEO_SRC)
    if os.path.exists(gp):
        Dg = {"_geo": _read(ROOT, GEO_SRC), "states": {"WI": {"n": "Wisconsin", "slug": "wisconsin", "c": [{"f": "55017", "n": "Chippewa County", "s": "chippewa-county"}]}}}
        gj = geo_json(Dg)
        nv = gj[f"{OUT_STATE_JSON}/nav.json"]
        gs = {k.split("/")[-1][:-5]: v for k, v in gj.items() if "/geo/" in k}
        st_, rec = locate(nv, gs, -91.29, 44.94)
        ck("Chippewa's coordinates (44.94, -91.29) land on 55017 with its page; Eau Claire's (44.73, -91.29) on 55035 without one",
           st_ == "WI" and rec[0] == "55017" and rec[2] == "chippewa-county" and locate(nv, gs, -91.29, 44.73)[1][0] == "55035"
           and locate(nv, gs, -91.29, 44.73)[1][2] is None, str(rec[:3] if rec else None))
    # ---- QR code (scripts/qr_svg.py): structure, and the exact symbol a QR reader decoded on 2026-10-09
    import hashlib
    import qr_svg
    qm = qr_svg.qr_matrix("https://agsist.com/arc-plc/wisconsin/chippewa-county")
    finder = all(qm[r][c] == (max(abs(r - 3), abs(c - 3)) not in (2, 4)) for r in range(7) for c in range(7))
    ck("QR: version 4 (33 x 33), finder pattern, timing row, dark module",
       len(qm) == 33 and finder and all(qm[6][x] == (x % 2 == 0) for x in range(8, 25)) and qm[33 - 8][8])
    ck("QR: the same symbol jsQR 1.4.0 decoded to the Chippewa URL",
       hashlib.sha256("".join("".join("1" if v else "0" for v in r) for r in qm).encode()).hexdigest()
       == "53f0997e75ec4eea2bde0ddde10308671f6c95b883fec9e80342241a9de69f15")
    # ---- the call (scen_call), by hand
    # t threshold, 10 years (9 df, t 2.262): a $10 gap with SE 4.60 is under 2.262 x 4.60 = 10.41 -> Leans; SE 4.40 -> 9.95 -> Pick
    ck("t, 9 df: $10 gap, SE 4.60 -> Leans PLC; SE 4.40 -> Pick PLC",
       scen_call(10, 0, 10, 4.60, 5, 0, 10) == "lean_plc" and scen_call(10, 0, 10, 4.40, 5, 0, 10) == "plc")
    # 11 years (10 df, t 2.228): SE 4.60 -> 10.25 still over $10 -> Leans; SE 4.45 -> 9.91 -> Pick
    ck("t, 10 df: SE 4.60 -> Leans; SE 4.45 -> Pick", scen_call(10, 0, 10, 4.60, 5, 0, 11) == "lean_plc" and scen_call(10, 0, 10, 4.45, 5, 0, 11) == "plc")
    ck("the old 2 x SE rule would have picked at SE 4.60 (2 x 4.60 = 9.20 < 10); t does not", 10 > 2 * 4.60 and scen_call(10, 0, 10, 4.60, 5, 0, 10) != "plc")
    # 3-win rule: a $10 gap with a tiny SE, but the leader paid more in only 2 years -> close
    ck("leader won 2 of 10 years -> Close, however big the gap", scen_call(0, 10, -10, 0.5, 0, 2, 10) == "close"
       and scen_call(0, 10, -10, 0.5, 0, 3, 10) == "arc")
    ck("Leans needs one SE: $4 gap, SE 4.10 -> Close; SE 3.90 -> Leans",
       scen_call(4, 0, 4, 4.10, 4, 0, 10) == "close" and scen_call(4, 0, 4, 3.90, 4, 0, 10) == "lean_plc")
    # both expected payments round to $0: neither expected to pay
    ck("$0.49 vs $0.20 -> neither; $0.50 vs $0 -> not neither", scen_call(0.49, 0.20, 0.29, 0.1, 3, 0, 10) == "none"
       and scen_call(0.50, 0, 0.5, 0.1, 3, 0, 10) == "close")
    ck("borrowed spread caps a Pick at Leans", scen_call(10, 0, 10, 1, 8, 0, 10, borrowed=True) == "lean_plc")
    # futures-center downgrade, by hand: ERP 4, BP 5, loan 2, PLC yield 100, benchmark 100, normal crop every year.
    # USDA start $3.00: prices 3.00 2.94 3.06 3.00 2.97 3.03 3.00 3.00 -> PLC (4 - p) x 85: mean 85; ARC capped 60 x .85 = 51 every year:
    # gap 34, SE about 1.2 -> Pick PLC. Futures start $4.50: PLC 0; ARC 7.65 (at $4.41) and 3.825 (at $4.455), mean 1.43 -> ARC leads: Pick drops to Leans.
    Rf = {2015: 1.0, 2016: 0.98, 2017: 1.02, 2018: 1.0, 2019: 0.99, 2020: 1.01, 2021: 1.0, 2022: 1.0}
    v1 = scen_eval(4.0, 5.0, 2.0, 100, [(1.0, 100, 1.0)], 3.0, Rf)
    v2 = scen_eval(4.0, 5.0, 2.0, 100, [(1.0, 100, 1.0)], 3.0, Rf, alt=4.5)
    ck("normal crop every year (dy = 1.0): PLC 85, ARC 51 -> Pick PLC", abs(v1["plc"] - 85) < 1e-9 and abs(v1["arc"] - 51) < 1e-9 and v1["verdict"] == "plc", str(v1)[:200])
    ck("futures start disagrees on the leader -> Leans PLC", v2["verdict"] == "lean_plc" and v2.get("alt_down") and abs(v2["alt"]["arc"] - 11.475 / 8) < 1e-9, str(v2.get("alt")))
    ck("futures start agrees -> stays Pick", scen_eval(4.0, 5.0, 2.0, 100, [(1.0, 100, 1.0)], 3.0, Rf, alt=3.1)["verdict"] == "plc")
    ck("a typed price ignores the borrowed cap and the futures check",
       scen_eval(4.0, 5.0, 2.0, 100, [(1.0, 100, 1.0)], 3.0, Rf, price=3.0, borrowed=True, alt=4.5)["verdict"] == "plc")
    ck("rescale_to: spread equals the target, mean 1", abs(sd1(rescale_to({1: 1.0, 2: 1.2, 3: 0.9, 4: 1.1}, 0.05).values()) - 0.05) < 2e-4)
    one = [{"w": "plc", "lo": 300, "hi": 359}, {"w": "arc", "lo": 360, "hi": 360}, {"w": "plc", "lo": 361, "hi": 400}, {"w": "none", "lo": 401, "hi": 500}]
    ck("a one-cent band reads 'at $3.60'; every band is printed, the later PLC band too",
       ranges_text(one) == ["PLC pays more at $3.59 or lower.", "ARC-CO pays more at $3.60.", "PLC pays more from $3.61 to $4.00.", "Neither pays at $4.01 or higher."],
       str(ranges_text(one)))
    # FSA's county payment rates are recomputed; a row that disagrees stops the build
    good = {"py": 2025, "rows": {("55017", "", "Corn", "non"): {"by": 177.34, "bp": 5.03, "ay": 176.84, "np": 4.16, "pay": 67.17}}, "years": []}
    ck("FSA 2025 Chippewa corn: 0.9 x 892.02 - 176.84 x 4.16 = 802.82 - 735.65 = 67.17, as printed", check_fsa_pay({2025: good}, {}) == {2025: 1})
    try:
        check_fsa_pay({2025: {**good, "rows": {k: dict(v, pay=60.0) for k, v in good["rows"].items()}}}, {})
        ck("refuses a county payment rate that does not recompute", False)
    except SystemExit:
        ck("refuses a county payment rate that does not recompute", True)
    # newest WASDE by report month, not by file name: Jan 2027 (0127) sorts after Oct 2026 (1026)
    ck("WASDE pick: wasde0127 after wasde1226 and wasde1026",
       max(["data/wasde-pdf/wasde1026.pdf", "data/wasde-pdf/wasde0127.pdf", "data/wasde-pdf/wasde1226.pdf"], key=wasde_key).endswith("0127.pdf"))
    # signup copy on fixed dates: Dec 11 is the last day of 2026 signup, Mar 15 of 2027
    d = dt.date
    s1, s2, s3, s4 = (signup_copy(x) for x in (d(2026, 10, 9), d(2026, 12, 11), d(2026, 12, 12), d(2027, 3, 16)))
    ck("signup Oct 9, 2026: 2026 open, 2027 upcoming", s1["state"] == {2026: "open", 2027: "upcoming"} and s1["box6"] == "2026 signup closes", str(s1["state"]))
    ck("signup Dec 11, 2026: 2026 still open", s2["state"][2026] == "open" and "closes Dec 11, 2026" in s2["quick"], s2["quick"])
    ck("signup Dec 12, 2026: 2026 closed, 2027 open, no 'closes Dec 11' left",
       s3["state"] == {2026: "closed", 2027: "open"} and "closes" not in s3["quick"] + s3["box6"] + s3["llms"] and s3["desc"] == "2027 signup ends Mar 15.", str(s3))
    ck("signup Mar 16, 2027: both closed", s4["state"] == {2026: "closed", 2027: "closed"} and "open" not in s4["quick"] + s4["faq"], str(s4))
    ck("signup copy has no em dash", "—" not in json.dumps([s1, s2, s3, s4]))
    sn = {"futures": {"corn": {"price": 5.12, "date": "2026-10-09"}, "soybeans": {"price": 12.45, "date": "2026-10-09"}}}
    ck("snapshot 10 days old passes, 11 days refused",
       snapshot_problem(sn, d(2026, 10, 19)) is None and "11 days" in (snapshot_problem(sn, d(2026, 10, 20)) or ""))
    ck("snapshot missing or without soybeans refused",
       snapshot_problem(None, d(2026, 10, 9)) and snapshot_problem({"futures": {"corn": sn["futures"]["corn"]}}, d(2026, 10, 9)))
    # ---- crops beyond corn, soybeans and wheat (scripts/arc_plc_crops.py)
    # pound crop, peanuts 2026: 0.21, 0.243, 0.268, 0.269, 0.261; drop 0.21 and 0.269; (0.243+0.268+0.261)/3 = 0.257333;
    # x 0.88 = 0.226453 -> 0.2265, under the $0.315 statutory price, which holds; benchmark price: every year raised to 0.315
    e = XC.erp_x([0.21, 0.243, 0.268, 0.269, 0.261], 0.315, 4)
    ck("peanuts 2026: 88% figure 0.2265, ERP 0.315 statutory, cap 0.3623, benchmark 0.315",
       e["pct_value"] == 0.2265 and e["erp"] == 0.315 and e["binding"] == "statutory" and e["cap"] == 0.3623
       and XC.bp_x([0.21, 0.243, 0.268, 0.269, 0.261], 0.315, 4)[0] == 0.315, str(e))
    # flaxseed (bushel, four decimals): 11.1, 25.9, 17.5, 12.1, 12.5 -> (12.1+12.5+17.5)/3 = 14.0333 x 0.88 = 12.3493; ERP 13.30
    ck("flaxseed 2026: 12.3493, ERP 13.30, benchmark 14.70", XC.erp_x([11.1, 25.9, 17.5, 12.1, 12.5], 13.3, 4)["pct_value"] == 12.3493
       and XC.bp_x([11.1, 25.9, 17.5, 12.1, 12.5], 13.3, 4)[0] == 14.7)
    # peanut county, by hand: ERP .315, loan .195, PLC yield 3,300 lb, benchmark 4,000 lb at $.315, county yield 3,400 lb.
    # ARC cap 0.12 x 1,260 = 151.20 (128.52/base acre) binds below (1,134 - 151.20)/3,400 = $.2891; PLC 2,805 x (.315 - p) beats
    # 128.52 below $.2692 (at .2691: 128.75 vs 128.52; at .2692: 128.47); ARC pays until 1,134 - 3,400p hits 0 at $.3335
    pr = ranges(0.315, 0.195, 3300, 4000, 0.315, 3400, scale=10000)
    ck("peanut ranges in hundredths of a cent", ranges_text(pr, 10000) == ["PLC pays more at $0.2691 or lower.", "ARC-CO pays more from $0.2692 to $0.3335.",
                                                                         "Neither pays at $0.3336 or higher."], str(ranges_text(pr, 10000)))
    ck("peanut payments at $0.24: PLC 210.375, ARC capped 128.52", abs(per_base(plc_rate(0.315, 0.24, 0.195) * 3300) - 210.375) < 1e-9
       and abs(per_base(arc_rate(4000, 0.315, 3600, 0.24, 0.195)) - 128.52) < 1e-9)
    # hundredweight crop: WASDE prints long grain rice in $/cwt; FSA prices it per pound. $14.00/cwt -> $0.1400/lb
    wx = XC.parse_wasde_x("U.S. Sorghum, Barley, and Oats Supply and Use 1/\nSORGHUM \nAvg. Farm Price ($/bu)  2/ 4.07 3.67 4.60 4.50\n"
                          "BARLEY \nAvg. Farm Price ($/bu)  2/ 6.31 5.46 5.70 5.70\nOATS \nAvg. Farm Price ($/bu)  2/ 3.35 3.23 3.35 3.35\n"
                          "LONG-GRAIN RICE \n  Avg. Farm Price ($/cwt)  6/ 14.00 10.40 13.50 14.00\n")
    ck("WASDE other crops: sorghum 4.50, barley 5.70, oats 3.35, long grain rice $14.00/cwt = $0.1400/lb",
       wx == {"sorghum": (4.5, "$/bu", 4.5), "barley": (5.7, "$/bu", 5.7), "oats": (3.35, "$/bu", 3.35), "rice_long": (0.14, "$/cwt", 14.0)}, str(wx))
    ck("long grain rice PLC rate at $0.14/lb: 0.169 - 0.14 = 0.029", abs(plc_rate(0.169, 0.14, 0.077) - 0.029) < 1e-12)
    xt = XC.load_tables(load_xlsx, ROOT)
    if xt:
        c26, _p = XC.year_params_x(xt, 2026)
        bad = [k for k, cd in c26.items() if abs(cd["erp"]["erp"] - cd["fsa_check"]["erp"]) > 1e-9 or abs(cd["bp"]["value"] - cd["fsa_check"]["bp"]) > 1e-9]
        ck(f"every 2026 ERP and benchmark price in FSA's tables reproduces ({len(c26)} crops)", not bad and len(c26) == len(XC.XCROPS), str(bad))
        c27, p27 = XC.year_params_x(xt, 2027)
        # long grain 2021-2025: 0.136, 0.167, 0.159, 0.14, 2025 projected; at any 2025 price 88% of the middle three is at most
        # 0.88 x (0.167+0.159+0.14)/3 = 0.1367 < 0.169, and every year is raised to 0.169: ERP and benchmark 0.169 whatever 2025 brings
        ck("2027 long grain rice fixed at 0.169 whatever the 2025 price; sorghum waits on its final 2025 price",
           c27.get("rice_long", {}).get("erp", {}).get("erp") == 0.169 and c27["rice_long"]["bp"]["value"] == 0.169 and "sorghum" in p27, str(p27.get("sorghum")))
    # 65 characters with the yield, so the yield goes first; a longer name then takes "Co."
    ck("title: Pottawattamie drops the yield, Prince George's takes Co.",
       county_title("Pottawattamie County", "IA", "Corn", "2026", "203.4 bu") == "Pottawattamie County, IA ARC or PLC 2026: Corn Benchmark"
       and county_title("Prince George's County", "MD", "Soybeans", "2026", "45.1 bu") == "Prince George's Co., MD ARC or PLC 2026: Soybeans Benchmark")
    ck("title: always 60 or fewer with 'ARC or PLC'", all(len(t) <= 60 and "ARC or PLC" in t for t in (
        county_title("Prince of Wales-Hyder Census Area", "AK", "Medium and short grain rice", "2027", "7,000 lb"),
        county_title("Story County", "IA", "Corn", "2026", "203.4 bu"))))
    print(f"\n{'FAIL' if fails else 'OK'}: {len(fails)} failure(s)")
    return 1 if fails else 0


def main():
    ap = argparse.ArgumentParser(add_help=False)
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--check", action="store_true")
    ap.add_argument("--snapshot-prices", action="store_true")
    ap.add_argument("--today", default="")
    a = ap.parse_args()
    os.chdir(ROOT)
    global BUILD_DATE
    BUILD_DATE = dt.date.fromisoformat(a.today) if a.today else dt.datetime.now(dt.timezone.utc).date()
    if a.snapshot_prices:
        print(json.dumps(snapshot_prices("."), indent=1))
    if a.selftest:
        return selftest()
    return main_build(".", check=a.check)


if __name__ == "__main__":
    sys.exit(main())
