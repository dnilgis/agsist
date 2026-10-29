#!/usr/bin/env python3
"""build_farm_credit.py -- farm loan repayment and renewals by Federal Reserve
district, from the district surveys of agricultural bankers.

Writes data/farm-credit/districts.json. The Farmland Atlas state pages show the
district (or both districts, for a split state) beside the Chapter 12 line, and
the Atlas "By state" hub shows every reporting district side by side
(scripts/build_atlas_pages.py reads the file; no file, no section).

WHAT IS SHOWN
  Two diffusion indexes per district, quarterly: Loan Repayment Rate and
  Renewals or Extensions (series 003 and 004 in the Kansas City Fed's table),
  and the latest Collateral Requirements reading (005). Each quarter bankers
  say whether the item was higher, lower or the same as a year earlier; the
  index is the percent saying higher minus the percent saying lower, plus 100
  (the Kansas City Fed's note; the St. Louis and Dallas Feds describe the same
  arithmetic, Dallas without the 100). Values are the table's own, rounded to
  2 decimals only to drop float noise (85.71000000000001 -> 85.71).

INPUT (raw branch, never deployed; the sandbox cannot reach the reserve banks)
  history/fed-ag-credit/<date>/kc-fed/FederalReserveDistrictSurveysofAgCreditConditionsHistoricalData*.xlsx
      the primary file: 7 districts x 20 series, 1991Q1 on.
  history/fed-ag-credit/<date>/01-chicagofed-credit-conditions-7th-district-xls.xls
  history/fed-ag-credit/<date>/10-dallasfed-aglending.xlsx
      the two districts' own files. Every Chicago and Dallas repayment and
      renewals quarter is checked against them. A quarter where the two files
      disagree (beyond the rounding of the shorter one) is WITHHELD: it is left
      out of the chart and listed under "withheld". A latest quarter that
      disagrees withholds the district's numbers altogether.

    bash scripts/pull_raw.sh history/fed-ag-credit/2026-10-28
    python3 scripts/build_farm_credit.py --raw history/fed-ag-credit/2026-10-28
    python3 scripts/build_farm_credit.py --check       the committed file passes every rule
    python3 scripts/build_farm_credit.py --selftest

NEXT RELEASE (by hand, no schedule; a new schedule needs Sig's OK): save the new
files to the raw branch under a new dated folder, then run the "Farm credit by
Fed district" workflow (.github/workflows/farm-credit.yml) with that folder.
Stdlib + openpyxl (xlsx) + xlrd (the Chicago .xls).
"""
import argparse
import datetime as dt
import glob
import hashlib
import json
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "farm-credit", "districts.json")
SCHEMA = "agsist-farm-credit/1"
CHART_FROM = "2010Q1"
KC_PAGE = "https://www.kansascityfed.org/agriculture/ag-credit-survey/"

# The seven districts in the Kansas City Fed's table, plus the five that have no
# farm credit survey there (named on split-state pages only).
DISTRICTS = {
    "CHI": {"name": "Chicago", "num": 7, "url": "https://www.chicagofed.org/"},
    "STL": {"name": "St. Louis", "num": 8, "url": "https://www.stlouisfed.org/"},
    "MIN": {"name": "Minneapolis", "num": 9, "url": "https://www.minneapolisfed.org/"},
    "KC": {"name": "Kansas City", "num": 10, "url": KC_PAGE},
    "DAL": {"name": "Dallas", "num": 11, "url": "https://www.dallasfed.org/"},
    "RICH": {"name": "Richmond", "num": 5, "url": "https://www.richmondfed.org/"},
    "SF": {"name": "San Francisco", "num": 12, "url": "https://www.frbsf.org/"},
}
NO_SURVEY = {"BOS": "Boston", "NY": "New York", "PHI": "Philadelphia", "CLE": "Cleveland", "ATL": "Atlanta"}

# Federal Reserve district of each state. A state wholly inside one district has
# one entry; a split state lists every district that covers part of it. The
# Tenth District line is the Kansas City Fed's own footnote ("Colorado, Kansas,
# Nebraska, Oklahoma, Wyoming, the northern half of New Mexico, and the western
# third of Missouri").
STATE_DISTRICTS = {
    "IA": ["CHI"], "AR": ["STL"], "MN": ["MIN"], "MT": ["MIN"], "ND": ["MIN"], "SD": ["MIN"],
    "CO": ["KC"], "KS": ["KC"], "NE": ["KC"], "OK": ["KC"], "WY": ["KC"], "TX": ["DAL"],
    "MD": ["RICH"], "NC": ["RICH"], "SC": ["RICH"], "VA": ["RICH"], "DC": ["RICH"],
    "AK": ["SF"], "AZ": ["SF"], "CA": ["SF"], "HI": ["SF"], "ID": ["SF"], "NV": ["SF"], "OR": ["SF"], "UT": ["SF"], "WA": ["SF"],
    "ME": ["BOS"], "MA": ["BOS"], "NH": ["BOS"], "RI": ["BOS"], "VT": ["BOS"], "NY": ["NY"], "DE": ["PHI"], "OH": ["CLE"],
    "AL": ["ATL"], "FL": ["ATL"], "GA": ["ATL"],
    "IL": ["CHI", "STL"], "IN": ["CHI", "STL"], "MI": ["CHI", "MIN"], "WI": ["CHI", "MIN"], "MO": ["KC", "STL"],
    "NM": ["KC", "DAL"], "KY": ["STL", "CLE"], "TN": ["STL", "ATL"], "MS": ["STL", "ATL"], "LA": ["DAL", "ATL"],
    "WV": ["RICH", "CLE"], "PA": ["CLE", "PHI"], "NJ": ["NY", "PHI"], "CT": ["NY", "BOS"],
}
SERIES = {"repay": ("003", "Loan Repayment Rate"), "renew": ("004", "Renewals or Extensions"),
          "collat": ("005", "Collateral Requirements")}
HOW = ("Each quarter, ag bankers in the district say whether each item was higher, lower or the same as a year "
       "earlier. The index is the percent saying higher, minus the percent saying lower, plus 100. Over 100 means "
       "more said higher; under 100, more said lower; 100 is an even split.")
WORDS = {
    "repay": "Under 100: more bankers said repayment rates fell than rose.",
    "renew": "Over 100: more bankers said renewals or extensions rose than fell.",
    "collat": "Over 100: more bankers asked for more collateral than less.",
}
KC_PANEL = "From 2025 Q3 the Kansas City Fed's results also count Farm Credit System lenders, not only commercial banks."
Q = re.compile(r"^(\d{4})Q([1-4])$")


def qkey(q):
    m = Q.match(q)
    if not m:
        raise ValueError("not a quarter: %r" % (q,))
    return int(m.group(1)) * 4 + int(m.group(2)) - 1


def qlabel(q):
    """'2026Q2' -> '2026 Q2'."""
    qkey(q)
    return q[:4] + " Q" + q[5]


def qrange(a, b):
    out = []
    for k in range(qkey(a), qkey(b) + 1):
        out.append("%dQ%d" % (k // 4, k % 4 + 1))
    return out


def num(v):
    """A cell as the table prints it, or None for '---', 'n/a' and blanks."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        r = round(float(v), 2)
        return int(r) if r == int(r) else r
    return None


def decimals(v):
    s = repr(float(v)).rstrip("0").rstrip(".")
    return len(s.split(".")[1]) if "." in s else 0


def agree(kc, own):
    """True when the Kansas City table's value is the district's own value, allowing
    only the rounding of whichever of the two prints fewer decimals."""
    if kc is None or own is None:
        return kc is None and own is None
    d = min(decimals(kc), decimals(own))
    return abs(kc - own) <= 0.5 * 10 ** -d + 1e-9


def read_kc(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Data" not in wb.sheetnames or "Description" not in wb.sheetnames:
        raise ValueError("%s: no Data or Description sheet" % path)
    desc = {}
    for r in wb["Description"].iter_rows(values_only=True):
        if r and isinstance(r[0], str) and re.match(r"^\d{3}[A-Z]+$", r[0].strip()):
            desc[r[0].strip()] = {"name": str(r[1] or "").strip(), "unit": str(r[2] or "").strip(), "district": str(r[3] or "").strip()}
    rows = list(wb["Data"].iter_rows(values_only=True))
    head = [str(h).strip() if h is not None else "" for h in rows[0]]
    if head[0] != "Date":
        raise ValueError("%s: Data sheet does not start with Date" % path)
    data = {}
    for r in rows[1:]:
        if not r or not isinstance(r[0], str) or not Q.match(r[0].strip()):
            continue
        q = r[0].strip()
        for i, h in enumerate(head[1:], 1):
            if h:
                data.setdefault(h, {})[q] = num(r[i]) if i < len(r) else None
    return desc, data


def read_chicago(path):
    """Chicago Fed's own file: YYYYQ codes, columns demand, funds, repayment, renewals."""
    import xlrd
    s = xlrd.open_workbook(path).sheet_by_index(0)
    hdr = " ".join(str(c) for i in range(min(6, s.nrows)) for c in s.row_values(i))
    if "Repayment" not in hdr or "Renewals" not in hdr:
        raise ValueError("%s: no Repayment or Renewals column" % path)
    out = {"003": {}, "004": {}}
    for i in range(s.nrows):
        c = str(s.cell_value(i, 0)).replace(".0", "")
        if not re.match(r"^\d{5}$", c):
            continue
        q = "%sQ%s" % (c[:4], c[4])
        out["003"][q] = num(s.cell_value(i, 3)) if s.cell_value(i, 3) != "" else None
        out["004"][q] = num(s.cell_value(i, 4)) if s.cell_value(i, 4) != "" else None
    return out


def read_dallas(path):
    """Dallas Fed's own lending file: index = greater minus less (no 100); add 100 to compare."""
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    top = next(i for i, r in enumerate(rows) if r and "Rate of Loan Repayment" in [str(c) for c in r])
    names = [str(c) if c is not None else "" for c in rows[top]]
    ci = {"003": names.index("Rate of Loan Repayment"), "004": names.index("Loan Renewals or Extensions")}
    for k, i in ci.items():
        if str(rows[top + 1][i]).strip() != "Index1":
            raise ValueError("%s: column %d is not an index" % (path, i))
    out = {"003": {}, "004": {}}
    for r in rows[top + 2:]:
        m = re.match(r"^(\d{4}): Q([1-4])$", str(r[0] or ""))
        if not m:
            continue
        q = "%sQ%s" % m.groups()
        for k, i in ci.items():
            v = r[i]
            out[k][q] = num(100 + v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None
    return out


def sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(1 << 16), b""):
            h.update(b)
    return h.hexdigest()


def assemble(desc, data, own, built, files):
    """own: {district: {'003': {q: v}, '004': {...}}} from the districts' own files."""
    quarters = sorted({q for s in data.values() for q in s}, key=qkey)
    latest = quarters[-1]
    for code in ("003", "004", "005"):
        for d in DISTRICTS:
            sid = code + d
            if sid not in desc:
                raise ValueError("%s is not in the Description sheet" % sid)
    for key, (code, nm) in SERIES.items():
        for d in DISTRICTS:
            got = desc[code + d]
            if got["name"] != nm or got["unit"] != "Diffusion Index" or got["district"] != DISTRICTS[d]["name"]:
                raise ValueError("%s%s reads %r, expected %s, Diffusion Index, %s" % (code, d, got, nm, DISTRICTS[d]["name"]))
    withheld, districts = [], {}
    for d, meta in DISTRICTS.items():
        ser = {k: {q: data.get(SERIES[k][0] + d, {}).get(q) for q in quarters} for k in SERIES}
        bad = set()
        if d in own:
            for k in ("repay", "renew"):
                code = SERIES[k][0]
                for q in quarters:
                    kc, ov = ser[k][q], own[d][code].get(q, "missing")
                    if ov == "missing" or (kc is None and ov is None):
                        continue
                    if not agree(kc, ov):
                        withheld.append({"district": d, "q": q, "series": k, "kc": kc, "own": ov})
                        bad.add((k, q))
        last = max((q for q in quarters if ser["repay"][q] is not None), key=qkey, default=None)
        cur = last == latest
        lat = None
        if cur:
            lat = {k: (None if (k, latest) in bad else ser[k][latest]) for k in SERIES}
            if lat["repay"] is None or lat["renew"] is None:
                lat = None
                cur = False
        chart_q = qrange(CHART_FROM, latest)
        chart = {k: [None if (k, q) in bad else ser[k].get(q) for q in chart_q] for k in ("repay", "renew")}
        first = {k: min((q for q in quarters if ser[k][q] is not None), key=qkey, default=None) for k in ("repay", "renew")}
        districts[d] = {
            "name": meta["name"], "num": meta["num"], "url": meta["url"], "current": cur, "last": last,
            "first": first, "latest": lat, "chart": chart,
            "checked_against_own_file": d in own,
            "note": KC_PANEL if d == "KC" else "",
            "mark": {"q": "2025Q3", "label": "Farm Credit lenders added"} if d == "KC" else None,
        }
    return {
        "schema": SCHEMA,
        "about": ("Federal Reserve district surveys of agricultural bankers, as compiled by the Federal Reserve Bank of "
                  "Kansas City. Diffusion indexes, quarterly. Built by scripts/build_farm_credit.py; Chicago and Dallas "
                  "quarters are checked against those banks' own files and any that disagree are withheld."),
        "source": "Federal Reserve district agricultural credit surveys, compiled by the Federal Reserve Bank of Kansas City",
        "source_url": KC_PAGE,
        "built": built,
        "latest": latest,
        "chart_from": CHART_FROM,
        "how": HOW,
        "words": WORDS,
        "series": {k: {"id": v[0], "name": v[1]} for k, v in SERIES.items()},
        "districts": districts,
        "no_survey": NO_SURVEY,
        "states": STATE_DISTRICTS,
        "withheld": withheld,
        "files": files,
    }


def dumps(doc):
    """Indented, with each list of numbers and each withheld row on one line (a smaller, readable diff)."""
    t = json.dumps(doc, indent=1, ensure_ascii=False)
    t = re.sub(r"\[\s*((?:-?[\d.]+|null)(?:,\s*(?:-?[\d.]+|null))*)\s*\]", lambda m: "[" + re.sub(r",\s*", ",", m.group(1)) + "]", t)
    t = re.sub(r"\{\s*(\"district\"[^{}]*?)\s*\}", lambda m: "{" + re.sub(r",\s*", ", ", re.sub(r"\s*\n\s*", " ", m.group(1))) + "}", t)
    return t + "\n"


def build(raw, out=OUT, built=None):
    kc = sorted(glob.glob(os.path.join(raw, "kc-fed", "FederalReserveDistrictSurveysofAgCreditConditionsHistoricalData*.xlsx")))
    chi = glob.glob(os.path.join(raw, "*chicagofed-credit-conditions*.xls"))
    dal = glob.glob(os.path.join(raw, "*dallasfed-aglending.xlsx"))
    if len(kc) != 1 or len(chi) != 1 or len(dal) != 1:
        raise SystemExit("need exactly one Kansas City table, one Chicago and one Dallas file in %s (found %d, %d, %d)" % (raw, len(kc), len(chi), len(dal)))
    desc, data = read_kc(kc[0])
    own = {"CHI": read_chicago(chi[0]), "DAL": read_dallas(dal[0])}
    files = [{"name": os.path.basename(p), "sha256": sha(p)} for p in (kc[0], chi[0], dal[0])]
    doc = assemble(desc, data, own, built or dt.date.today().isoformat(), files)
    problems = check_doc(doc)
    if problems:
        raise SystemExit("not written:\n  " + "\n  ".join(problems))
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w", encoding="utf-8") as f:
        f.write(dumps(doc))
    cur = [d for d, x in doc["districts"].items() if x["current"]]
    print("farm credit: %s, %d current districts (%s), %d quarters withheld" % (doc["latest"], len(cur), ", ".join(cur), len(doc["withheld"])))
    return doc


def check_doc(doc):
    p = []
    if doc.get("schema") != SCHEMA:
        p.append("schema")
    lat = doc.get("latest") or ""
    if not Q.match(lat):
        p.append("latest quarter")
        return p
    n = qkey(lat) - qkey(doc["chart_from"]) + 1
    for d, x in doc["districts"].items():
        for k in ("repay", "renew"):
            if len(x["chart"][k]) != n:
                p.append("%s %s chart has %d quarters, expected %d" % (d, k, len(x["chart"][k]), n))
        if x["current"]:
            if not x["latest"] or x["latest"]["repay"] is None or x["latest"]["renew"] is None:
                p.append("%s current without both latest readings" % d)
            elif x["chart"]["repay"][-1] != x["latest"]["repay"] or x["chart"]["renew"][-1] != x["latest"]["renew"]:
                p.append("%s latest is not the chart's last point" % d)
            if not str(x["url"]).startswith("https://"):
                p.append("%s link" % d)
        for k in ("repay", "renew"):
            for v in x["chart"][k]:
                if v is not None and not (0 <= v <= 200):
                    p.append("%s %s value %r outside 0 to 200" % (d, k, v))
    for st, ds in doc["states"].items():
        for d in ds:
            if d not in doc["districts"] and d not in doc["no_survey"]:
                p.append("%s district %s unknown" % (st, d))
    txt = json.dumps(doc, ensure_ascii=False)
    if "—" in txt:
        p.append("em dash in the file")
    return p


def selftest():
    fails = []

    def check(name, cond):
        if not cond:
            fails.append(name)
            print("FAIL", name)
    check("quarter keys", qkey("2026Q2") - qkey("2010Q1") == 65 and qrange("2025Q3", "2026Q2") == ["2025Q3", "2025Q4", "2026Q1", "2026Q2"])
    check("label", qlabel("2026Q2") == "2026 Q2")
    check("float noise dropped, the table's value kept", num(85.71000000000001) == 85.71 and num(73) == 73 and num(124.0) == 124 and num("---") is None and num("n/a") is None)
    # Dallas own file prints 2 decimals, the Kansas City table sometimes 1 or 0 (real rows)
    check("rounding agrees", agree(111, 110.96) and agree(97.6, 97.65) and agree(96.15, 96.15) and agree(94.1, 94.12))
    check("a real disagreement does not", not agree(82, 64) and not agree(105.7, 115.68) and not agree(89, 86.72))
    check("missing both agrees, one missing does not", agree(None, None) and not agree(73, None))
    check("no em dash in the copy", all("—" not in s for s in [HOW, KC_PANEL] + list(WORDS.values())))
    check("every state and DC mapped once", len(STATE_DISTRICTS) == 51)
    # a tiny table run through assemble: a disagreement withholds the point, latest kept
    desc = {}
    for d, m in DISTRICTS.items():
        for k, (code, nm) in SERIES.items():
            desc[code + d] = {"name": nm, "unit": "Diffusion Index", "district": m["name"]}
    qs = qrange("2010Q1", "2010Q4")
    data = {"003CHI": dict(zip(qs, [90, 82, 70, 73])), "004CHI": dict(zip(qs, [110, 120, 130, 124])), "005CHI": dict(zip(qs, [1, 1, 1, 116])),
            "003SF": dict(zip(qs, [75, None, None, None])), "004SF": dict(zip(qs, [130, None, None, None]))}
    own = {"CHI": {"003": dict(zip(qs, [90, 64, 70, 73])), "004": dict(zip(qs, [110, 120, 130, 124]))}}
    doc = assemble(desc, data, own, "2026-10-28", [])
    c = doc["districts"]["CHI"]
    check("disagreeing quarter withheld", c["chart"]["repay"] == [90, None, 70, 73] and doc["withheld"] == [{"district": "CHI", "q": "2010Q2", "series": "repay", "kc": 82, "own": 64}])
    check("latest kept", c["current"] and c["latest"] == {"repay": 73, "renew": 124, "collat": 116})
    check("a district that stopped is not current", not doc["districts"]["SF"]["current"] and doc["districts"]["SF"]["last"] == "2010Q1" and doc["districts"]["SF"]["latest"] is None)
    own["CHI"]["003"]["2010Q4"] = 70
    doc2 = assemble(desc, data, own, "2026-10-28", [])
    check("latest disagreement withholds the district", not doc2["districts"]["CHI"]["current"] and doc2["districts"]["CHI"]["latest"] is None)
    print("selftest: %d failed" % len(fails) if fails else "selftest: all passed")
    return 1 if fails else 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", help="the dated folder from the raw branch")
    ap.add_argument("--out", default=OUT)
    ap.add_argument("--check", action="store_true")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        return selftest()
    if a.check:
        doc = json.load(open(a.out, encoding="utf-8"))
        p = check_doc(doc)
        print("\n".join(p) if p else "farm credit file: ok (%s)" % doc["latest"])
        return 1 if p else 0
    if not a.raw:
        ap.error("--raw is required to build")
    build(a.raw, a.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
